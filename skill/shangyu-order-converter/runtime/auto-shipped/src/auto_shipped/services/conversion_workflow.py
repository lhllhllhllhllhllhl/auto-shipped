from __future__ import annotations

import hashlib
from copy import deepcopy
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


WORKFLOW_SCHEMA_VERSION = "conversion-workflow/1.0"

_ORDER_SCOPED_RULES = {
    "GY-STORE-ASSIGNMENT",
    "GY-PLATFORM-ORDER-NUMBER",
    "GY-BUYER-MEMBER",
    "GY-PAYMENT-VALUES",
    "GY-WAREHOUSE-ASSIGNMENT",
    "GY-PAYMENT-METHOD",
    "GY-ORDER-TYPE",
    "GY-ORDER-FLAGS",
    "GY-RECIPIENT-COMPLETENESS",
    "GY-JD-ZTO-CONTACT-SUFFIX",
    "GY-LOGISTICS-SELECTION",
    "GY-SELLER-REMARK-POLICY",
    "GY-DELIVERY-INSTRUCTION-ADDRESS",
    "GY-DAILY-GOODS-WHITE-LABEL",
}
_ITEM_SCOPED_RULES = {
    "GY-PRODUCT-MAPPING",
    "GY-ITEM-EXPANSION",
}
_SUPPORTED_RULES = _ORDER_SCOPED_RULES | _ITEM_SCOPED_RULES
_COMPLETED_EVENT_STATUSES = {"completed", "not_applicable"}

_REQUIRED_NONEMPTY_FIELDS: dict[str, tuple[str, ...]] = {
    "GY-STORE-ASSIGNMENT": ("店铺",),
    "GY-PLATFORM-ORDER-NUMBER": ("平台单号",),
    "GY-BUYER-MEMBER": ("买家会员",),
    "GY-WAREHOUSE-ASSIGNMENT": ("仓库名称",),
    "GY-PAYMENT-METHOD": ("支付方式",),
    "GY-ORDER-TYPE": ("订单类型",),
    "GY-PRODUCT-MAPPING": ("商品代码",),
    "GY-RECIPIENT-COMPLETENESS": ("收货人", "收货地址"),
    "GY-LOGISTICS-SELECTION": ("物流公司",),
    "GY-DELIVERY-INSTRUCTION-ADDRESS": ("收货地址",),
}
_REQUIRED_PRESENT_FIELDS: dict[str, tuple[str, ...]] = {
    "GY-PAYMENT-VALUES": ("支付金额", "价格", "运费"),
    "GY-ORDER-FLAGS": ("是否手机订单", "是否货到付款", "是否分销商订单"),
    "GY-JD-ZTO-CONTACT-SUFFIX": ("联系电话", "联系手机"),
    "GY-SELLER-REMARK-POLICY": ("卖家备注",),
    "GY-DAILY-GOODS-WHITE-LABEL": ("卖家备注",),
    "GY-ITEM-EXPANSION": ("商品代码", "数量"),
}


class ConversionWorkflowError(ValueError):
    def __init__(
        self,
        code: str,
        message: str,
        *,
        details: dict[str, Any] | None = None,
    ) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.details = details or {}

    def to_dict(self) -> dict[str, Any]:
        return {
            "code": self.code,
            "message": self.message,
            **({"details": self.details} if self.details else {}),
        }


def supported_conversion_rule_ids() -> frozenset[str]:
    """Return conversion rules with an explicit runtime completion verifier."""

    return frozenset(_SUPPORTED_RULES)


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _source_item_scope_keys(orders: Iterable[Any]) -> set[str]:
    return {
        f"order:{order_index}/item:{item_index}"
        for order_index, order in enumerate(orders, start=1)
        for item_index, _item in enumerate(order.items, start=1)
    }


def _order_scope_keys(orders: Iterable[Any]) -> set[str]:
    return {
        f"order:{order_index}"
        for order_index, _order in enumerate(orders, start=1)
    }


def _validate_output_postconditions(
    required_rule_ids: list[str],
    lines: list[dict[str, Any]],
    *,
    expected_order_count: int,
) -> list[dict[str, Any]]:
    if not lines:
        raise ConversionWorkflowError(
            "CONVERSION_WORKFLOW_NO_OUTPUT_LINES",
            "业务规则已经执行，但没有形成可交付的管易明细行。",
        )
    platform_numbers = {
        str(line.get("平台单号") or "").strip()
        for line in lines
        if str(line.get("平台单号") or "").strip()
    }
    if len(platform_numbers) != expected_order_count:
        raise ConversionWorkflowError(
            "CONVERSION_WORKFLOW_ORDER_COVERAGE_MISMATCH",
            "业务结果未覆盖本批全部订单。",
            details={
                "expected_order_count": expected_order_count,
                "actual_order_count": len(platform_numbers),
            },
        )

    if "GY-RECIPIENT-COMPLETENESS" in required_rule_ids:
        missing_contact_rows = [
            row_number
            for row_number, line in enumerate(lines, start=1)
            if _is_blank(line.get("联系电话")) and _is_blank(line.get("联系手机"))
        ]
        if missing_contact_rows:
            raise ConversionWorkflowError(
                "CONVERSION_WORKFLOW_RECIPIENT_CONTACT_MISSING",
                "收件信息规则已执行，但成品存在没有联系电话和联系手机的订单行。",
                details={"output_row_numbers": missing_contact_rows},
            )

    checks: list[dict[str, Any]] = []
    for rule_id in required_rule_ids:
        missing_fields: set[str] = set()
        blank_fields: set[str] = set()
        for field in _REQUIRED_PRESENT_FIELDS.get(rule_id, ()):
            if any(field not in line for line in lines):
                missing_fields.add(field)
        for field in _REQUIRED_NONEMPTY_FIELDS.get(rule_id, ()):
            if any(field not in line for line in lines):
                missing_fields.add(field)
            elif any(_is_blank(line.get(field)) for line in lines):
                blank_fields.add(field)
        if missing_fields or blank_fields:
            raise ConversionWorkflowError(
                "CONVERSION_WORKFLOW_RULE_POSTCONDITION_FAILED",
                f"规则 {rule_id} 已登记执行，但输出结果未满足必要后置条件。",
                details={
                    "rule_id": rule_id,
                    "missing_fields": sorted(missing_fields),
                    "blank_fields": sorted(blank_fields),
                },
            )
        checks.append(
            {
                "rule_id": rule_id,
                "status": "completed",
                "checked_output_rows": len(lines),
                "checked_fields": sorted(
                    set(_REQUIRED_PRESENT_FIELDS.get(rule_id, ()))
                    | set(_REQUIRED_NONEMPTY_FIELDS.get(rule_id, ()))
                ),
            }
        )

    for row_number, line in enumerate(lines, start=1):
        quantity = line.get("数量")
        try:
            valid_quantity = float(quantity) > 0
        except (TypeError, ValueError):
            valid_quantity = False
        if not valid_quantity:
            raise ConversionWorkflowError(
                "CONVERSION_WORKFLOW_INVALID_OUTPUT_QUANTITY",
                "管易明细行存在无效数量。",
                details={"output_row_number": row_number},
            )
    return checks


def _business_summary(lines: list[dict[str, Any]]) -> dict[str, Any]:
    orders: dict[str, dict[str, str]] = {}
    products: dict[tuple[str, str], float] = {}
    for line in lines:
        order_no = str(line.get("平台单号") or "").strip()
        if order_no:
            orders.setdefault(
                order_no,
                {
                    "store": str(line.get("店铺") or "").strip(),
                    "logistics": str(line.get("物流公司") or "").strip(),
                },
            )
        product_key = (
            str(line.get("商品代码") or "").strip(),
            str(line.get("规格代码") or "").strip(),
        )
        try:
            quantity = float(line.get("数量"))
        except (TypeError, ValueError):
            quantity = 0.0
        products[product_key] = products.get(product_key, 0.0) + quantity

    def counted(field: str) -> list[dict[str, Any]]:
        counts: dict[str, int] = {}
        for order in orders.values():
            value = order[field]
            counts[value] = counts.get(value, 0) + 1
        return [
            {"value": value, "order_count": count}
            for value, count in sorted(counts.items())
        ]

    return {
        "stores": counted("store"),
        "logistics": counted("logistics"),
        "products": [
            {
                "product_code": product_code,
                "spec_code": spec_code,
                "quantity": int(quantity) if quantity.is_integer() else quantity,
            }
            for (product_code, spec_code), quantity in sorted(products.items())
        ],
    }


def build_conversion_workflow(
    *,
    orders: list[Any],
    required_rule_ids: list[str],
    rule_events: list[dict[str, Any]],
    lines: list[dict[str, Any]],
    source_profile_id: str,
    company_id: str,
    route_id: str,
    mapping_revision: str | int | None = None,
) -> dict[str, Any]:
    """Enforce the per-batch execution directory before any workbook is delivered."""

    unsupported = sorted(set(required_rule_ids) - _SUPPORTED_RULES)
    if unsupported:
        raise ConversionWorkflowError(
            "CONVERSION_WORKFLOW_RULE_CHECK_MISSING",
            "规则包包含尚未登记运行校验器的转换规则。",
            details={"rule_ids": unsupported},
        )

    order_keys = _order_scope_keys(orders)
    item_keys = _source_item_scope_keys(orders)
    event_checks: list[dict[str, Any]] = []
    for rule_id in required_rule_ids:
        expected_keys = item_keys if rule_id in _ITEM_SCOPED_RULES else order_keys
        matching = [event for event in rule_events if event.get("rule_id") == rule_id]
        actual_keys = {str(event.get("scope_key") or "") for event in matching}
        incomplete = [
            event
            for event in matching
            if str(event.get("status") or "") not in _COMPLETED_EVENT_STATUSES
        ]
        if actual_keys != expected_keys or incomplete:
            raise ConversionWorkflowError(
                "CONVERSION_WORKFLOW_RULE_EXECUTION_INCOMPLETE",
                f"规则 {rule_id} 未对本批全部适用对象完成判断。",
                details={
                    "rule_id": rule_id,
                    "expected_scope_count": len(expected_keys),
                    "actual_scope_count": len(actual_keys),
                    "missing_scope_keys": sorted(expected_keys - actual_keys),
                    "unexpected_scope_keys": sorted(actual_keys - expected_keys),
                },
            )
        event_checks.append(
            {
                "rule_id": rule_id,
                "status": "completed",
                "scope": "source_item" if rule_id in _ITEM_SCOPED_RULES else "order",
                "scope_count": len(expected_keys),
                "outcomes": sorted({str(event.get("outcome") or "") for event in matching}),
            }
        )

    output_checks = _validate_output_postconditions(
        required_rule_ids,
        lines,
        expected_order_count=len(orders),
    )
    workflow_identity = "|".join(
        [source_profile_id, company_id, route_id, *required_rule_ids]
    )
    input_fingerprints = sorted(
        {
            str(getattr(getattr(order, "source", None), "file_fingerprint", "") or "")
            for order in orders
            if str(
                getattr(getattr(order, "source", None), "file_fingerprint", "") or ""
            )
        }
    )
    return {
        "schema_version": WORKFLOW_SCHEMA_VERSION,
        "workflow_id": hashlib.sha256(workflow_identity.encode("utf-8")).hexdigest()[:16],
        "status": "pending_output_verification",
        "context": {
            "source_profile_id": source_profile_id,
            "company_id": company_id,
            "route_id": route_id,
            "mapping_revision": mapping_revision,
            "input_fingerprints": input_fingerprints,
        },
        "counts": {
            "orders": len(orders),
            "source_items": len(item_keys),
            "output_rows": len(lines),
        },
        "business_summary": _business_summary(lines),
        "required_rule_ids": list(required_rule_ids),
        "stages": [
            {"stage_id": "source_intake", "status": "completed"},
            {"stage_id": "company_workflow_routing", "status": "completed"},
            {
                "stage_id": "business_rule_directory",
                "status": "completed",
                "rule_checks": event_checks,
                "postcondition_checks": output_checks,
            },
            {"stage_id": "output_render_readback", "status": "pending"},
            {"stage_id": "platform_preflight", "status": "pending"},
        ],
    }


def _normalized_cell(value: Any) -> Any:
    if value is None:
        return ""
    return value


def verify_rendered_output(
    workbook_path: str | Path,
    *,
    expected_lines: list[dict[str, Any]],
    platform_profile: dict[str, Any],
) -> dict[str, Any]:
    """Re-read the saved workbook and compare every business cell with the plan."""

    path = Path(workbook_path)
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        sheet = workbook[str(platform_profile["sheet_name"])]
        header_row = int(platform_profile["header_row"])
        columns = list(platform_profile["columns"])
        actual_lines: list[dict[str, Any]] = []
        for row_no in range(header_row + 1, sheet.max_row + 1):
            values = [
                sheet.cell(row_no, column_no).value
                for column_no in range(1, len(columns) + 1)
            ]
            if not any(not _is_blank(value) for value in values):
                continue
            actual_lines.append(dict(zip(columns, values, strict=True)))
    finally:
        workbook.close()

    if len(actual_lines) != len(expected_lines):
        raise ConversionWorkflowError(
            "CONVERSION_WORKFLOW_RENDERED_ROW_COUNT_MISMATCH",
            "成品Excel回读行数与确定性业务结果不一致。",
            details={
                "expected_rows": len(expected_lines),
                "actual_rows": len(actual_lines),
            },
        )
    for row_number, (expected, actual) in enumerate(
        zip(expected_lines, actual_lines, strict=True), start=1
    ):
        different_fields = [
            column
            for column in columns
            if _normalized_cell(expected.get(column))
            != _normalized_cell(actual.get(column))
        ]
        if different_fields:
            raise ConversionWorkflowError(
                "CONVERSION_WORKFLOW_RENDERED_VALUE_MISMATCH",
                "成品Excel回读内容与确定性业务结果不一致。",
                details={
                    "output_row_number": row_number,
                    "different_fields": different_fields,
                },
            )

    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    return {
        "status": "completed",
        "row_count": len(actual_lines),
        "column_count": len(platform_profile["columns"]),
        "output_sha256": digest,
    }


def complete_conversion_workflow(
    workflow: dict[str, Any],
    *,
    readback: dict[str, Any],
    preflight_checks: dict[str, bool],
) -> dict[str, Any]:
    failed_checks = sorted(name for name, passed in preflight_checks.items() if not passed)
    if failed_checks:
        raise ConversionWorkflowError(
            "CONVERSION_WORKFLOW_PREFLIGHT_INCOMPLETE",
            "平台预检未全部通过，不能交付成品。",
            details={"failed_checks": failed_checks},
        )
    result = deepcopy(workflow)
    for stage in result["stages"]:
        if stage["stage_id"] == "output_render_readback":
            stage.update(readback)
        elif stage["stage_id"] == "platform_preflight":
            stage.update(
                {
                    "status": "completed",
                    "check_count": len(preflight_checks),
                    "checks": dict(preflight_checks),
                }
            )
    result["status"] = "complete"
    return result


def build_batch_workflow(
    *,
    source_workflows: list[dict[str, Any]],
    expected_lines: list[dict[str, Any]],
) -> dict[str, Any]:
    incomplete = [
        index
        for index, workflow in enumerate(source_workflows, start=1)
        if workflow.get("status") != "complete"
    ]
    if incomplete:
        raise ConversionWorkflowError(
            "BATCH_CONVERSION_WORKFLOW_SOURCE_INCOMPLETE",
            "批次包含尚未完成内部执行目录的来源文件。",
            details={"source_indexes": incomplete},
        )
    return {
        "schema_version": WORKFLOW_SCHEMA_VERSION,
        "workflow_id": hashlib.sha256(
            "|".join(str(item.get("workflow_id") or "") for item in source_workflows).encode(
                "utf-8"
            )
        ).hexdigest()[:16],
        "status": "pending_output_verification",
        "counts": {"sources": len(source_workflows), "output_rows": len(expected_lines)},
        "business_summary": _business_summary(expected_lines),
        "source_workflow_ids": [item.get("workflow_id") for item in source_workflows],
        "stages": [
            {"stage_id": "source_workflows", "status": "completed"},
            {"stage_id": "batch_merge", "status": "completed"},
            {"stage_id": "output_render_readback", "status": "pending"},
            {"stage_id": "platform_preflight", "status": "pending"},
        ],
    }
