from __future__ import annotations

import hashlib
import json
import math
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from openpyxl import Workbook

from auto_shipped.companies import (
    CompanyRegistryError,
    load_company_registry,
    resolve_company_abbreviation_references,
    resolve_company_by_identity,
)
from auto_shipped.domain import ClarificationRequest, deduplicate_clarifications
from auto_shipped.services.order_number_allocations import reserve_minute_sequence
from auto_shipped.source_adapters.adaptive_excel import inspect_excel_structure
from auto_shipped.source_adapters.value_validation import normalize_contact


TEXT_DRAFT_SCHEMA_VERSION = "1.0"
TEXT_DRAFT_TYPE = "agent_text_order_draft"
TEXT_DRAFT_CONFIRMATION = "CONFIRM_TEXT_ORDER_DRAFT"
TEXT_ORDER_NUMBER_CONFIRMATION = "CONFIRM_TEXT_ORDER_NUMBER_GENERATION"
PROVENANCE_VALUES = {"explicit", "user_confirmed", "inferred", "missing"}
CRITICAL_PROVENANCE_FIELDS = (
    "source_channel",
    "source_order_no",
    "recipient_name",
    "contact",
    "address",
    "items",
)
RESERVED_EXTENSION_KEYS = {
    "adaptive_plan_id",
    "adaptive_structure_signature",
    "source_channel",
    "source_order_no",
    "recipient_name",
    "contact",
    "province",
    "city",
    "district",
    "address",
    "product_code",
    "product_name",
    "source_spec",
    "quantity",
    "unit",
    "carrier",
    "source_note",
    "ordered_at",
}


class TextOrderDraftError(ValueError):
    """Raised when the text-order draft file itself cannot be read safely."""


@dataclass(slots=True)
class TextOrderIntakeResult:
    status: str
    draft_file: str
    order_count: int = 0
    company_id: str | None = None
    rule_source_profile_id: str | None = None
    masked_preview: list[dict[str, Any]] = field(default_factory=list)
    clarifications: list[ClarificationRequest] = field(default_factory=list)
    outputs: list[str] = field(default_factory=list)
    order_number_allocation: dict[str, Any] | None = None

    def to_public_dict(self) -> dict[str, Any]:
        return {
            "status": self.status,
            "draft_file": self.draft_file,
            "order_count": self.order_count,
            "company_id": self.company_id,
            "rule_source_profile_id": self.rule_source_profile_id,
            "masked_preview": self.masked_preview,
            "clarifications": [asdict(item) for item in self.clarifications],
            "outputs": self.outputs,
            "order_number_allocation": self.order_number_allocation,
            "privacy": "recipient_fields_masked_raw_text_not_returned",
        }


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _load_draft(path: str | Path) -> tuple[Path, dict[str, Any]]:
    draft_path = Path(path).expanduser().resolve()
    try:
        payload = json.loads(draft_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise TextOrderDraftError(f"无法读取文字订单草稿：{exc}") from exc
    if not isinstance(payload, dict):
        raise TextOrderDraftError("文字订单草稿根节点必须是对象。")
    if payload.get("schema_version") != TEXT_DRAFT_SCHEMA_VERSION:
        raise TextOrderDraftError("文字订单草稿schema_version必须是1.0。")
    if payload.get("draft_type") != TEXT_DRAFT_TYPE:
        raise TextOrderDraftError("文字订单草稿draft_type不正确。")
    if not _text(payload.get("draft_id")):
        raise TextOrderDraftError("文字订单草稿缺少draft_id。")
    return draft_path, payload


def _mask_name(value: Any) -> str:
    text = _text(value)
    return "" if not text else text[:1] + "**"


def _mask_contact(value: Any) -> str:
    digits = re.sub(r"\D", "", _text(value))
    if len(digits) >= 7:
        return f"{digits[:3]}****{digits[-4:]}"
    return "***" if digits else ""


def _mask_address(value: Any) -> str:
    text = _text(value)
    if not text:
        return ""
    return text[:6] + "…" if len(text) > 6 else text[:2] + "…"


def _positive_number(value: Any) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) and number > 0 else None


def _today_in_timezone(timezone_name: str) -> date | None:
    try:
        return datetime.now(ZoneInfo(timezone_name)).date()
    except ZoneInfoNotFoundError:
        return None


def _now_in_timezone(timezone_name: str) -> datetime | None:
    try:
        return datetime.now(ZoneInfo(timezone_name))
    except ZoneInfoNotFoundError:
        return None


def _active_workflow_profiles(
    registry: dict[str, Any], company_id: str
) -> list[str]:
    company = next(
        (
            item
            for item in registry.get("companies", [])
            if item.get("company_id") == company_id
            and item.get("enabled", True) is not False
            and item.get("status") == "active"
        ),
        None,
    )
    if company is None:
        return []
    return sorted(
        {
            _text(item.get("source_profile_id"))
            for item in company.get("workflows", [])
            if item.get("runtime_enabled") is True
            and _text(item.get("source_profile_id"))
        }
    )


def _resolve_profile(
    payload: dict[str, Any],
    registry: dict[str, Any],
    requests: list[ClarificationRequest],
) -> tuple[str | None, str | None]:
    company_identity = _text(payload.get("company_id")) or None
    if company_identity is None:
        requests.append(
            ClarificationRequest(
                code="TEXT_ORDER_COMPANY_REQUIRED",
                scope="batch",
                field="company_id",
                question="这些文字订单属于哪家公司或业务来源？",
                reason="平台名称不能替代公司来源；不同公司的店铺、单号、商品和物流规则不同。",
                answer_type="text",
            )
        )
        return None, None

    identity_resolution = resolve_company_by_identity(registry, company_identity)
    company_id = (
        identity_resolution.company.company_id
        if identity_resolution.company is not None
        else company_identity
    )
    if identity_resolution.status != "resolved":
        requests.append(
            ClarificationRequest(
                code=identity_resolution.code or "TEXT_ORDER_COMPANY_WORKFLOW_UNAVAILABLE",
                scope="batch",
                field="company_id",
                question="该公司尚未登记可用于文字订单的业务流程，请先补充公司规则。",
                reason=identity_resolution.reason or f"company={company_identity}",
                answer_type="text",
            )
        )
        return company_id, None

    profiles = _active_workflow_profiles(registry, company_id)
    if not profiles:
        requests.append(
            ClarificationRequest(
                code="TEXT_ORDER_COMPANY_WORKFLOW_UNAVAILABLE",
                scope="batch",
                field="company_id",
                question="该公司还没有可用于文字订单的已启用业务规则，请先确认来源或补充公司流程。",
                reason=f"company_id={company_id}",
                answer_type="text",
            )
        )
        return company_id, None

    requested = _text(payload.get("rule_source_profile_id"))
    if requested:
        if requested not in profiles:
            requests.append(
                ClarificationRequest(
                    code="TEXT_ORDER_RULE_PROFILE_INVALID",
                    scope="batch",
                    field="rule_source_profile_id",
                    question="指定的业务规则不属于该公司，请确认订单类型。",
                    reason=f"可用规则：{', '.join(profiles)}",
                    answer_type="single_choice",
                    choices=tuple(profiles),
                )
            )
            return company_id, None
        return company_id, requested
    if len(profiles) > 1:
        requests.append(
            ClarificationRequest(
                code="TEXT_ORDER_RULE_PROFILE_REQUIRED",
                scope="batch",
                field="rule_source_profile_id",
                question="该公司有多种订单流程，请确认这批文字订单使用哪一种。",
                reason="候选来源规则：" + ", ".join(profiles),
                answer_type="single_choice",
                choices=tuple(profiles),
            )
        )
        return company_id, None
    return company_id, profiles[0]


def _load_source_profile(
    company_registry_path: str | Path,
    profile_id: str | None,
) -> dict[str, Any]:
    if not profile_id:
        return {}
    source_profiles_dir = Path(company_registry_path).resolve().parents[1] / "source_profiles"
    for path in sorted(source_profiles_dir.rglob("*.json")):
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        if payload.get("profile_id") == profile_id:
            return payload
    return {}


def _date_sequence_policy(
    payload: dict[str, Any],
    requests: list[ClarificationRequest],
    source_profile: dict[str, Any] | None = None,
) -> dict[str, Any] | None:
    explicit_policy = payload.get("order_number_policy")
    configured_policy = (
        ((source_profile or {}).get("text_intake") or {}).get("order_number_policy")
        if explicit_policy is None
        else None
    )
    policy = explicit_policy or configured_policy or {"strategy": "source"}
    if not isinstance(policy, dict):
        policy = {}
    strategy = _text(policy.get("strategy")) or "source"
    if strategy == "source":
        return {"strategy": "source"}
    if strategy in {"current_date_sequence", "current_datetime_sequence"}:
        if policy.get("confirmed") is not True:
            requests.append(
                ClarificationRequest(
                    code="TEXT_ORDER_NUMBER_POLICY_UNCONFIRMED",
                    scope="batch",
                    field="order_number_policy",
                    question="当前公司的当日顺序号规则尚未确认，请先确认。",
                    reason="来源配置中的自动单号策略未标记为已确认。",
                    answer_type="confirmation",
                )
            )
            return None
        timezone_name = _text(policy.get("timezone")) or "Asia/Shanghai"
        batch_datetime = _now_in_timezone(timezone_name)
        try:
            start = int(policy.get("sequence_start"))
            width = int(policy.get("minimum_sequence_width"))
        except (TypeError, ValueError):
            start = -1
            width = 0
        prefix = _text(policy.get("prefix"))
        configured_format = _text(
            policy.get("datetime_format") or policy.get("date_format")
        )
        stamp_format = configured_format or (
            "%Y%m%d%H%M"
            if strategy == "current_datetime_sequence"
            else "%Y%m%d"
        )
        allowed_formats = (
            {"%Y%m%d%H%M", "%y%m%d%H%M"}
            if strategy == "current_datetime_sequence"
            else {"%Y%m%d", "%y%m%d"}
        )
        if (
            batch_datetime is None
            or start < 0
            or width < 1
            or width > 8
            or stamp_format not in allowed_formats
            or not re.fullmatch(r"[0-9A-Za-z_-]{0,16}", prefix)
        ):
            requests.append(
                ClarificationRequest(
                    code="TEXT_ORDER_NUMBER_POLICY_INVALID",
                    scope="batch",
                    field="order_number_policy",
                    question="当前公司的时间顺序号配置无效，请修复来源规则。",
                    reason="前缀、时间格式、起始序号、最小宽度或时区未通过校验。",
                    answer_type="text",
                )
            )
            return None
        return {
            "strategy": (
                "datetime_sequence"
                if strategy == "current_datetime_sequence"
                else "date_sequence"
            ),
            "batch_datetime": batch_datetime,
            "business_date": batch_datetime.date(),
            "prefix": prefix,
            "stamp_format": stamp_format,
            "date_format": stamp_format,
            "sequence_start": start,
            "sequence_width": width,
            "sequence_scope": _text(policy.get("sequence_scope")) or "current_batch",
            "timezone": timezone_name,
        }
    if strategy != "date_sequence":
        requests.append(
            ClarificationRequest(
                code="TEXT_ORDER_NUMBER_POLICY_INVALID",
                scope="batch",
                field="order_number_policy",
                question="文字订单的单号策略无法识别，请选择使用来源单号或确认日期顺序号规则。",
                reason=f"strategy={strategy!r}",
                answer_type="text",
            )
        )
        return None
    if policy.get("confirmation_token") != TEXT_ORDER_NUMBER_CONFIRMATION:
        requests.append(
            ClarificationRequest(
                code="TEXT_ORDER_NUMBER_GENERATION_CONFIRMATION_REQUIRED",
                scope="batch",
                field="order_number_policy",
                question="部分文字订单没有来源单号。是否确认按指定日期、前缀和顺序生成唯一来源单号？",
                reason="系统不能自行决定缺失订单号的生成规则。",
                answer_type="confirmation",
            )
        )
        return None
    try:
        business_date = date.fromisoformat(_text(policy.get("business_date")))
        start = int(policy.get("sequence_start"))
        width = int(policy.get("sequence_width"))
    except (TypeError, ValueError):
        business_date = None
        start = -1
        width = 0
    prefix = _text(policy.get("prefix"))
    date_format = _text(policy.get("date_format")) or "%y%m%d"
    if (
        business_date is None
        or start < 0
        or width < 1
        or width > 8
        or date_format not in {"%Y%m%d", "%y%m%d"}
        or not re.fullmatch(r"[0-9A-Za-z_-]{0,16}", prefix)
    ):
        requests.append(
            ClarificationRequest(
                code="TEXT_ORDER_NUMBER_POLICY_INVALID",
                scope="batch",
                field="order_number_policy",
                question="日期顺序号参数不完整，请确认日期、前缀、起始序号和位数。",
                reason="单号参数未通过格式校验。",
                answer_type="text",
            )
        )
        return None
    return {
        "strategy": "date_sequence",
        "business_date": business_date,
        "prefix": prefix,
        "date_format": date_format,
        "sequence_start": start,
        "sequence_width": width,
    }


def _generated_order_number(policy: dict[str, Any], index: int) -> str:
    if policy.get("strategy") == "datetime_sequence":
        stamp = policy["batch_datetime"].strftime(policy["stamp_format"])
    else:
        business_date: date = policy["business_date"]
        stamp = business_date.strftime(policy["date_format"])
    return (
        f"{policy['prefix']}{stamp}"
        f"{policy['sequence_start'] + index:0{policy['sequence_width']}d}"
    )


def validate_text_order_draft(
    draft_path: str | Path,
    *,
    company_registry_path: str | Path,
    require_confirmation: bool = True,
) -> tuple[TextOrderIntakeResult, dict[str, Any] | None]:
    path, payload = _load_draft(draft_path)
    requests: list[ClarificationRequest] = []
    source = payload.get("source") or {}
    if not isinstance(source, dict) or source.get("kind") not in {
        "plain_text",
        "image_ocr",
    }:
        requests.append(
            ClarificationRequest(
                code="TEXT_ORDER_SOURCE_INVALID",
                scope="file",
                question="文字订单来源类型无效，请重新提取当前文字或截图。",
                reason="source.kind必须为plain_text或image_ocr。",
                answer_type="file",
            )
        )
    if not re.fullmatch(r"[0-9a-f]{64}", _text(source.get("sha256"))):
        requests.append(
            ClarificationRequest(
                code="TEXT_ORDER_SOURCE_HASH_REQUIRED",
                scope="file",
                question="缺少原始文字或截图的SHA-256，请重新生成草稿。",
                reason="草稿必须绑定本次输入，避免确认后输入被替换。",
                answer_type="file",
            )
        )

    try:
        registry = load_company_registry(company_registry_path)
    except (OSError, json.JSONDecodeError, CompanyRegistryError) as exc:
        raise TextOrderDraftError(f"无法读取公司库：{exc}") from exc
    company_id, rule_profile = _resolve_profile(payload, registry, requests)
    source_profile = _load_source_profile(company_registry_path, rule_profile)
    if company_id and source_profile:
        identity_resolution = resolve_company_by_identity(registry, company_id)
        if (
            identity_resolution.status == "resolved"
            and identity_resolution.company is not None
        ):
            try:
                source_profile = resolve_company_abbreviation_references(
                    source_profile,
                    identity_resolution.company,
                )
            except CompanyRegistryError as exc:
                requests.append(
                    ClarificationRequest(
                        code="TEXT_ORDER_COMPANY_ABBREVIATION_INVALID",
                        scope="batch",
                        field="company_id",
                        question="当前公司简称无法用于文字订单编号，请先修复公司库。",
                        reason=str(exc),
                        answer_type="text",
                    )
                )
    order_number_policy = _date_sequence_policy(payload, requests, source_profile)

    orders = payload.get("orders")
    if not isinstance(orders, list) or not orders:
        requests.append(
            ClarificationRequest(
                code="TEXT_ORDER_ROWS_REQUIRED",
                scope="file",
                question="没有识别到任何文字订单，请重新提供文字或截图。",
                reason="orders必须至少包含一项。",
                answer_type="file",
            )
        )
        orders = []

    seen_refs: set[str] = set()
    seen_order_numbers: set[str] = set()
    masked_preview: list[dict[str, Any]] = []
    materialized_orders: list[dict[str, Any]] = []
    for index, raw_order in enumerate(orders):
        if not isinstance(raw_order, dict):
            requests.append(
                ClarificationRequest(
                    code="TEXT_ORDER_ROW_INVALID",
                    scope="order",
                    source_ref=f"order-{index + 1}",
                    question="存在无法识别的订单记录，请重新提取。",
                    reason="订单记录必须是对象。",
                    answer_type="file",
                )
            )
            continue
        order_ref = _text(raw_order.get("order_ref")) or f"order-{index + 1}"
        if order_ref in seen_refs:
            requests.append(
                ClarificationRequest(
                    code="TEXT_ORDER_REF_DUPLICATED",
                    scope="order",
                    source_ref=order_ref,
                    question="文字订单内部编号重复，请重新编号。",
                    reason="order_ref必须唯一。",
                    answer_type="file",
                )
            )
        seen_refs.add(order_ref)

        provenance = raw_order.get("field_provenance") or {}
        if not isinstance(provenance, dict):
            provenance = {}
        for field_name in CRITICAL_PROVENANCE_FIELDS:
            value = _text(provenance.get(field_name)) or "missing"
            if value not in PROVENANCE_VALUES:
                value = "missing"
            if field_name == "source_order_no" and (
                _text(raw_order.get("source_order_no"))
                or (
                    order_number_policy is not None
                    and order_number_policy.get("strategy")
                    in {"date_sequence", "datetime_sequence"}
                )
            ):
                continue
            if value in {"inferred", "missing"}:
                requests.append(
                    ClarificationRequest(
                        code="TEXT_ORDER_CRITICAL_FIELD_UNCONFIRMED",
                        scope="field",
                        source_ref=order_ref,
                        field=field_name,
                        question=f"订单{order_ref}的{field_name}不是原文明确值，请确认。",
                        reason=f"field_provenance={value}",
                        answer_type="text",
                    )
                )

        channel = _text(raw_order.get("source_channel"))
        if not channel:
            requests.append(
                ClarificationRequest(
                    code="TEXT_ORDER_CHANNEL_REQUIRED",
                    scope="order",
                    source_ref=order_ref,
                    field="source_channel",
                    question=f"订单{order_ref}来自哪个平台或业务渠道？",
                    reason="文字订单中的平台标签会影响店铺和业务规则，不能省略。",
                    answer_type="text",
                )
            )

        source_order_no = _text(raw_order.get("source_order_no"))
        source_order_no_generated = False
        if not source_order_no and order_number_policy is not None:
            if order_number_policy.get("strategy") in {
                "date_sequence",
                "datetime_sequence",
            }:
                source_order_no = _generated_order_number(order_number_policy, index)
                source_order_no_generated = True
            else:
                requests.append(
                    ClarificationRequest(
                        code="TEXT_ORDER_NUMBER_REQUIRED",
                        scope="order",
                        source_ref=order_ref,
                        field="source_order_no",
                        question=f"订单{order_ref}没有可确认的来源单号，请补充或确认统一生成规则。",
                        reason="平台单号必须可唯一追溯且不能由Agent临时猜测。",
                        answer_type="text",
                    )
                )
        if re.search(r"\d(?:\.\d+)?[eE][+-]?\d+", source_order_no):
            requests.append(
                ClarificationRequest(
                    code="TEXT_ORDER_NUMBER_SCIENTIFIC_NOTATION",
                    scope="order",
                    source_ref=order_ref,
                    field="source_order_no",
                    question=f"订单{order_ref}的来源单号疑似科学计数法，请提供完整文本单号。",
                    reason="科学计数法可能已经丢失订单号精度。",
                    answer_type="text",
                )
            )
        if source_order_no:
            if source_order_no in seen_order_numbers:
                requests.append(
                    ClarificationRequest(
                        code="TEXT_ORDER_NUMBER_DUPLICATED",
                        scope="order",
                        source_ref=order_ref,
                        field="source_order_no",
                        question="多张文字订单使用了同一个来源单号，请确认是否属于同一订单的多个商品。",
                        reason="当前草稿按每个order_ref作为独立订单。",
                        answer_type="confirmation",
                    )
                )
            seen_order_numbers.add(source_order_no)

        recipient = raw_order.get("recipient") or {}
        if not isinstance(recipient, dict):
            recipient = {}
        for field_name, label in (
            ("name", "收件人"),
            ("contact", "联系方式"),
            ("address", "收货地址"),
        ):
            if not _text(recipient.get(field_name)):
                requests.append(
                    ClarificationRequest(
                        code="TEXT_ORDER_RECIPIENT_FIELD_REQUIRED",
                        scope="field",
                        source_ref=order_ref,
                        field=field_name,
                        question=f"订单{order_ref}缺少{label}，请补充。",
                        reason="管易发货订单必须具备完整收件信息。",
                        answer_type="text",
                    )
                )
        contact = _text(recipient.get("contact"))
        if contact and normalize_contact(contact) is None:
            requests.append(
                ClarificationRequest(
                    code="TEXT_ORDER_CONTACT_INVALID",
                    scope="field",
                    source_ref=order_ref,
                    field="contact",
                    question=f"订单{order_ref}的联系方式无法识别，请确认完整手机号或座机号。",
                    reason="联系方式必须通过确定性格式校验。",
                    answer_type="text",
                )
            )

        raw_items = raw_order.get("items")
        if not isinstance(raw_items, list) or not raw_items:
            requests.append(
                ClarificationRequest(
                    code="TEXT_ORDER_ITEMS_REQUIRED",
                    scope="order",
                    source_ref=order_ref,
                    field="items",
                    question=f"订单{order_ref}缺少商品和数量，请补充。",
                    reason="无法进行商品映射和物流判断。",
                    answer_type="text",
                )
            )
            raw_items = []
        valid_items: list[dict[str, Any]] = []
        for item_index, raw_item in enumerate(raw_items, start=1):
            if not isinstance(raw_item, dict):
                continue
            code = _text(raw_item.get("source_product_code"))
            name = _text(raw_item.get("source_product_name"))
            quantity = _positive_number(raw_item.get("quantity"))
            unit = _text(raw_item.get("unit"))
            source_spec = _text(raw_item.get("source_spec"))
            if not source_spec and unit in {"根", "袋", "裸棒", "彩袋"}:
                source_spec = unit
            if not code and not name:
                requests.append(
                    ClarificationRequest(
                        code="TEXT_ORDER_PRODUCT_REQUIRED",
                        scope="field",
                        source_ref=order_ref,
                        field=f"items[{item_index}]",
                        question=f"订单{order_ref}的第{item_index}个商品缺少名称或编码，请确认。",
                        reason="无法进行正式商品映射。",
                        answer_type="text",
                    )
                )
            if quantity is None:
                requests.append(
                    ClarificationRequest(
                        code="TEXT_ORDER_QUANTITY_INVALID",
                        scope="field",
                        source_ref=order_ref,
                        field=f"items[{item_index}].quantity",
                        question=f"订单{order_ref}的第{item_index}个商品数量无效，请确认。",
                        reason="商品数量必须是大于0的数字。",
                        answer_type="text",
                    )
                )
            valid_items.append(
                {
                    "source_product_code": code,
                    "source_product_name": name,
                    "source_spec": source_spec,
                    "quantity": quantity or 0,
                    "unit": unit,
                }
            )

        uncertainties = raw_order.get("uncertainties") or []
        if not isinstance(uncertainties, list):
            uncertainties = []
        for uncertainty in uncertainties:
            if not isinstance(uncertainty, dict):
                continue
            requests.append(
                ClarificationRequest(
                    code="TEXT_ORDER_UNCERTAINTY_REQUIRES_CONFIRMATION",
                    scope="field",
                    source_ref=order_ref,
                    field=_text(uncertainty.get("field")) or None,
                    question=_text(uncertainty.get("question"))
                    or f"订单{order_ref}存在待确认信息。",
                    reason=_text(uncertainty.get("reason")) or "Agent标记为不确定。",
                    answer_type="text",
                )
            )

        extensions = raw_order.get("source_extensions") or {}
        if not isinstance(extensions, dict):
            extensions = {}
        cleaned_extensions: dict[str, str] = {}
        for key, value in extensions.items():
            clean_key = _text(key)
            if (
                not re.fullmatch(r"[a-z][a-z0-9_]{0,63}", clean_key)
                or clean_key in RESERVED_EXTENSION_KEYS
            ):
                requests.append(
                    ClarificationRequest(
                        code="TEXT_ORDER_EXTENSION_KEY_INVALID",
                        scope="field",
                        source_ref=order_ref,
                        field=clean_key or "source_extensions",
                        question="文字订单扩展字段名称不合法，请使用稳定的小写英文标识。",
                        reason="扩展字段只能使用小写字母、数字和下划线。",
                        answer_type="file",
                    )
                )
                continue
            cleaned_extensions[clean_key] = _text(value)

        masked_preview.append(
            {
                "order_ref": order_ref,
                "source_channel": channel,
                "source_order_no": "present" if source_order_no else "missing",
                "recipient": {
                    "name": _mask_name(recipient.get("name")),
                    "contact": _mask_contact(contact),
                    "address": _mask_address(recipient.get("address")),
                },
                "items": [
                    {
                        "product": item["source_product_name"]
                        or item["source_product_code"],
                        "spec": item["source_spec"],
                        "quantity": item["quantity"],
                        "unit": item["unit"],
                    }
                    for item in valid_items
                ],
                "source_carrier": _text(raw_order.get("source_carrier")),
            }
        )
        materialized_orders.append(
            {
                "order_ref": order_ref,
                "source_channel": channel,
                "source_order_no": source_order_no,
                "_source_order_no_generated": source_order_no_generated,
                "_order_index": index,
                "recipient": {
                    "name": _text(recipient.get("name")),
                    "contact": contact,
                    "province": _text(recipient.get("province")),
                    "city": _text(recipient.get("city")),
                    "district": _text(recipient.get("district")),
                    "address": _text(recipient.get("address")),
                },
                "items": valid_items,
                "source_carrier": _text(raw_order.get("source_carrier")),
                "source_note": _text(raw_order.get("source_note")),
                "ordered_at": _text(raw_order.get("ordered_at")),
                "source_extensions": cleaned_extensions,
            }
        )

    if not requests and require_confirmation and (
        (payload.get("confirmation") or {}).get("token") != TEXT_DRAFT_CONFIRMATION
    ):
        requests.append(
            ClarificationRequest(
                code="TEXT_ORDER_DRAFT_CONFIRMATION_REQUIRED",
                scope="batch",
                question="请核对脱敏预览中的订单数、平台、商品与数量；确认无误后再生成转换输入。",
                reason="AI文字识别结果必须经过一次人确认，不能直接进入管易Excel。",
                answer_type="confirmation",
                next_action=(
                    "用户确认后将confirmation.token设为"
                    f"{TEXT_DRAFT_CONFIRMATION}并重新校验。"
                ),
            )
        )

    requests = deduplicate_clarifications(requests)
    result = TextOrderIntakeResult(
        status="needs_input" if requests else "ready",
        draft_file=path.name,
        order_count=len(orders),
        company_id=company_id,
        rule_source_profile_id=rule_profile,
        masked_preview=masked_preview,
        clarifications=requests,
    )
    if requests:
        return result, None
    return result, {
        "draft_id": _text(payload.get("draft_id")),
        "company_id": company_id,
        "rule_source_profile_id": rule_profile,
        "source": source,
        "order_number_policy": order_number_policy,
        "orders": materialized_orders,
    }


def prepare_text_order_draft(
    draft_path: str | Path,
    output_dir: str | Path,
    *,
    company_registry_path: str | Path,
    order_number_state_dir: str | Path | None = None,
) -> TextOrderIntakeResult:
    result, materialized = validate_text_order_draft(
        draft_path,
        company_registry_path=company_registry_path,
        require_confirmation=True,
    )
    if materialized is None:
        return result

    order_number_policy = materialized.get("order_number_policy") or {}
    if order_number_policy.get("strategy") == "datetime_sequence" and any(
        order.get("_source_order_no_generated")
        for order in materialized["orders"]
    ):
        allocation = reserve_minute_sequence(
            profile_id=materialized["rule_source_profile_id"],
            prefix=order_number_policy["prefix"],
            batch_minute=order_number_policy["batch_datetime"].strftime(
                order_number_policy["stamp_format"]
            ),
            source_sha256=_text(materialized["source"].get("sha256")),
            sequence_count=len(materialized["orders"]),
            sequence_width=order_number_policy["sequence_width"],
            stamp_format=order_number_policy["stamp_format"],
            sequence_start=order_number_policy["sequence_start"],
            state_dir=order_number_state_dir,
        )
        allocated_policy = {
            **order_number_policy,
            "batch_datetime": datetime.strptime(
                allocation.batch_minute,
                allocation.stamp_format,
            ),
            "prefix": allocation.prefix,
            "stamp_format": allocation.stamp_format,
            "sequence_start": allocation.sequence_start,
            "sequence_width": allocation.sequence_width,
        }
        for order in materialized["orders"]:
            if order.get("_source_order_no_generated"):
                order["source_order_no"] = _generated_order_number(
                    allocated_policy,
                    int(order["_order_index"]),
                )
        result.order_number_allocation = allocation.to_public_dict()

    output_root = Path(output_dir).expanduser().resolve()
    output_root.mkdir(parents=True, exist_ok=True)
    safe_id = re.sub(r"[^0-9A-Za-z_-]+", "_", materialized["draft_id"]).strip("_")
    safe_id = safe_id or "text-orders"
    workbook_path = output_root / f"{safe_id}_文字订单标准化.xlsx"
    plan_path = output_root / f"{safe_id}_文字订单自适应计划.json"

    extension_keys = sorted(
        {
            key
            for order in materialized["orders"]
            for key in order["source_extensions"]
        }
    )
    headers = [
        "来源单号",
        "平台标签",
        "收件人",
        "联系方式",
        "省",
        "市",
        "区",
        "完整地址",
        "商品编码",
        "商品名称",
        "商品规格",
        "数量",
        "单位",
        "指定物流",
        "客户备注",
        "下单时间",
        *[f"扩展:{key}" for key in extension_keys],
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TextOrders"
    sheet.append(headers)
    for order in materialized["orders"]:
        recipient = order["recipient"]
        for item in order["items"]:
            sheet.append(
                [
                    order["source_order_no"],
                    order["source_channel"],
                    recipient["name"],
                    recipient["contact"],
                    recipient["province"],
                    recipient["city"],
                    recipient["district"],
                    recipient["address"],
                    item["source_product_code"],
                    item["source_product_name"],
                    item["source_spec"],
                    item["quantity"],
                    item["unit"],
                    order["source_carrier"],
                    order["source_note"],
                    order["ordered_at"],
                    *[
                        order["source_extensions"].get(key, "")
                        for key in extension_keys
                    ],
                ]
            )
    for row_no in range(2, sheet.max_row + 1):
        sheet.cell(row_no, 1).number_format = "@"
        sheet.cell(row_no, 4).number_format = "@"
        sheet.cell(row_no, 9).number_format = "@"
    workbook.save(workbook_path)

    structure = inspect_excel_structure(workbook_path)
    sheet_structure = structure["sheets"][0]
    columns = {
        str(column["header"]): str(column["column"])
        for column in sheet_structure["columns"]
    }
    field_headers = {
        "source_order_no": "来源单号",
        "recipient_name": "收件人",
        "contact": "联系方式",
        "province": "省",
        "city": "市",
        "district": "区",
        "address": "完整地址",
        "product_code": "商品编码",
        "product_name": "商品名称",
        "source_spec": "商品规格",
        "quantity": "数量",
        "unit": "单位",
        "carrier": "指定物流",
        "source_note": "客户备注",
        "ordered_at": "下单时间",
    }
    field_mapping = {
        field: {
            "columns": [columns[header]],
            "expected_headers": [header],
            "separator": "",
        }
        for field, header in field_headers.items()
    }
    extension_mapping = {
        "source_channel": {
            "columns": [columns["平台标签"]],
            "expected_headers": ["平台标签"],
            "separator": "",
        },
        **{
            key: {
                "columns": [columns[f"扩展:{key}"]],
                "expected_headers": [f"扩展:{key}"],
                "separator": "",
            }
            for key in extension_keys
        },
    }
    plan = {
        "schema_version": "1.0",
        "plan_type": "adaptive_excel_source",
        "plan_id": f"text-{hashlib.sha256((structure['source_sha256'] + materialized['draft_id']).encode('utf-8')).hexdigest()[:20]}",
        "status": "confirmed",
        "source_file": workbook_path.name,
        "source_sha256": structure["source_sha256"],
        "company_id": materialized["company_id"],
        "rule_source_profile_id": materialized["rule_source_profile_id"],
        "sheet_name": "TextOrders",
        "header_row": 1,
        "structure_signature": sheet_structure["structure_signature"],
        "field_mapping": field_mapping,
        "extension_mapping": extension_mapping,
        "batch_defaults": {},
        "order_number": {"strategy": "source_column"},
        "unresolved_fields": [],
    }
    plan_path.write_text(
        json.dumps(plan, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    result.outputs = [str(workbook_path), str(plan_path)]
    return result
