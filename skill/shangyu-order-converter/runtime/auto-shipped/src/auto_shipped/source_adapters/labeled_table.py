from __future__ import annotations

import hashlib
import re
import uuid
import warnings
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from auto_shipped.domain import (
    ClarificationRequest,
    ParsedItem,
    ParsedOrder,
    ParsedRecipient,
    ParsedSource,
    deduplicate_clarifications,
)

from .tiantian_parsed import ParseResult
from .value_validation import normalize_contact


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _positive_quantity(value: Any) -> float | None:
    try:
        number = float(_text(value))
    except (TypeError, ValueError):
        return None
    return number if number > 0 else None


def _cell_text(sheet, reference: str) -> str:
    if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", reference):
        return ""
    return _text(sheet[reference].value)


def _combined_recipient_contact(value: str) -> tuple[str, str]:
    match = re.search(r"1[3-9]\d{9}(?:-\d{4})?", value)
    if match is None:
        return value.strip(), ""
    name = (value[: match.start()] + value[match.end() :]).strip()
    name = re.sub(r"^[\s:：,，;；]+|[\s:：,，;；]+$", "", name)
    return name, match.group(0)


class LabeledTableOrderParsedAdapter:
    """Parse one shipment whose items are a table and recipient facts are labels.

    The parser is deliberately configuration driven.  It is suitable for purchase,
    inbound and requisition styled workbooks when the Skill invocation itself means
    "create a shipment" and the workbook still contains complete recipient facts.
    """

    adapter_id = "labeled_table_order_parsed"

    def parse(
        self,
        source_path: str | Path,
        profile: dict[str, Any],
        layout: dict[str, Any] | None = None,
    ) -> ParseResult:
        path = Path(source_path).expanduser().resolve()
        settings = dict(layout or profile.get("labeled_table") or {})
        sheet_name = str(settings.get("sheet_name") or "Sheet1")
        header_row = int(settings.get("header_row") or 1)
        item_fields = settings.get("item_fields") or {}
        batch_fields = settings.get("batch_fields") or {}
        requests: list[ClarificationRequest] = []

        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Workbook contains no default style.*")
            workbook = load_workbook(path, read_only=False, data_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                return ParseResult(
                    clarifications=[
                        ClarificationRequest(
                            code="MISSING_SOURCE_SHEET",
                            scope="file",
                            source_ref=path.name,
                            question=f"文件中缺少工作表“{sheet_name}”，请确认文件是否完整。",
                            reason="标签式商品表解析器找不到配置工作表。",
                            answer_type="file",
                        )
                    ]
                )
            sheet = workbook[sheet_name]
            headers = {
                _text(sheet.cell(header_row, column).value): column
                for column in range(1, sheet.max_column + 1)
                if _text(sheet.cell(header_row, column).value)
            }

            def column(field: str) -> int | None:
                header = str(item_fields.get(field) or "").strip()
                return headers.get(header) if header else None

            product_column = column("product_name")
            product_code_column = column("product_code")
            quantity_column = column("quantity")
            if (product_column is None and product_code_column is None) or quantity_column is None:
                return ParseResult(
                    clarifications=[
                        ClarificationRequest(
                            code="LABELED_TABLE_ITEM_FIELDS_MISSING",
                            scope="file",
                            source_ref=path.name,
                            question="商品名称/编码或数量列无法识别，请确认表头或更新字段映射。",
                            reason="标签式商品表至少需要一个商品标识列和一个数量列。",
                            answer_type="file",
                        )
                    ]
                )

            def batch_value(field: str) -> str:
                reference = str(batch_fields.get(field) or "").strip().upper()
                return _cell_text(sheet, reference) if reference else ""

            name = batch_value("recipient_name")
            contact = batch_value("contact")
            combined = batch_value("combined_recipient_contact")
            if combined:
                combined_name, combined_contact = _combined_recipient_contact(combined)
                name = name or combined_name
                contact = contact or combined_contact
            address = batch_value("address")
            for field, value, label in (
                ("recipient_name", name, "收件人"),
                ("contact", contact, "联系方式"),
                ("address", address, "收货地址"),
            ):
                if not value:
                    requests.append(
                        ClarificationRequest(
                            code="MISSING_RECIPIENT_FIELD",
                            scope="order",
                            source_ref=path.name,
                            field=label,
                            question=f"文件缺少可识别的{label}，请补充后重新提供。",
                            reason=f"标签式商品表的{field}为空。",
                            answer_type="file",
                        )
                    )
            normalized_contact = normalize_contact(contact) if contact else None
            if contact and normalized_contact is None:
                requests.append(
                    ClarificationRequest(
                        code="INVALID_RECIPIENT_CONTACT",
                        scope="order",
                        source_ref=path.name,
                        field="联系方式",
                        question="收件联系方式格式无法识别，请修正后重新提供。",
                        reason="无效联系方式不能进入管易上传Excel。",
                        answer_type="file",
                    )
                )

            stop_labels = {
                str(value).strip()
                for value in settings.get("stop_labels") or ("合计", "总计")
                if str(value).strip()
            }
            spec_headers = [
                str(value).strip()
                for value in item_fields.get("source_spec_headers") or ()
                if str(value).strip()
            ]
            items: list[ParsedItem] = []
            row_numbers: list[int] = []
            for row_no in range(header_row + 1, sheet.max_row + 1):
                row_labels = {
                    _text(sheet.cell(row_no, column_no).value)
                    for column_no in range(1, sheet.max_column + 1)
                }
                if row_labels.intersection({"收货地址", "收件地址", "联系人", "收货联系人"}):
                    break
                product_name = (
                    _text(sheet.cell(row_no, product_column).value)
                    if product_column is not None
                    else ""
                )
                product_code = (
                    _text(sheet.cell(row_no, product_code_column).value)
                    if product_code_column is not None
                    else ""
                )
                if product_name in stop_labels or product_code in stop_labels:
                    break
                if not product_name and not product_code:
                    continue
                quantity = _positive_quantity(sheet.cell(row_no, quantity_column).value)
                if quantity is None:
                    requests.append(
                        ClarificationRequest(
                            code="INVALID_SOURCE_QUANTITY",
                            scope="field",
                            source_ref=f"{sheet_name}!{row_no}:{row_no}",
                            field=str(item_fields.get("quantity") or "数量"),
                            question="商品数量缺失或不是大于0的数字，请修正后重新提供。",
                            reason="管易每个商品行必须有有效数量。",
                            answer_type="file",
                        )
                    )
                    continue
                source_spec = " ".join(
                    _text(sheet.cell(row_no, headers[header]).value)
                    for header in spec_headers
                    if header in headers and _text(sheet.cell(row_no, headers[header]).value)
                )
                row_numbers.append(row_no)
                items.append(
                    ParsedItem(
                        source_product_code=product_code,
                        source_product_name=product_name,
                        source_spec=source_spec,
                        source_display_spec=source_spec,
                        quantity=quantity,
                        source_line_no=row_no,
                    )
                )
        finally:
            workbook.close()

        requests = deduplicate_clarifications(requests)
        if requests:
            return ParseResult(clarifications=requests)
        if not items or normalized_contact is None:
            return ParseResult(
                clarifications=[
                    ClarificationRequest(
                        code="LABELED_TABLE_NO_ITEMS",
                        scope="file",
                        source_ref=path.name,
                        question="没有找到可转换的商品行，请确认商品表区域。",
                        reason="标签式商品表未产生有效商品明细。",
                        answer_type="file",
                    )
                ]
            )

        mobile, phone = normalized_contact
        fingerprint = _sha256(path)
        record_key = str(
            uuid.uuid5(
                uuid.NAMESPACE_URL,
                f"{profile['profile_id']}:{fingerprint}:single-order",
            )
        )
        return ParseResult(
            orders=[
                ParsedOrder(
                    record_key=record_key,
                    source=ParsedSource(
                        profile_id=str(profile["profile_id"]),
                        source_label=str(profile.get("source_label") or profile["profile_id"]),
                        order_type=str(profile.get("order_type") or "standard_order"),
                        order_no=None,
                        source_order_no=None,
                        file_name=path.name,
                        file_fingerprint=f"sha256:{fingerprint}",
                        sheet_name=sheet_name,
                        row_numbers=tuple(row_numbers),
                    ),
                    recipient=ParsedRecipient(
                        name=name,
                        mobile=mobile,
                        phone=phone,
                        raw_address=address,
                    ),
                    items=items,
                    source_extensions={
                        "source_layout": "single_order_labeled_item_table",
                        "selected_order_identity_field": "generated_central_daily_sequence",
                    },
                )
            ]
        )
