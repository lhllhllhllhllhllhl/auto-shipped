from __future__ import annotations

import hashlib
import json
import re
import unicodedata
import uuid
import warnings
from dataclasses import replace
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from auto_shipped.domain import (
    ClarificationRequest,
    ParsedItem,
    ParsedOrder,
    ParsedRecipient,
    ParsedShipmentFacts,
    ParsedSource,
    deduplicate_clarifications,
)

from .tiantian_parsed import ParseResult
from .value_validation import identifier_requires_text, normalize_contact


ADAPTIVE_PLAN_SCHEMA_VERSION = "1.0"
ADAPTIVE_PLAN_TYPE = "adaptive_excel_source"
ADAPTIVE_ASSUMPTION_CONFIRMATION = "CONFIRM_ADAPTIVE_BATCH_ASSUMPTIONS"


class AdaptiveExcelPlanError(ValueError):
    """Raised when an adaptive plan cannot safely describe the source workbook."""


FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "source_order_no": ("订单编号", "订单号", "来源单号", "平台单号"),
    "recipient_name": ("收件人姓名", "收件人", "收货人", "姓名"),
    "contact": ("收件人手机", "收件人电话", "联系方式", "联系电话", "手机", "电话"),
    "province": ("省", "省份"),
    "city": ("市", "城市"),
    "district": ("区", "区县", "县区"),
    "detail": ("详细地址", "收货详细地址"),
    "address": ("完整地址", "收货地址", "省市区详细地址"),
    "product_code": ("商品编码", "商品代码", "来源商品编码"),
    "product_name": ("商品名称", "礼品名称", "套餐名称", "商品", "品名"),
    "source_spec": ("规格编码", "规格代码", "商品规格", "礼品规格", "规格"),
    "quantity": ("数量", "商品数量", "件数"),
    "unit": ("单位", "计量单位"),
    "carrier": ("快递公司", "物流公司", "指定物流"),
    "tracking_no": ("快递单号", "物流单号"),
    "source_note": ("备注", "客户备注", "买家备注", "卖家备注", "送货时间"),
    "ordered_at": ("订单生成时间", "订单创建时间", "下单时间"),
}

ADDRESS_COMPONENT_ALIASES = {
    "省",
    "省份",
    "市",
    "城市",
    "区",
    "区县",
    "县区",
    "镇街道",
    "街道",
    "详细地址",
    "收货详细地址",
}


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalized(value: Any) -> str:
    text = unicodedata.normalize("NFKC", _text(value)).casefold()
    return "".join(
        character
        for character in text
        if not character.isspace()
        and not unicodedata.category(character).startswith("P")
    )


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _value_kind(value: Any) -> str:
    if value in (None, ""):
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number"
    if hasattr(value, "isoformat") and not isinstance(value, str):
        return "date"
    return "text"


def _header_score(values: tuple[Any, ...]) -> tuple[int, int]:
    texts = [_text(value) for value in values]
    nonempty = [value for value in texts if value]
    alias_tokens = {
        _normalized(alias)
        for aliases in FIELD_ALIASES.values()
        for alias in aliases
    } | {_normalized(value) for value in ADDRESS_COMPONENT_ALIASES}
    alias_hits = sum(_normalized(value) in alias_tokens for value in nonempty)
    return alias_hits, len(nonempty)


def inspect_excel_structure(source_path: str | Path) -> dict[str, Any]:
    """Return workbook structure and type counts without returning any data value."""

    path = Path(source_path).expanduser().resolve()
    if not path.is_file():
        raise AdaptiveExcelPlanError(f"找不到Excel文件：{path.name}")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise AdaptiveExcelPlanError("自适应结构探针只支持.xlsx或.xlsm文件。")
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style.*")
        workbook = load_workbook(path, read_only=False, data_only=True)
    try:
        sheets: list[dict[str, Any]] = []
        for sheet in workbook.worksheets:
            scan_end = min(max(sheet.max_row, 1), 20)
            scored_rows: list[tuple[tuple[int, int], int]] = []
            for row_no in range(1, scan_end + 1):
                values = tuple(
                    sheet.cell(row_no, column_no).value
                    for column_no in range(1, sheet.max_column + 1)
                )
                scored_rows.append((_header_score(values), row_no))
            _, header_row = max(scored_rows, key=lambda item: (item[0], -item[1]))
            headers = [
                _text(sheet.cell(header_row, column_no).value)
                for column_no in range(1, sheet.max_column + 1)
            ]
            columns: list[dict[str, Any]] = []
            for column_no, header in enumerate(headers, start=1):
                counts = {
                    "blank": 0,
                    "text": 0,
                    "number": 0,
                    "date": 0,
                    "boolean": 0,
                }
                for row_no in range(header_row + 1, sheet.max_row + 1):
                    counts[_value_kind(sheet.cell(row_no, column_no).value)] += 1
                columns.append(
                    {
                        "column": get_column_letter(column_no),
                        "header": header,
                        "nonempty_data_rows": sum(
                            count for kind, count in counts.items() if kind != "blank"
                        ),
                        "value_type_counts": counts,
                    }
                )
            signature_source = {
                "sheet_name": sheet.title,
                "header_row": header_row,
                "headers": headers,
            }
            signature = hashlib.sha256(
                json.dumps(
                    signature_source,
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                ).encode("utf-8")
            ).hexdigest()
            sheets.append(
                {
                    "sheet_name": sheet.title,
                    "row_count": sheet.max_row,
                    "column_count": sheet.max_column,
                    "candidate_header_row": header_row,
                    "structure_signature": signature,
                    "columns": columns,
                }
            )
        return {
            "schema_version": "1.0",
            "source_file": path.name,
            "source_sha256": _sha256(path),
            "sheets": sheets,
            "privacy": "headers_and_type_counts_only_no_data_values",
        }
    finally:
        workbook.close()


def _unique_alias_column(
    columns: list[dict[str, Any]],
    aliases: tuple[str, ...],
) -> dict[str, Any] | None:
    normalized_aliases = {_normalized(value) for value in aliases}
    matches = [
        column
        for column in columns
        if column.get("header")
        and _normalized(column.get("header")) in normalized_aliases
    ]
    return matches[0] if len(matches) == 1 else None


def propose_field_mapping(sheet_structure: dict[str, Any]) -> dict[str, Any]:
    columns = list(sheet_structure.get("columns") or [])
    mapping: dict[str, Any] = {}
    for field, aliases in FIELD_ALIASES.items():
        column = _unique_alias_column(columns, aliases)
        if column is None:
            continue
        mapping[field] = {
            "columns": [column["column"]],
            "expected_headers": [column["header"]],
            "separator": "",
        }

    component_columns = [
        column
        for column in columns
        if _normalized(column.get("header"))
        in {_normalized(value) for value in ADDRESS_COMPONENT_ALIASES}
    ]
    if component_columns and (
        "address" not in mapping or len(component_columns) > 1
    ):
        mapping["address"] = {
            "columns": [column["column"] for column in component_columns],
            "expected_headers": [column["header"] for column in component_columns],
            "separator": "",
        }
    return mapping


def load_adaptive_plan(path: str | Path) -> dict[str, Any]:
    plan_path = Path(path).expanduser().resolve()
    try:
        plan = json.loads(plan_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise AdaptiveExcelPlanError(f"无法读取自适应来源计划：{exc}") from exc
    if not isinstance(plan, dict):
        raise AdaptiveExcelPlanError("自适应来源计划根节点必须是对象。")
    if plan.get("schema_version") != ADAPTIVE_PLAN_SCHEMA_VERSION:
        raise AdaptiveExcelPlanError("自适应来源计划schema_version必须是1.0。")
    if plan.get("plan_type") != ADAPTIVE_PLAN_TYPE:
        raise AdaptiveExcelPlanError("自适应来源计划plan_type不正确。")
    return plan


def _mapping_value(
    sheet,
    row_no: int,
    spec: dict[str, Any] | None,
) -> str:
    if not spec:
        return ""
    columns = list(spec.get("columns") or [])
    separator = str(spec.get("separator") or "")
    values = [
        _text(sheet[f"{column}{row_no}"].value)
        for column in columns
        if re.fullmatch(r"[A-Z]{1,3}", str(column or ""))
    ]
    return separator.join(value for value in values if value)


def _positive_quantity(value: Any) -> float | None:
    try:
        number = float(_text(value))
    except (TypeError, ValueError):
        return None
    return number if number > 0 else None


def _plan_default(plan: dict[str, Any], field: str) -> Any:
    return (plan.get("batch_defaults") or {}).get(field)


def _generated_order_number(
    plan: dict[str, Any],
    sequence_no: int,
) -> str | None:
    settings = plan.get("order_number") or {}
    if settings.get("strategy") != "date_sequence":
        return None
    try:
        business_date = date.fromisoformat(str(settings.get("business_date") or ""))
    except ValueError:
        return None
    prefix = str(settings.get("prefix") or "")
    if not re.fullmatch(r"[0-9A-Za-z_-]{0,16}", prefix):
        return None
    width = int(settings.get("sequence_width") or 3)
    start = int(settings.get("sequence_start") or 1)
    if width < 1 or width > 8 or start < 0:
        return None
    return f"{prefix}{business_date.strftime('%y%m%d')}{start + sequence_no - 1:0{width}d}"


def _merge_order_rows(
    orders: list[ParsedOrder],
    requests: list[ClarificationRequest],
) -> list[ParsedOrder]:
    """Merge repeated source order numbers into one multi-item canonical order."""

    grouped: dict[str, list[ParsedOrder]] = {}
    order_keys: list[str] = []
    for order in orders:
        source_order_no = str(order.source.source_order_no or "")
        if source_order_no not in grouped:
            grouped[source_order_no] = []
            order_keys.append(source_order_no)
        grouped[source_order_no].append(order)

    merged: list[ParsedOrder] = []
    for source_order_no in order_keys:
        group = grouped[source_order_no]
        first = group[0]
        if len(group) == 1:
            merged.append(first)
            continue
        conflicts: list[str] = []
        for candidate in group[1:]:
            for field, first_value, candidate_value in (
                ("recipient", first.recipient, candidate.recipient),
                ("ordered_at", first.ordered_at, candidate.ordered_at),
                ("source_note", first.source_note, candidate.source_note),
                ("shipment_facts", first.shipment_facts, candidate.shipment_facts),
                ("source_extensions", first.source_extensions, candidate.source_extensions),
            ):
                if first_value != candidate_value:
                    conflicts.append(field)
        row_numbers = tuple(
            row_no
            for candidate in group
            for row_no in candidate.source.row_numbers
        )
        if conflicts:
            requests.append(
                ClarificationRequest(
                    code="ADAPTIVE_ORDER_FIELDS_CONFLICT",
                    scope="order",
                    source_ref=(
                        f"{first.source.sheet_name}!"
                        f"{min(row_numbers)}:{max(row_numbers)}"
                    ),
                    question="同一来源单号的多条商品行存在订单字段冲突，请修正后重试。",
                    reason="冲突字段：" + ", ".join(sorted(set(conflicts))),
                    answer_type="file",
                )
            )
            continue
        first.source = replace(first.source, row_numbers=row_numbers)
        first.items = [item for candidate in group for item in candidate.items]
        merged.append(first)
    return merged


class AdaptiveExcelParsedAdapter:
    """Parse a one-off Excel layout through a validated, non-PII mapping plan."""

    adapter_id = "adaptive_excel_parsed"

    def parse(
        self,
        source_path: str | Path,
        profile: dict[str, Any],
        plan: dict[str, Any],
    ) -> ParseResult:
        path = Path(source_path).expanduser().resolve()
        requests: list[ClarificationRequest] = []
        expected_sha = str(plan.get("source_sha256") or "")
        actual_sha = _sha256(path)
        if expected_sha != actual_sha:
            return ParseResult(
                clarifications=[
                    ClarificationRequest(
                        code="ADAPTIVE_PLAN_SOURCE_CHANGED",
                        scope="file",
                        source_ref=path.name,
                        question="自适应字段映射对应的Excel已经变化，请重新分析当前文件。",
                        reason="来源文件SHA-256与映射计划不一致。",
                        answer_type="file",
                        next_action="重新运行自适应来源结构分析。",
                    )
                ]
            )
        if plan.get("company_id") != profile.get("company_id"):
            raise AdaptiveExcelPlanError("自适应计划公司与业务规则来源公司不一致。")
        if plan.get("rule_source_profile_id") != profile.get("profile_id"):
            raise AdaptiveExcelPlanError("自适应计划绑定的规则来源配置不一致。")

        mapping = plan.get("field_mapping") or {}
        if not isinstance(mapping, dict):
            raise AdaptiveExcelPlanError("field_mapping必须是对象。")
        extension_mapping = plan.get("extension_mapping") or {}
        if not isinstance(extension_mapping, dict):
            raise AdaptiveExcelPlanError("extension_mapping必须是对象。")
        uses_assumptions = bool(plan.get("batch_defaults")) or (
            (plan.get("order_number") or {}).get("strategy") == "date_sequence"
        )
        assumptions_confirmed = (
            plan.get("assumption_confirmation_token")
            == ADAPTIVE_ASSUMPTION_CONFIRMATION
        )
        if uses_assumptions and not assumptions_confirmed:
            requests.append(
                ClarificationRequest(
                    code="ADAPTIVE_BATCH_ASSUMPTIONS_UNCONFIRMED",
                    scope="batch",
                    question="请确认这批订单补充的商品、数量和单号生成规则后再继续。",
                    reason="自适应计划包含原Excel没有的批次默认值或生成规则。",
                    answer_type="confirmation",
                    next_action=(
                        "用户明确确认后，把assumption_confirmation_token设为"
                        f"{ADAPTIVE_ASSUMPTION_CONFIRMATION}。"
                    ),
                )
            )

        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Workbook contains no default style.*")
            workbook = load_workbook(path, read_only=False, data_only=True)
        try:
            sheet_name = str(plan.get("sheet_name") or "")
            if sheet_name not in workbook.sheetnames:
                raise AdaptiveExcelPlanError("自适应计划指定的工作表不存在。")
            sheet = workbook[sheet_name]
            header_row = int(plan.get("header_row") or 1)
            for field, spec in {**mapping, **extension_mapping}.items():
                columns = list((spec or {}).get("columns") or [])
                expected_headers = list((spec or {}).get("expected_headers") or [])
                if len(columns) != len(expected_headers) or not columns:
                    raise AdaptiveExcelPlanError(f"字段{field}的列映射结构不完整。")
                for column, expected_header in zip(columns, expected_headers, strict=True):
                    actual_header = _text(sheet[f"{column}{header_row}"].value)
                    if actual_header != str(expected_header):
                        raise AdaptiveExcelPlanError(
                            f"字段{field}的列{column}表头已变化。"
                        )

            row_numbers: list[int] = []
            for row_no in range(header_row + 1, sheet.max_row + 1):
                if any(
                    _mapping_value(sheet, row_no, mapping.get(field))
                    for field in ("recipient_name", "contact", "address")
                ):
                    row_numbers.append(row_no)
            if not row_numbers:
                return ParseResult(
                    clarifications=[
                        ClarificationRequest(
                            code="ADAPTIVE_NO_ORDER_ROWS",
                            scope="file",
                            source_ref=path.name,
                            question="没有找到同时具备收件信息的数据行，请确认表头和数据区域。",
                            reason="自适应字段映射没有识别出订单行。",
                            answer_type="file",
                        )
                    ]
                )

            missing_batch_fields: set[str] = set()
            orders: list[ParsedOrder] = []
            for sequence_no, row_no in enumerate(row_numbers, start=1):
                row_ref = f"{sheet_name}!{row_no}:{row_no}"
                source_order_value = _mapping_value(
                    sheet, row_no, mapping.get("source_order_no")
                )
                source_order_raw = (
                    sheet[f"{mapping['source_order_no']['columns'][0]}{row_no}"].value
                    if mapping.get("source_order_no")
                    else None
                )
                source_order_no = source_order_value or _generated_order_number(
                    plan, sequence_no
                )
                if not source_order_no:
                    missing_batch_fields.add("source_order_no")
                elif source_order_value and identifier_requires_text(source_order_raw):
                    requests.append(
                        ClarificationRequest(
                            code="SOURCE_IDENTIFIER_PRECISION_UNSAFE",
                            scope="field",
                            source_ref=row_ref,
                            field="平台单号",
                            question="来源订单号可能已经被Excel转为数值或科学计数法，请改为文本后重新导出。",
                            reason="无法保证订单号原始数字完整。",
                            answer_type="file",
                        )
                    )

                product_code = _mapping_value(
                    sheet, row_no, mapping.get("product_code")
                ) or _text(_plan_default(plan, "product_code"))
                product_name = _mapping_value(
                    sheet, row_no, mapping.get("product_name")
                ) or _text(_plan_default(plan, "product_name"))
                if not product_code and not product_name:
                    missing_batch_fields.add("product_identity")

                quantity_value = _mapping_value(
                    sheet, row_no, mapping.get("quantity")
                )
                if not quantity_value:
                    quantity_value = _plan_default(plan, "quantity")
                quantity = _positive_quantity(quantity_value)
                if quantity is None:
                    missing_batch_fields.add("quantity")

                name = _mapping_value(sheet, row_no, mapping.get("recipient_name"))
                contact = _mapping_value(sheet, row_no, mapping.get("contact"))
                raw_address = _mapping_value(sheet, row_no, mapping.get("address"))
                for field, value, label in (
                    ("recipient_name", name, "收件人"),
                    ("contact", contact, "联系方式"),
                    ("address", raw_address, "收货地址"),
                ):
                    if not value:
                        requests.append(
                            ClarificationRequest(
                                code="MISSING_RECIPIENT_FIELD",
                                scope="order",
                                source_ref=row_ref,
                                field=label,
                                question=f"该行缺少{label}，请补充或修正字段映射。",
                                reason="管易订单需要完整收件信息。",
                                answer_type="file",
                            )
                        )
                normalized_contact = normalize_contact(contact) if contact else None
                if contact and normalized_contact is None:
                    requests.append(
                        ClarificationRequest(
                            code="INVALID_RECIPIENT_CONTACT",
                            scope="order",
                            source_ref=row_ref,
                            field="联系方式",
                            question="收件联系方式格式无法识别，请修正为11位大陆手机号或常用座机号码。",
                            reason="无效联系方式不能进入管易上传Excel。",
                            answer_type="file",
                        )
                    )
                if (
                    not source_order_no
                    or not (product_code or product_name)
                    or quantity is None
                    or not name
                    or normalized_contact is None
                    or not raw_address
                ):
                    continue
                mobile, phone = normalized_contact
                carrier = _mapping_value(sheet, row_no, mapping.get("carrier"))
                tracking_no = _mapping_value(
                    sheet, row_no, mapping.get("tracking_no")
                )
                record_key = str(
                    uuid.uuid5(
                        uuid.NAMESPACE_URL,
                        f"adaptive:{profile['profile_id']}:{actual_sha}:{source_order_no}",
                    )
                )
                orders.append(
                    ParsedOrder(
                        record_key=record_key,
                        source=ParsedSource(
                            profile_id=str(profile["profile_id"]),
                            source_label=str(
                                profile.get("source_label") or profile["profile_id"]
                            ),
                            order_type=str(profile.get("order_type") or "standard_order"),
                            source_order_no=source_order_no,
                            file_name=path.name,
                            file_fingerprint=f"sha256:{actual_sha}",
                            sheet_name=sheet_name,
                            row_numbers=(row_no,),
                        ),
                        recipient=ParsedRecipient(
                            name=name,
                            mobile=mobile,
                            phone=phone,
                            province=_mapping_value(
                                sheet, row_no, mapping.get("province")
                            ) or _text(_plan_default(plan, "province")),
                            city=_mapping_value(sheet, row_no, mapping.get("city"))
                            or _text(_plan_default(plan, "city")),
                            district=_mapping_value(
                                sheet, row_no, mapping.get("district")
                            ) or _text(_plan_default(plan, "district")),
                            detail=_mapping_value(sheet, row_no, mapping.get("detail")),
                            raw_address=raw_address,
                        ),
                        items=[
                            ParsedItem(
                                source_product_code=product_code,
                                source_product_name=product_name,
                                source_spec=_mapping_value(
                                    sheet, row_no, mapping.get("source_spec")
                                ) or _text(_plan_default(plan, "source_spec")),
                                source_display_spec=_mapping_value(
                                    sheet, row_no, mapping.get("source_spec")
                                ),
                                quantity=quantity,
                                unit=_mapping_value(
                                    sheet, row_no, mapping.get("unit")
                                ),
                                source_line_no=row_no,
                            )
                        ],
                        ordered_at=_mapping_value(
                            sheet, row_no, mapping.get("ordered_at")
                        ) or None,
                        source_note=_mapping_value(
                            sheet, row_no, mapping.get("source_note")
                        ) or None,
                        shipment_facts=(
                            ParsedShipmentFacts(
                                carrier_name=carrier,
                                tracking_no=tracking_no,
                            )
                            if carrier or tracking_no
                            else None
                        ),
                        source_extensions={
                            "adaptive_plan_id": str(plan.get("plan_id") or ""),
                            "adaptive_structure_signature": str(
                                plan.get("structure_signature") or ""
                            ),
                            **{
                                str(field): _mapping_value(sheet, row_no, spec)
                                for field, spec in extension_mapping.items()
                            },
                        },
                    )
                )
        finally:
            workbook.close()

        if "source_order_no" in missing_batch_fields:
            requests.append(
                ClarificationRequest(
                    code="ADAPTIVE_ORDER_NUMBER_REQUIRED",
                    scope="batch",
                    field="平台单号",
                    question="原表没有可用订单号。请确认是否按日期加顺序号为这批订单生成唯一来源单号。",
                    reason="管易平台单号必须能唯一标识每张订单。",
                    answer_type="confirmation",
                )
            )
        if "product_identity" in missing_batch_fields:
            requests.append(
                ClarificationRequest(
                    code="ADAPTIVE_PRODUCT_REQUIRED",
                    scope="batch",
                    field="商品",
                    question="原表没有商品信息。请说明这批订单对应的商品或套餐；如果全部相同，只需确认一次。",
                    reason="无法执行商品映射或套餐展开。",
                    answer_type="text",
                )
            )
        if "quantity" in missing_batch_fields:
            requests.append(
                ClarificationRequest(
                    code="ADAPTIVE_QUANTITY_REQUIRED",
                    scope="batch",
                    field="数量",
                    question="原表没有可确认的商品数量。请说明每单数量；如果全部相同，只需确认一次。",
                    reason="管易每个商品行必须有大于0的数量。",
                    answer_type="text",
                )
            )
        requests = deduplicate_clarifications(requests)
        if requests:
            return ParseResult(orders=[], clarifications=requests)
        orders = _merge_order_rows(orders, requests)
        requests = deduplicate_clarifications(requests)
        if requests:
            return ParseResult(orders=[], clarifications=requests)
        return ParseResult(orders=orders)
