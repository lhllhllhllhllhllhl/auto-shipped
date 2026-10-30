from __future__ import annotations

import unicodedata
import warnings
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from auto_shipped.domain import ClarificationRequest, deduplicate_clarifications

from .configured_excel import ConfiguredExcelParsedAdapter
from .tiantian_parsed import ParseResult


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalized(value: Any) -> str:
    return unicodedata.normalize("NFKC", _text(value)).replace(" ", "")


class BailihuiDropshipParsedAdapter:
    """Parse 百礼汇待发货代发采购 exports through the shared row adapter."""

    adapter_id = "bailihui_dropship_parsed"

    def parse(self, source_path: str | Path, profile: dict[str, Any]) -> ParseResult:
        path = Path(source_path)
        detection = profile["detection"]
        sheet_name = str(detection["sheet_name"])
        header_row = int(detection.get("header_row", 1))
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore", message="Workbook contains no default style.*"
            )
            workbook = load_workbook(path, read_only=False, data_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                return ParseResult(
                    clarifications=[
                        ClarificationRequest(
                            code="MISSING_SOURCE_SHEET",
                            scope="file",
                            source_ref=path.name,
                            question=(
                                f"文件中缺少工作表“{sheet_name}”，请从百礼汇供应商后台"
                                "重新导出代发订单。"
                            ),
                            reason="百礼汇来源适配器无法找到固定订单工作表。",
                            answer_type="file",
                        )
                    ]
                )
            sheet = workbook[sheet_name]
            headers = [_text(cell.value) for cell in sheet[header_row]]
            missing = [
                name for name in detection["required_headers"] if name not in headers
            ]
            if missing:
                return ParseResult(
                    clarifications=[
                        ClarificationRequest(
                            code="MISSING_SOURCE_HEADERS",
                            scope="file",
                            source_ref=path.name,
                            question=(
                                "百礼汇Excel缺少必要列，请重新导出正确版本或确认后台格式是否升级。"
                            ),
                            reason=f"缺少列：{', '.join(missing)}",
                            answer_type="file",
                        )
                    ]
                )

            rows: list[tuple[int, dict[str, Any]]] = []
            for row_no, cells in enumerate(
                sheet.iter_rows(min_row=header_row + 1, values_only=True),
                start=header_row + 1,
            ):
                row = {
                    headers[index]: value
                    for index, value in enumerate(cells)
                    if index < len(headers) and headers[index]
                }
                if any(value not in (None, "") for value in row.values()):
                    rows.append((row_no, row))
        finally:
            workbook.close()

        clarifications = self._validate_export_rows(rows, profile, sheet_name)
        if clarifications:
            return ParseResult(
                clarifications=deduplicate_clarifications(clarifications)
            )
        return ConfiguredExcelParsedAdapter().parse_rows(rows, path, profile)

    @staticmethod
    def _validate_export_rows(
        rows: list[tuple[int, dict[str, Any]]],
        profile: dict[str, Any],
        sheet_name: str,
    ) -> list[ClarificationRequest]:
        requests: list[ClarificationRequest] = []
        selection = profile.get("order_row_policy") or {}
        purchase_type_field = str(selection.get("purchase_type_field") or "")
        allowed_purchase_types = {
            _normalized(value)
            for value in selection.get("allowed_purchase_types") or []
        }
        status_field = str(selection.get("status_field") or "")
        allowed_statuses = {
            _normalized(value) for value in selection.get("allowed_statuses") or []
        }
        if not rows:
            return [
                ClarificationRequest(
                    code="BAILIHUI_NO_ORDER_ROWS",
                    scope="file",
                    question="百礼汇文件中没有订单行，请确认筛选结果后重新导出。",
                    reason="空文件不能生成管易导入订单。",
                    answer_type="file",
                )
            ]

        for row_no, row in rows:
            purchase_type = _normalized(row.get(purchase_type_field))
            if allowed_purchase_types and purchase_type not in allowed_purchase_types:
                requests.append(
                    ClarificationRequest(
                        code="BAILIHUI_PURCHASE_TYPE_NOT_DROPSHIP",
                        scope="field",
                        source_ref=f"{sheet_name}!{row_no}:{row_no}",
                        field=purchase_type_field,
                        question="文件中包含非代发采购，请重新筛选百礼汇代发订单后导出。",
                        reason=(
                            f"当前值：{_text(row.get(purchase_type_field)) or '空'}；"
                            "转换模块只处理代发采购。"
                        ),
                        answer_type="file",
                    )
                )
            status = _normalized(row.get(status_field))
            if allowed_statuses and status not in allowed_statuses:
                requests.append(
                    ClarificationRequest(
                        code="BAILIHUI_ORDER_STATUS_NOT_PENDING_SHIPMENT",
                        scope="field",
                        source_ref=f"{sheet_name}!{row_no}:{row_no}",
                        field=status_field,
                        question="文件中包含非待发货订单，请重新筛选待发货订单后导出。",
                        reason=(
                            f"当前值：{_text(row.get(status_field)) or '空'}；"
                            "已处理订单不得重复生成管易导入文件。"
                        ),
                        answer_type="file",
                    )
                )
        return requests
