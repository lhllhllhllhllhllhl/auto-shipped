from __future__ import annotations

import hashlib
import json
import math
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


@dataclass(frozen=True, slots=True)
class PreflightIssue:
    code: str
    message: str
    row: int | None = None
    column: str | None = None
    severity: str = "error"

    def to_dict(self) -> dict[str, Any]:
        return {
            "code": self.code,
            "severity": self.severity,
            "message": self.message,
            "row": self.row,
            "column": self.column,
        }


@dataclass(slots=True)
class GuanyiPreflightResult:
    status: str
    file_path: str
    profile_id: str
    profile_version: int
    file_sha256: str | None = None
    size_bytes: int | None = None
    order_count: int = 0
    item_row_count: int = 0
    checks: dict[str, bool] = field(default_factory=dict)
    issues: list[PreflightIssue] = field(default_factory=list)

    def to_manifest(self) -> dict[str, Any]:
        file = Path(self.file_path)
        return {
            "schema_version": "upload-batch/1.0",
            "status": self.status,
            "platform": "guanyi",
            "operation": "custom_order_import",
            "validated_at": datetime.now(timezone.utc).isoformat(),
            "profile": {
                "profile_id": self.profile_id,
                "version": self.profile_version,
            },
            "artifact": {
                "path": str(file.resolve()),
                "file_name": file.name,
                "sha256": self.file_sha256,
                "size_bytes": self.size_bytes,
            },
            "counts": {
                "orders": self.order_count,
                "item_rows": self.item_row_count,
            },
            "checks": self.checks,
            "issues": [issue.to_dict() for issue in self.issues],
            "safety": {
                "manifest_contains_recipient_pii": False,
                "workbook_contains_recipient_pii": True,
                "upload_authorized": False,
            },
        }

    def write_manifest(self, output_path: str | Path) -> Path:
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(
            json.dumps(self.to_manifest(), ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        return output


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _is_finite_number(value: Any) -> bool:
    return (
        isinstance(value, (int, float))
        and not isinstance(value, bool)
        and math.isfinite(float(value))
    )


def _same_value(left: Any, right: Any) -> bool:
    if _is_blank(left) and _is_blank(right):
        return True
    return left == right


def _default_checks() -> dict[str, bool]:
    return {
        "file_is_xlsx": False,
        "workbook_readable": False,
        "sheet_structure_valid": False,
        "headers_match_exactly": False,
        "data_rows_present": False,
        "formulas_absent": False,
        "cell_errors_absent": False,
        "required_values_present": False,
        "required_any_values_present": False,
        "numeric_values_valid": False,
        "pattern_values_valid": False,
        "allowed_values_valid": False,
        "required_text_suffixes_valid": False,
        "identifier_formats_preserved": False,
        "order_fields_consistent": False,
    }


def preflight_custom_import(
    workbook_path: str | Path,
    platform_profile: dict[str, Any],
) -> GuanyiPreflightResult:
    path = Path(workbook_path)
    result = GuanyiPreflightResult(
        status="blocked",
        file_path=str(path),
        profile_id=str(platform_profile.get("profile_id", "unknown")),
        profile_version=int(platform_profile.get("version", 0)),
        checks=_default_checks(),
    )

    if path.suffix.lower() != ".xlsx":
        result.issues.append(
            PreflightIssue("FILE_TYPE_UNSUPPORTED", "只允许上传 .xlsx 文件。")
        )
        return result
    result.checks["file_is_xlsx"] = True

    if not path.is_file():
        result.issues.append(
            PreflightIssue("WORKBOOK_NOT_FOUND", "待上传文件不存在。")
        )
        return result

    result.size_bytes = path.stat().st_size
    result.file_sha256 = _sha256(path)
    try:
        workbook = load_workbook(
            path,
            read_only=False,
            data_only=False,
            keep_links=False,
        )
    except (OSError, ValueError, KeyError, BadZipFile, InvalidFileException) as exc:
        result.issues.append(
            PreflightIssue("WORKBOOK_UNREADABLE", f"无法读取待上传工作簿：{exc}")
        )
        return result

    result.checks["workbook_readable"] = True
    try:
        try:
            sheet_name = platform_profile["sheet_name"]
            header_row = int(platform_profile["header_row"])
            columns = list(platform_profile["columns"])
        except (KeyError, TypeError, ValueError):
            result.issues.append(
                PreflightIssue(
                    "PLATFORM_PROFILE_INVALID",
                    "管易平台配置缺少工作表、表头行或列定义。",
                )
            )
            return result
        allow_extra_sheets = bool(
            platform_profile.get("preflight", {}).get("allow_extra_sheets", False)
        )
        if sheet_name not in workbook.sheetnames:
            result.issues.append(
                PreflightIssue(
                    "REQUIRED_SHEET_MISSING",
                    "工作簿缺少管易导入工作表。",
                )
            )
            return result
        if not allow_extra_sheets and workbook.sheetnames != [sheet_name]:
            result.issues.append(
                PreflightIssue(
                    "UNEXPECTED_WORKSHEETS",
                    "工作簿包含模板之外的工作表。",
                )
            )
        else:
            result.checks["sheet_structure_valid"] = True

        sheet = workbook[sheet_name]
        headers = [
            sheet.cell(header_row, column_no).value
            for column_no in range(1, len(columns) + 1)
        ]
        extra_headers = [
            sheet.cell(header_row, column_no).value
            for column_no in range(len(columns) + 1, sheet.max_column + 1)
            if not _is_blank(sheet.cell(header_row, column_no).value)
        ]
        if headers != columns or extra_headers:
            result.issues.append(
                PreflightIssue(
                    "TEMPLATE_HEADERS_MISMATCH",
                    "工作簿表头与已登记的管易模板不一致。",
                    row=header_row,
                )
            )
            return result
        result.checks["headers_match_exactly"] = True

        column_numbers = {name: index + 1 for index, name in enumerate(columns)}
        settings = platform_profile.get("preflight", {})
        configured_columns = {
            *settings.get(
                "required_nonempty_columns",
                platform_profile.get("required_columns_from_template_style", []),
            ),
            *settings.get("nonnegative_numeric_columns", []),
            *settings.get("optional_nonnegative_numeric_columns", []),
            *settings.get("positive_numeric_columns", []),
            *settings.get("binary_flag_columns", []),
            *settings.get("order_consistency_columns", []),
            *[
                column
                for group in settings.get("required_any_nonempty_groups", [])
                for column in group
            ],
            *settings.get("pattern_columns", {}).keys(),
            *settings.get("allowed_values", {}).keys(),
            *settings.get("required_text_suffixes", {}).keys(),
            *platform_profile.get("identifier_columns", []),
            platform_profile.get("order_group_key", "平台单号"),
        }
        unknown_columns = sorted(configured_columns - set(columns))
        if unknown_columns:
            result.issues.append(
                PreflightIssue(
                    "PLATFORM_PROFILE_COLUMN_UNKNOWN",
                    "管易平台预检配置引用了模板中不存在的列。",
                )
            )
            return result

        content_rows: list[int] = []
        for row_no in range(header_row + 1, sheet.max_row + 1):
            values = [
                sheet.cell(row_no, column_no).value
                for column_no in range(1, len(columns) + 1)
            ]
            if any(not _is_blank(value) for value in values):
                content_rows.append(row_no)
            else:
                result.issues.append(
                    PreflightIssue(
                        "EMPTY_DATA_ROW",
                        "数据区域包含空白行。",
                        row=row_no,
                    )
                )
        if not content_rows:
            result.issues.append(
                PreflightIssue("NO_DATA_ROWS", "工作簿没有可上传的订单明细行。")
            )
            return result
        result.checks["data_rows_present"] = True
        result.item_row_count = len(content_rows)

        formula_found = False
        cell_error_found = False
        for row_no in content_rows:
            for column_no in range(1, len(columns) + 1):
                cell = sheet.cell(row_no, column_no)
                if cell.data_type == "f":
                    formula_found = True
                    result.issues.append(
                        PreflightIssue(
                            "FORMULA_NOT_ALLOWED",
                            "待上传数据不能包含公式。",
                            row=row_no,
                            column=columns[column_no - 1],
                        )
                    )
                elif cell.data_type == "e":
                    cell_error_found = True
                    result.issues.append(
                        PreflightIssue(
                            "CELL_ERROR_NOT_ALLOWED",
                            "待上传数据不能包含错误单元格。",
                            row=row_no,
                            column=columns[column_no - 1],
                        )
                    )
        result.checks["formulas_absent"] = not formula_found
        result.checks["cell_errors_absent"] = not cell_error_found

        required_columns = settings.get(
            "required_nonempty_columns",
            platform_profile.get("required_columns_from_template_style", []),
        )
        missing_required = False
        for row_no in content_rows:
            for column in required_columns:
                value = sheet.cell(row_no, column_numbers[column]).value
                if _is_blank(value):
                    missing_required = True
                    result.issues.append(
                        PreflightIssue(
                            "REQUIRED_VALUE_MISSING",
                            "必填字段为空。",
                            row=row_no,
                            column=column,
                        )
                    )
        result.checks["required_values_present"] = not missing_required

        missing_required_group = False
        for row_no in content_rows:
            for group in settings.get("required_any_nonempty_groups", []):
                if not group:
                    continue
                values = [
                    sheet.cell(row_no, column_numbers[column]).value
                    for column in group
                ]
                if all(_is_blank(value) for value in values):
                    missing_required_group = True
                    result.issues.append(
                        PreflightIssue(
                            "REQUIRED_VALUE_GROUP_MISSING",
                            "这一组字段至少需要填写一项。",
                            row=row_no,
                            column="/".join(group),
                        )
                    )
        result.checks["required_any_values_present"] = not missing_required_group

        numeric_invalid = False
        nonnegative_columns = settings.get("nonnegative_numeric_columns", [])
        optional_nonnegative_columns = settings.get(
            "optional_nonnegative_numeric_columns", []
        )
        positive_columns = settings.get("positive_numeric_columns", [])
        for row_no in content_rows:
            for column in nonnegative_columns:
                value = sheet.cell(row_no, column_numbers[column]).value
                if not _is_finite_number(value) or float(value) < 0:
                    numeric_invalid = True
                    result.issues.append(
                        PreflightIssue(
                            "NONNEGATIVE_NUMBER_REQUIRED",
                            "字段必须是大于或等于零的数字。",
                            row=row_no,
                            column=column,
                        )
                    )
            for column in optional_nonnegative_columns:
                value = sheet.cell(row_no, column_numbers[column]).value
                if _is_blank(value):
                    continue
                if not _is_finite_number(value) or float(value) < 0:
                    numeric_invalid = True
                    result.issues.append(
                        PreflightIssue(
                            "OPTIONAL_NONNEGATIVE_NUMBER_INVALID",
                            "字段可以留空；填写时必须是大于或等于零的数字。",
                            row=row_no,
                            column=column,
                        )
                    )
            for column in positive_columns:
                value = sheet.cell(row_no, column_numbers[column]).value
                if not _is_finite_number(value) or float(value) <= 0:
                    numeric_invalid = True
                    result.issues.append(
                        PreflightIssue(
                            "POSITIVE_NUMBER_REQUIRED",
                            "字段必须是大于零的数字。",
                            row=row_no,
                            column=column,
                        )
                    )
            for column in settings.get("binary_flag_columns", []):
                value = sheet.cell(row_no, column_numbers[column]).value
                if value not in (0, 1):
                    numeric_invalid = True
                    result.issues.append(
                        PreflightIssue(
                            "BINARY_FLAG_REQUIRED",
                            "标志字段只能填写 0 或 1。",
                            row=row_no,
                            column=column,
                        )
                    )
        result.checks["numeric_values_valid"] = not numeric_invalid

        pattern_invalid = False
        compiled_patterns: dict[str, re.Pattern[str]] = {}
        try:
            compiled_patterns = {
                column: re.compile(str(pattern))
                for column, pattern in settings.get("pattern_columns", {}).items()
            }
        except re.error:
            result.issues.append(
                PreflightIssue(
                    "PLATFORM_PROFILE_PATTERN_INVALID",
                    "管易平台预检配置包含无效的字段格式规则。",
                )
            )
            return result
        for row_no in content_rows:
            for column, pattern in compiled_patterns.items():
                value = sheet.cell(row_no, column_numbers[column]).value
                if _is_blank(value):
                    continue
                if not pattern.fullmatch(str(value).strip()):
                    pattern_invalid = True
                    result.issues.append(
                        PreflightIssue(
                            "FIELD_PATTERN_INVALID",
                            "字段值不符合已登记格式。",
                            row=row_no,
                            column=column,
                        )
                    )
        result.checks["pattern_values_valid"] = not pattern_invalid

        allowed_value_invalid = False
        for row_no in content_rows:
            for column, allowed in settings.get("allowed_values", {}).items():
                value = sheet.cell(row_no, column_numbers[column]).value
                if _is_blank(value):
                    continue
                if value not in allowed:
                    allowed_value_invalid = True
                    result.issues.append(
                        PreflightIssue(
                            "FIELD_VALUE_NOT_ALLOWED",
                            "字段值不在已登记的管易标准值中。",
                            row=row_no,
                            column=column,
                        )
                    )
        result.checks["allowed_values_valid"] = not allowed_value_invalid

        suffix_invalid = False
        for row_no in content_rows:
            for column, suffix in settings.get("required_text_suffixes", {}).items():
                value = sheet.cell(row_no, column_numbers[column]).value
                if _is_blank(value):
                    continue
                required_suffix = str(suffix)
                if not required_suffix or not str(value).endswith(required_suffix):
                    suffix_invalid = True
                    result.issues.append(
                        PreflightIssue(
                            "REQUIRED_TEXT_SUFFIX_MISSING",
                            "系统生成的平台单号缺少已确认的机器人尾缀。",
                            row=row_no,
                            column=column,
                        )
                    )
        result.checks["required_text_suffixes_valid"] = not suffix_invalid

        identifier_invalid = False
        for row_no in content_rows:
            for column in platform_profile.get("identifier_columns", []):
                cell = sheet.cell(row_no, column_numbers[column])
                if not _is_blank(cell.value) and cell.number_format != "@":
                    identifier_invalid = True
                    result.issues.append(
                        PreflightIssue(
                            "IDENTIFIER_TEXT_FORMAT_REQUIRED",
                            "标识符字段必须保持文本格式。",
                            row=row_no,
                            column=column,
                        )
                    )
        result.checks["identifier_formats_preserved"] = not identifier_invalid

        group_column = platform_profile.get("order_group_key", "平台单号")
        consistency_columns = settings.get("order_consistency_columns", [])
        grouped_first_rows: dict[str, int] = {}
        inconsistent = False
        for row_no in content_rows:
            group_value = sheet.cell(row_no, column_numbers[group_column]).value
            if _is_blank(group_value):
                continue
            key = str(group_value)
            first_row = grouped_first_rows.get(key)
            if first_row is None:
                grouped_first_rows[key] = row_no
                continue
            for column in consistency_columns:
                first = sheet.cell(first_row, column_numbers[column]).value
                current = sheet.cell(row_no, column_numbers[column]).value
                if not _same_value(first, current):
                    inconsistent = True
                    result.issues.append(
                        PreflightIssue(
                            "ORDER_FIELD_INCONSISTENT",
                            "同一平台单号的订单级字段不一致。",
                            row=row_no,
                            column=column,
                        )
                    )
        result.order_count = len(grouped_first_rows)
        result.checks["order_fields_consistent"] = not inconsistent

        result.status = "ready" if not result.issues else "blocked"
        return result
    finally:
        workbook.close()
