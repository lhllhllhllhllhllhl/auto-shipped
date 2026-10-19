from __future__ import annotations

import hashlib
from copy import copy
from dataclasses import dataclass, field, replace
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from openpyxl import load_workbook

from auto_shipped.catalog import (
    ProductCatalog,
    ProductResolution,
    WhiteLabelLibrary,
)
from auto_shipped.domain import (
    ClarificationRequest,
    DecisionItem,
    MappingReuseCandidate,
    ParsedOrder,
    deduplicate_clarifications,
    deduplicate_decisions,
    resolve_parsed_source_identity,
)
from auto_shipped.platforms.guanyi.order_policies import (
    GuanyiOrderPolicyError,
    ResolvedOrderItem,
    apply_contact_suffix_policy,
    compose_seller_remark,
    route_order_notes,
    resolve_sam_product_type_code,
    select_buyer_member,
    select_logistics_carrier,
)
from auto_shipped.platforms.guanyi.item_expansion import (
    ResolvedLineItem,
    expand_source_item,
)
from auto_shipped.platforms.guanyi.store_assignment import (
    GuanyiStoreAssignmentError,
    StoreAssignmentNotice,
    resolve_store_assignment,
)


class GuanyiCustomImportError(ValueError):
    """Raised when a blocked or incompatible batch reaches workbook rendering."""


@dataclass(slots=True)
class GuanyiBuildResult:
    lines: list[dict[str, Any]] = field(default_factory=list)
    clarifications: list[ClarificationRequest] = field(default_factory=list)
    decisions: list[DecisionItem] = field(default_factory=list)
    mapping_reuse_candidates: list[MappingReuseCandidate] = field(
        default_factory=list
    )
    mapping_evidence: list[dict[str, Any]] = field(default_factory=list)
    store_assignment_notices: list[dict[str, Any]] = field(default_factory=list)

    @property
    def status(self) -> str:
        if self.clarifications:
            return "needs_input"
        return "preview_ready" if self.decisions else "ready"


def _store_notice_record(
    notice: StoreAssignmentNotice,
    *,
    platform_order_no: str,
) -> dict[str, Any]:
    return {
        "code": notice.code,
        "scope": "order",
        "source_ref": f"订单:{platform_order_no}",
        "store": notice.store,
        "rule_id": notice.rule_id,
        "message": notice.message,
        "product": {
            "product_code": notice.product_code,
            "spec_code": notice.spec_code,
        },
        "blocking": False,
    }


def _deduplicate_store_assignment_notices(
    notices: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    unique: dict[tuple[str, str, str, str, str, str], dict[str, Any]] = {}
    for notice in notices:
        product = notice.get("product") or {}
        key = (
            str(notice.get("code") or ""),
            str(notice.get("source_ref") or ""),
            str(notice.get("store") or ""),
            str(notice.get("rule_id") or ""),
            str(product.get("product_code") or ""),
            str(product.get("spec_code") or ""),
        )
        unique[key] = notice
    return list(unique.values())


def _rule_value(rules: dict[str, Any], key: str) -> Any:
    return rules["defaults"][key]["value"]


def _field_strategy(rules: dict[str, Any], key: str, default: str) -> str:
    policy = (rules.get("field_output_policy") or {}).get(key) or {}
    return str(policy.get("strategy") or default)


def _output_value(strategy: str, value: Any, value_strategy: str, field: str) -> Any:
    if strategy == value_strategy:
        return value or ""
    if strategy == "blank":
        return ""
    raise GuanyiCustomImportError(f"未知{field}输出策略: {strategy}")


def _product_name_output_value(strategy: str, line_item: ResolvedLineItem) -> str:
    product = line_item.product
    item = line_item.source_item
    if strategy == "catalog_product_name":
        return product.product_name or ""
    if strategy == "expansion_component_name_or_catalog":
        return line_item.output_product_name or product.product_name or ""
    if strategy == "blank":
        return ""
    if strategy == "source_name_with_display_spec_for_standard_items":
        if line_item.expansion_rule_id:
            return product.product_name or ""
        source_name = str(item.source_product_name or "").strip()
        source_display_spec = str(item.source_display_spec or "").strip().strip("/")
        if not source_name:
            return product.product_name or ""
        if not source_display_spec or source_display_spec in source_name:
            return source_name
        return f"{source_name.rstrip('/')}/{source_display_spec}"
    raise GuanyiCustomImportError(f"未知商品名称输出策略: {strategy}")


def _rule_clarifications(rules: dict[str, Any]) -> list[ClarificationRequest]:
    requests: list[ClarificationRequest] = []
    platform_number = rules.get("platform_order_number", {})
    if not platform_number.get("confirmed", False):
        requests.append(
            ClarificationRequest(
                code="CONFIRM_PLATFORM_ORDER_NUMBER_STRATEGY",
                scope="batch",
                field="平台单号",
                question=platform_number.get("question")
                or "请确认当前来源应如何生成管易平台单号。",
                reason="平台单号是导入和防重的关键字段，不能使用未确认规则。",
                answer_type="confirmation",
            )
        )

    used_defaults = {
        "payment_amount": "支付金额",
        "unit_price": "价格",
        "freight": "运费",
        "warehouse_name": "仓库名称",
        "payment_method": "支付方式",
        "order_type": "订单类型",
        "is_mobile_order": "是否手机订单",
        "is_cod": "是否货到付款",
        "is_distributor_order": "是否分销商订单",
    }
    if not rules.get("buyer_member_policy"):
        used_defaults["buyer_member"] = "买家会员"
    store_policy = rules.get("store_assignment_policy") or {}
    if store_policy:
        if not store_policy.get("confirmed", False):
            requests.append(
                ClarificationRequest(
                    code="CONFIRM_GUANYI_STORE_ASSIGNMENT_POLICY",
                    scope="batch",
                    field="店铺",
                    question=store_policy.get("question")
                    or "请确认当前来源的管易店铺分配规则。",
                    reason="店铺不能仅由来源公司默认，当前商品店铺规则尚未确认。",
                    answer_type="confirmation",
                )
            )
    else:
        used_defaults["store"] = "店铺"
    defaults = rules.get("defaults", {})
    for key, column in used_defaults.items():
        rule = defaults.get(key)
        if not rule:
            requests.append(
                ClarificationRequest(
                    code="MISSING_GUANYI_REQUIRED_RULE",
                    scope="batch",
                    field=column,
                    question=f"请提供管易必填列“{column}”的取值规则。",
                    reason="平台规则配置中没有该必填字段。",
                    answer_type="text",
                )
            )
        elif not rule.get("confirmed", False):
            requests.append(
                ClarificationRequest(
                    code=f"CONFIRM_GUANYI_{key.upper()}",
                    scope="batch",
                    field=column,
                    question=rule.get("question") or f"请确认“{column}”的默认值。",
                    reason="转换过程会写入该平台字段，当前默认值尚未确认。",
                    answer_type="confirmation",
                )
            )

    used_field_policies = {
        "product_name": "商品名称",
        "spec_name": "规格名称",
        "order_created_at": "订单创建时间",
        "contact_columns": "联系电话/联系手机",
        "recipient_region": "省/市/区",
    }
    field_policies = rules.get("field_output_policy") or {}
    for key, column in used_field_policies.items():
        rule = field_policies.get(key)
        if rule and not rule.get("confirmed", False):
            requests.append(
                ClarificationRequest(
                    code=f"CONFIRM_GUANYI_FIELD_POLICY_{key.upper()}",
                    scope="batch",
                    field=column,
                    question=rule.get("question") or f"请确认“{column}”的输出规则。",
                    reason="该字段输出策略尚未确认。",
                    answer_type="confirmation",
                )
            )
    for policy_key, field_name in {
        "buyer_member_policy": "买家会员",
        "item_expansion_policy": "组合商品展开",
        "logistics_policy": "物流公司",
        "note_routing_policy": "收货地址/卖家备注",
        "seller_remark_policy": "卖家备注",
    }.items():
        policy = rules.get(policy_key)
        if policy_key == "item_expansion_policy" and not policy:
            continue
        if not policy:
            requests.append(
                ClarificationRequest(
                    code=f"MISSING_GUANYI_{policy_key.upper()}",
                    scope="batch",
                    field=field_name,
                    question=f"请补充“{field_name}”的业务策略。",
                    reason=f"平台规则缺少 {policy_key}。",
                    answer_type="text",
                )
            )
        elif not policy.get("confirmed", False):
            requests.append(
                ClarificationRequest(
                    code=f"CONFIRM_GUANYI_{policy_key.upper()}",
                    scope="batch",
                    field=field_name,
                    question=policy.get("question") or f"请确认“{field_name}”的业务策略。",
                    reason=f"{policy_key} 尚未确认。",
                    answer_type="confirmation",
                )
            )
    contact_suffix_policy = rules.get("contact_suffix_policy")
    if contact_suffix_policy and not contact_suffix_policy.get("confirmed", False):
        requests.append(
            ClarificationRequest(
                code="CONFIRM_GUANYI_CONTACT_SUFFIX_POLICY",
                scope="batch",
                field="联系电话/联系手机",
                question=contact_suffix_policy.get("question")
                or "请确认京东中通订单的联系电话尾码规则。",
                reason="联系电话尾码规则尚未确认。",
                answer_type="confirmation",
            )
        )
    return requests


def _resolve_product_with_profile_aliases(
    catalog: ProductCatalog,
    item: ParsedItem,
    *,
    source_profile_id: str,
    source_profile_aliases: tuple[str, ...],
    source_fingerprint: str,
) -> ProductResolution:
    """Reuse exact mappings from explicitly compatible legacy source profiles.

    The current source profile always wins.  Legacy profiles are consulted only
    when the current profile did not produce a confirmed result, and they may be
    reused only when every confirmed legacy hit points to the same Guanyi SKU.
    """

    kwargs = {
        "source_item_text": item.source_item_text,
        "source_brand": item.source_brand,
        "source_product_name": item.source_product_name,
        "source_spec": item.source_spec,
        "unit": item.unit,
        "source_fingerprint": source_fingerprint,
    }
    primary = catalog.resolve(
        item.source_product_code,
        source_profile_id=source_profile_id,
        **kwargs,
    )
    if primary.confirmed or not source_profile_aliases:
        return primary

    confirmed_by_target: dict[tuple[str, str], ProductResolution] = {}
    for alias in dict.fromkeys(source_profile_aliases):
        candidate = catalog.resolve(
            item.source_product_code,
            source_profile_id=alias,
            **kwargs,
        )
        if candidate.confirmed and candidate.product is not None:
            key = (candidate.product.product_code, candidate.product.spec_code)
            confirmed_by_target[key] = candidate
    if len(confirmed_by_target) == 1:
        candidate = next(iter(confirmed_by_target.values()))
        return replace(
            candidate,
            reason=(
                f"{candidate.reason}；通过公司库声明的旧来源入口兼容复用"
            ),
            match_method=f"compatible_source_profile:{candidate.match_method}",
        )
    if len(confirmed_by_target) > 1:
        return ProductResolution(
            status="ambiguous",
            external_sku=primary.external_sku,
            product=None,
            confirmed=False,
            reason="同公司兼容的旧来源入口指向多个不同管易商品，必须确认。",
            match_method="compatible_source_profile_conflict",
            identifier_type=primary.identifier_type,
            candidates=tuple(
                resolution.product
                for resolution in confirmed_by_target.values()
                if resolution.product is not None
            ),
        )
    return primary


def _platform_order_no(
    order: ParsedOrder,
    rules: dict[str, Any],
    platform_profile: dict[str, Any],
) -> str:
    number_rule = rules["platform_order_number"]
    strategy = number_rule["strategy"]
    identity = resolve_parsed_source_identity(order.source).value
    if strategy == "preferred_order_identity":
        base_number = identity
    elif strategy == "preferred_order_identity_with_affixes":
        prefix = str(number_rule.get("prefix") or "")
        suffix = str(number_rule.get("suffix") or "")
        base_number = f"{prefix}{identity}{suffix}"
    else:
        raise GuanyiCustomImportError(f"未知平台单号策略: {strategy}")

    identity = platform_profile.get("automation_identity_policy") or {}
    if not identity.get("confirmed", False):
        raise GuanyiCustomImportError("管易自动化订单标识策略尚未确认")
    identity_strategy = str(identity.get("strategy") or "")
    if identity_strategy == "preserve_base":
        return base_number
    if identity_strategy == "append_suffix":
        automation_suffix = str(identity.get("suffix") or "")
        if not automation_suffix:
            raise GuanyiCustomImportError("管易自动化订单尾缀不能为空")
        return f"{base_number}{automation_suffix}"
    raise GuanyiCustomImportError(
        f"未知管易自动化订单标识策略: {identity_strategy or '未配置'}"
    )


def _product_resolution_key(
    order: ParsedOrder,
    item: Any,
) -> tuple[str, str, str, str, str, str, str]:
    return (
        order.source.profile_id,
        order.source.file_fingerprint,
        str(item.source_product_code or ""),
        str(item.source_item_text or ""),
        str(item.source_brand or ""),
        str(item.source_product_name or ""),
        str(item.source_spec or ""),
        str(item.unit or ""),
    )


def _product_source_label(item: Any) -> str:
    base = str(
        item.source_product_code
        or item.source_item_text
        or item.source_product_name
        or "未提供商品标识"
    )
    source_spec = str(item.source_spec or "").strip()
    if source_spec and source_spec not in base:
        return f"{base} / {source_spec}"
    return base


def _mapping_evidence(
    catalog: ProductCatalog,
    item: Any,
    resolution: Any,
) -> dict[str, Any]:
    identity = resolution.canonical_identity or catalog.canonicalize_source_identity(
        source_item_text=item.source_item_text,
        source_brand=item.source_brand,
        source_product_name=item.source_product_name,
        source_spec=item.source_spec,
        unit=item.unit,
    )
    product = resolution.product
    return {
        "mapping_revision": catalog.mapping_revision,
        "mapping_snapshot_sha256": catalog.mapping_snapshot_sha256,
        "source_item_text": str(item.source_item_text or ""),
        "initial_split": {
            "brand": str(item.source_brand or ""),
            "product_name": str(item.source_product_name or ""),
            "source_spec": str(item.source_spec or ""),
            "unit": str(item.unit or ""),
        },
        "canonical_identity": identity.to_dict(),
        "mapping_ids": list(resolution.mapping_ids),
        "match_method": str(resolution.match_method or ""),
        "target": {
            "product_code": product.product_code if product else "",
            "spec_code": product.spec_code if product else "",
        },
    }


def _fingerprint_sha256(order: ParsedOrder) -> str:
    fingerprint = str(order.source.file_fingerprint or "")
    if fingerprint.startswith("sha256:"):
        fingerprint = fingerprint.removeprefix("sha256:")
    if len(fingerprint) == 64 and all(
        character in "0123456789abcdef" for character in fingerprint.casefold()
    ):
        return fingerprint.casefold()
    raise ValueError("SAM无跟团号订单缺少有效的来源文件SHA-256。")


def _allocate_sam_buyer_member_sequences(
    orders: list[ParsedOrder],
    resolved_items_by_order: dict[int, list[ResolvedOrderItem]],
    rules: dict[str, Any],
    *,
    business_datetime: datetime | None,
    state_dir: str | Path | None,
    source_company_id: str,
    gateway: Any = None,
) -> tuple[
    dict[int, tuple[str, str, str]],
    datetime | None,
    list[ClarificationRequest],
]:
    from auto_shipped.services.order_number_allocations import (
        OrderNumberAllocationError,
        reserve_daily_sequence,
        reserve_minute_sequence,
    )

    policy = rules.get("buyer_member_policy") or {}
    if (
        policy.get("strategy") != "sam_product_type_date_group"
        or policy.get("on_missing_group_number")
        not in {"generate_daily_sequence", "generate_minute_sequence"}
    ):
        return {}, business_datetime, []

    group_field = str(policy.get("group_number_field") or "").strip()
    missing_group_orders = [
        order
        for order in orders
        if not str(order.source_extensions.get(group_field) or "").strip()
    ]
    if not missing_group_orders:
        return {}, business_datetime, []

    clarifications: list[ClarificationRequest] = []
    type_codes: dict[int, str] = {}
    for order in missing_group_orders:
        try:
            type_codes[id(order)] = resolve_sam_product_type_code(
                policy,
                resolved_items_by_order.get(id(order), []),
            )
        except GuanyiOrderPolicyError as exc:
            clarifications.append(
                ClarificationRequest(
                    code="GUANYI_BUYER_MEMBER_POLICY_INVALID",
                    scope="order",
                    source_ref=f"订单:{resolve_parsed_source_identity(order.source).value}",
                    field="买家会员",
                    question="SAM商品类型无法唯一确定，请先补充或修正商品类型规则。",
                    reason=str(exc),
                    answer_type="text",
                )
            )
    if clarifications:
        return {}, business_datetime, clarifications

    timezone_name = str(policy.get("timezone") or "Asia/Shanghai")
    try:
        timezone = ZoneInfo(timezone_name)
    except ZoneInfoNotFoundError:
        return {}, business_datetime, [
            ClarificationRequest(
                code="GUANYI_BUYER_MEMBER_POLICY_INVALID",
                scope="batch",
                field="买家会员",
                question="SAM买家会员时区配置无效，请先修正规则。",
                reason=f"未知时区：{timezone_name}",
                answer_type="text",
            )
        ]
    frozen_datetime = business_datetime or datetime.now(timezone)
    if frozen_datetime.tzinfo is None:
        frozen_datetime = frozen_datetime.replace(tzinfo=timezone)
    else:
        frozen_datetime = frozen_datetime.astimezone(timezone)
    daily_mode = policy.get("on_missing_group_number") == "generate_daily_sequence"
    stamp_format = str(
        policy.get("fallback_stamp_format")
        or ("%y%m%d" if daily_mode else "%y%m%d%H%M")
    )
    batch_stamp = frozen_datetime.strftime(stamp_format)
    sequence_width = int(
        policy.get("fallback_sequence_width") or (3 if daily_mode else 2)
    )
    sequence_start = int(policy.get("fallback_sequence_start") or 1)
    platform_prefix = str(policy.get("platform_prefix") or "sam").strip()

    grouped: dict[tuple[str, str, str], list[ParsedOrder]] = {}
    for order in missing_group_orders:
        try:
            source_sha256 = _fingerprint_sha256(order)
        except ValueError as exc:
            clarifications.append(
                ClarificationRequest(
                    code="SAM_BUYER_MEMBER_SEQUENCE_ALLOCATION_FAILED",
                    scope="order",
                    source_ref=f"订单:{resolve_parsed_source_identity(order.source).value}",
                    field="买家会员",
                    question="无法为SAM无跟团号订单安全分配会员序号，请重新提供原始文件。",
                    reason=str(exc),
                    answer_type="file",
                )
            )
            continue
        grouped.setdefault(
            (
                str(order.source.profile_id),
                source_sha256,
                type_codes[id(order)],
            ),
            [],
        ).append(order)
    if clarifications:
        return {}, frozen_datetime, clarifications

    sequences: dict[int, tuple[str, str, str]] = {}
    for (profile_id, source_sha256, type_code), source_orders in grouped.items():
        try:
            if daily_mode:
                allocation = reserve_daily_sequence(
                    purpose="sam_buyer_member",
                    profile_id=profile_id,
                    company_id=source_company_id,
                    prefix=f"{platform_prefix}{type_code}",
                    business_date=frozen_datetime.strftime("%Y%m%d"),
                    source_sha256=source_sha256,
                    sequence_count=len(source_orders),
                    sequence_width=sequence_width,
                    sequence_start=sequence_start,
                    stamp_format=stamp_format,
                    sequence_separator=str(
                        policy.get("fallback_sequence_separator") or "_"
                    ),
                    source_partition=type_code,
                    state_dir=state_dir,
                    gateway=gateway,
                )
                allocated_stamp = datetime.strptime(
                    allocation.business_date,
                    "%Y%m%d",
                ).strftime(allocation.stamp_format)
                allocated_separator = allocation.sequence_separator
            else:
                allocation = reserve_minute_sequence(
                    purpose="sam_buyer_member",
                    profile_id=profile_id,
                    prefix=platform_prefix,
                    batch_minute=batch_stamp,
                    source_sha256=source_sha256,
                    sequence_count=len(source_orders),
                    sequence_width=sequence_width,
                    sequence_start=sequence_start,
                    stamp_format=stamp_format,
                    state_dir=state_dir,
                )
                allocated_stamp = allocation.batch_minute
                allocated_separator = str(
                    policy.get("fallback_sequence_separator") or "_"
                )
        except OrderNumberAllocationError as exc:
            return {}, frozen_datetime, [
                ClarificationRequest(
                    code="SAM_BUYER_MEMBER_SEQUENCE_ALLOCATION_FAILED",
                    scope="batch",
                    field="买家会员",
                    question="无法为SAM无跟团号订单安全分配会员序号，请核对来源文件后重试。",
                    reason=str(exc),
                    answer_type="file",
                )
            ]
        for offset, order in enumerate(source_orders):
            sequence_no = allocation.sequence_start + offset
            sequences[id(order)] = (
                allocated_stamp,
                f"{sequence_no:0{allocation.sequence_width}d}",
                allocated_separator,
            )
    return sequences, frozen_datetime, []


def _product_candidate_summary(candidates: tuple[Any, ...]) -> str:
    rendered: list[str] = []
    for index, product in enumerate(candidates[:5], start=1):
        name = str(product.product_name or "未命名商品")
        spec_name = str(product.spec_name or "无规格名称")
        spec_code = str(product.spec_code or "无规格代码")
        rendered.append(
            f"{index}. {name} / {spec_name}（{product.product_code} / {spec_code}）"
        )
    return "；".join(rendered)


def build_custom_import_lines(
    orders: list[ParsedOrder],
    catalog: ProductCatalog,
    rules: dict[str, Any],
    platform_profile: dict[str, Any],
    *,
    white_label_library: WhiteLabelLibrary | None = None,
    business_date: date | None = None,
    business_datetime: datetime | None = None,
    buyer_member_sequence_state_dir: str | Path | None = None,
    buyer_member_sequence_gateway: Any = None,
    source_company_id: str = "",
    source_profile_aliases: tuple[str, ...] = (),
    official_mapping_ready: bool = True,
    official_mapping_error: str = "",
) -> GuanyiBuildResult:
    result = GuanyiBuildResult()
    result.clarifications.extend(_rule_clarifications(rules))

    resolved_products: dict[tuple[str, str, str, str, str, str, str], Any] = {}
    resolved_lines_by_order: dict[int, list[ResolvedLineItem]] = {}
    official_mapping_clarification_added = False
    for order in orders:
        order_lines: list[ResolvedLineItem] = []
        for item in order.items:
            expansion = expand_source_item(
                item,
                catalog,
                rules,
                company_id=source_company_id,
                source_profile_id=order.source.profile_id,
                source_profile_aliases=source_profile_aliases,
                source_fingerprint=order.source.file_fingerprint,
                registered_rules=catalog.item_expansion_rules,
            )
            if expansion.failed:
                if not official_mapping_ready:
                    if not official_mapping_clarification_added:
                        result.clarifications.append(
                            ClarificationRequest(
                                code="FEISHU_OFFICIAL_MAPPING_UNAVAILABLE",
                                scope="batch",
                                field="飞书正式商品映射",
                                question=(
                                    "无法读取并验证飞书正式商品映射和组合商品规则。"
                                    "请先完成飞书授权或正式快照同步后重试。"
                                ),
                                reason=(
                                    official_mapping_error
                                    or "正式映射快照不存在、读取失败或未通过完整性校验。"
                                ),
                                answer_type="text",
                                next_action="sync_feishu_official_mappings",
                            )
                        )
                        official_mapping_clarification_added = True
                else:
                    result.clarifications.append(
                        ClarificationRequest(
                            code=expansion.error_code,
                            scope="batch",
                            source_ref=f"商品:{_product_source_label(item)}",
                            field="组合商品展开",
                            question=(
                                "当前组合商品没有唯一、已确认的展开方式。请确认目标商品、"
                                "规格和数量倍数，并说明仅本次使用还是以后沿用。"
                            ),
                            reason=expansion.reason,
                            answer_type="text",
                            next_action="use_once_or_add_item_expansion_if_reusable",
                        )
                    )
                continue
            if expansion.applied:
                order_lines.extend(expansion.lines)
                continue

            if not official_mapping_ready:
                if not official_mapping_clarification_added:
                    result.clarifications.append(
                        ClarificationRequest(
                            code="FEISHU_OFFICIAL_MAPPING_UNAVAILABLE",
                            scope="batch",
                            field="飞书正式商品映射",
                            question=(
                                "无法读取并验证飞书正式商品映射和组合商品规则。"
                                "请先完成飞书授权或正式快照同步后重试。"
                            ),
                            reason=(
                                official_mapping_error
                                or "正式映射快照不存在、读取失败或未通过完整性校验。"
                            ),
                            answer_type="text",
                            next_action="sync_feishu_official_mappings",
                        )
                    )
                    official_mapping_clarification_added = True
                continue

            resolution_key = _product_resolution_key(order, item)
            resolution = resolved_products.get(resolution_key)
            if resolution is None:
                resolution = _resolve_product_with_profile_aliases(
                    catalog,
                    item,
                    source_profile_id=order.source.profile_id,
                    source_profile_aliases=source_profile_aliases,
                    source_fingerprint=order.source.file_fingerprint,
                )
                resolved_products[resolution_key] = resolution
            result.mapping_evidence.append(
                _mapping_evidence(catalog, item, resolution)
            )
            source_label = _product_source_label(item)
            source_ref = f"商品:{source_label}"
            for warning in resolution.review_warnings:
                result.decisions.append(
                    DecisionItem(
                        code=str(warning.get("code") or "REVIEW_PRODUCT_MAPPING"),
                        source_ref=source_ref,
                        field="来源商品名称与编码映射",
                        proposed_value={
                            "source_profile_id": order.source.profile_id,
                            "company_id": source_company_id,
                            "source_identifier_type": resolution.identifier_type,
                            "source_product_code": item.source_product_code,
                            "source_item_text": item.source_item_text,
                            "source_brand": item.source_brand,
                            "source_product_name": item.source_product_name,
                            "source_spec": item.source_spec,
                            "source_fingerprint": order.source.file_fingerprint,
                            "product_code": (
                                resolution.product.product_code
                                if resolution.product
                                else ""
                            ),
                            "spec_code": (
                                resolution.product.spec_code
                                if resolution.product
                                else ""
                            ),
                            "match_method": resolution.match_method,
                            "canonical_identity": (
                                resolution.canonical_identity.to_dict()
                                if resolution.canonical_identity
                                else catalog.canonicalize_source_identity(
                                    source_item_text=item.source_item_text,
                                    source_brand=item.source_brand,
                                    source_product_name=item.source_product_name,
                                    source_spec=item.source_spec,
                                    unit=item.unit,
                                ).to_dict()
                            ),
                        },
                        confirmation_question=(
                            "来源编码/条码已优先按正式映射生成预览，但来源商品名称"
                            "与已登记身份存在明显差异。请确认本批仍使用该编码映射，"
                            "还是原订单商品字段提取有误。"
                        ),
                        reason=str(warning.get("message") or resolution.reason),
                        persistence_options=("current_batch",),
                        next_action="confirm_code_mapping_for_current_batch_or_fix_source_extraction",
                    )
                )
            if resolution.match_method == "shipping_expression_forbidden":
                result.clarifications.append(
                    ClarificationRequest(
                        code="SHIPPING_PRODUCT_FORBIDDEN",
                        scope="batch",
                        source_ref=source_ref,
                        field="发货商品",
                        question="裸棒属于试吃商品，普通发货流程已禁用；请确认改用哪个正式发货商品。",
                        reason=resolution.reason,
                        answer_type="text",
                    )
                )
            elif resolution.match_method == "package_semantics_unregistered":
                result.clarifications.append(
                    ClarificationRequest(
                        code="CONFIRM_PACKAGE_SEMANTICS",
                        scope="batch",
                        source_ref=source_ref,
                        field="包装含义",
                        question=(
                            f"请确认来源商品“{source_label}”中的包装单位表示彩袋单棒装，"
                            "还是表示家庭装中的总根数；并说明仅本次使用还是以后沿用。"
                        ),
                        reason=resolution.reason,
                        answer_type="text",
                        next_action="record_package_semantics_proposal_if_reusable",
                    )
                )
            elif (
                resolution.match_method == "fuzzy_name_candidates"
                and resolution.confidence_tier == "blocking"
            ):
                candidate_text = _product_candidate_summary(resolution.candidates)
                result.clarifications.append(
                    ClarificationRequest(
                        code="CONFIRM_LOW_CONFIDENCE_GUANYI_PRODUCT_CANDIDATE",
                        scope="batch",
                        source_ref=source_ref,
                        field="商品代码/规格代码",
                        question=(
                            f"来源商品“{source_label}”没有安全的唯一商品匹配。"
                            + (
                                f"系统只找到以下低可信候选（仅供查找，未作为默认值）：{candidate_text}。"
                                if candidate_text
                                else ""
                            )
                            + "请提供或确认正确的管易商品代码、规格代码、条码或准确名称，"
                            "并说明仅本次使用，还是以后沿用。"
                        ),
                        reason=resolution.reason,
                        answer_type="text",
                        next_action="resolve_target_hint_then_use_once_or_persist",
                    )
                )
            elif len(resolution.candidates) > 1:
                result.clarifications.append(
                    ClarificationRequest(
                        code="CONFIRM_GUANYI_PRODUCT_CANDIDATE",
                        scope="batch",
                        source_ref=source_ref,
                        field="商品代码/规格代码",
                        question=(
                            f"请确认来源商品“{source_label}”对应以下哪个管易商品规格："
                            f"{_product_candidate_summary(resolution.candidates)}。"
                            "同时请说明仅本次使用，还是以后都沿用。"
                        ),
                        reason=resolution.reason or "商品存在一个或多个待确认候选。",
                        answer_type="text",
                        next_action="use_once_or_submit_mapping_if_reusable",
                    )
                )
            elif resolution.status == "unconfirmed" and (
                resolution.product or len(resolution.candidates) == 1
            ):
                proposed_product = resolution.product or resolution.candidates[0]
                future_allowed = resolution.match_method != "source_mapping_identity_guard"
                persistence_options = (
                    ("current_batch", "future")
                    if future_allowed
                    else ("current_batch",)
                )
                global_alias_eligible = (
                    future_allowed and resolution.identifier_type == "product_name"
                )
                if global_alias_eligible:
                    scope_prompt = (
                        "如果判断正确，请选择仅本次使用、以后仅当前来源沿用，"
                        "还是以后所有来源沿用。"
                    )
                elif future_allowed:
                    scope_prompt = "如果判断正确，请选择仅本次使用，还是以后遇到相同来源时沿用。"
                else:
                    scope_prompt = "该来源已有正式映射但身份信息不一致，只能先按本批次确认；长期修改请联系管理员。"
                result.decisions.append(
                    DecisionItem(
                        code="CONFIRM_GUANYI_PRODUCT_MAPPING",
                        source_ref=source_ref,
                        field="商品代码/规格代码",
                        proposed_value={
                            "source_profile_id": order.source.profile_id,
                            "company_id": source_company_id,
                            "source_identifier_type": resolution.identifier_type,
                            "source_product_code": item.source_product_code,
                            "source_product_name": item.source_product_name,
                            "source_spec": item.source_spec,
                            "source_fingerprint": order.source.file_fingerprint,
                            "product_code": proposed_product.product_code,
                            "spec_code": proposed_product.spec_code,
                            "product_name": proposed_product.product_name,
                            "spec_name": proposed_product.spec_name,
                            "match_method": resolution.match_method,
                            "confidence_tier": resolution.confidence_tier,
                            "candidate_evidence": [
                                evidence.public_evidence()
                                for evidence in resolution.candidate_evidence
                                if evidence.product == proposed_product
                            ],
                            "global_alias_eligible": global_alias_eligible,
                        },
                        confirmation_question=(
                            f"系统暂按“{source_label}”→管易商品代码"
                            f"“{proposed_product.product_code}”、规格代码"
                            f"“{proposed_product.spec_code}”生成预览。{scope_prompt}"
                        ),
                        reason=resolution.reason or "当前商品映射被标记为待确认。",
                        persistence_options=persistence_options,
                        future_scope_options=(
                            ("source", "global") if global_alias_eligible else ()
                        ),
                        next_action="use_once_or_submit_mapping_if_reusable",
                    )
                )
                resolution = replace(
                    resolution,
                    status="confirmed",
                    product=proposed_product,
                    confirmed=False,
                )
            elif resolution.status != "confirmed":
                result.clarifications.append(
                    ClarificationRequest(
                        code="PROVIDE_GUANYI_PRODUCT_MAPPING",
                        scope="batch",
                        source_ref=source_ref,
                        field="商品代码/规格代码",
                        question=(
                            f"请提供来源商品“{source_label}”对应的管易商品代码和规格代码，"
                            "并说明仅本次使用，还是以后都沿用。"
                        ),
                        reason=resolution.reason or "商品主数据中没有唯一匹配。",
                        answer_type="text",
                        next_action="use_once_or_submit_mapping_if_reusable",
                    )
                )
            if resolution.product and resolution.status == "confirmed":
                if (
                    resolution.mapping_authority == "user_batch_confirmation"
                    and resolution.reuse_scope == "current_batch"
                ):
                    identifier_type = resolution.identifier_type or "product_code"
                    source_value = (
                        item.source_product_name
                        if identifier_type == "product_name"
                        else item.source_product_code
                    )
                    result.mapping_reuse_candidates.append(
                        MappingReuseCandidate(
                            source_profile_id=order.source.profile_id,
                            company_id=source_company_id,
                            identifier_type=identifier_type,
                            source_value=str(source_value or "").strip(),
                            source_spec=str(item.source_spec or "").strip(),
                            source_fingerprint=order.source.file_fingerprint,
                            product_code=resolution.product.product_code,
                            spec_code=resolution.product.spec_code,
                            product_name=resolution.product.product_name,
                            spec_name=resolution.product.spec_name,
                            future_scope_options=(
                                ("source", "global")
                                if identifier_type == "product_name"
                                else ("source",)
                            ),
                        )
                    )
                forbidden_reason = catalog.shipping_forbidden_reason(resolution.product)
                if forbidden_reason:
                    result.clarifications.append(
                        ClarificationRequest(
                            code="SHIPPING_PRODUCT_FORBIDDEN",
                            scope="batch",
                            source_ref=source_ref,
                            field="发货商品",
                            question="该商品属于试吃裸棒，普通发货流程已禁用；请确认改用哪个正式发货商品。",
                            reason=forbidden_reason,
                            answer_type="text",
                        )
                    )
                    continue
                target_quantity = item.quantity
                if resolution.quantity_strategy == "divide_by_sticks_per_target_unit":
                    sticks_per_unit = resolution.sticks_per_target_unit or 0
                    quotient = item.quantity / sticks_per_unit if sticks_per_unit else 0
                    if sticks_per_unit <= 0 or abs(quotient - round(quotient)) > 1e-9:
                        result.clarifications.append(
                            ClarificationRequest(
                                code="CONFIRM_PACKAGE_QUANTITY_CONVERSION",
                                scope="batch",
                                source_ref=source_ref,
                                field="包装数量换算",
                                question=(
                                    f"来源数量{item.quantity:g}无法按每件{sticks_per_unit or '未知'}根"
                                    "完整换算，请确认是否需要家庭装与单根装混合拆分。"
                                ),
                                reason="包装语义规则只允许整件换算，不能自动决定余数商品。",
                                answer_type="text",
                            )
                        )
                        continue
                    target_quantity = int(round(quotient))
                elif resolution.quantity_strategy == "requires_bundle_expansion":
                    result.clarifications.append(
                        ClarificationRequest(
                            code="GUANYI_ITEM_EXPANSION_UNCONFIRMED",
                            scope="batch",
                            source_ref=source_ref,
                            field="组合商品展开",
                            question="该包装含义需要组合商品展开，请先登记具体组成和数量倍数。",
                            reason="包装语义只描述了组合类型，没有可执行的组成规则。",
                            answer_type="text",
                        )
                    )
                    continue
                order_lines.append(
                    ResolvedLineItem(
                        source_item=item,
                        product=resolution.product,
                        quantity=target_quantity,
                    )
                )
        resolved_lines_by_order[id(order)] = order_lines

    result.clarifications = deduplicate_clarifications(result.clarifications)
    result.decisions = deduplicate_decisions(result.decisions)
    evidence_by_key: dict[str, dict[str, Any]] = {}
    for evidence in result.mapping_evidence:
        key = hashlib.sha256(
            repr(
                (
                    evidence.get("source_item_text"),
                    evidence.get("initial_split"),
                    evidence.get("canonical_identity"),
                    evidence.get("mapping_ids"),
                    evidence.get("match_method"),
                    evidence.get("target"),
                )
            ).encode("utf-8")
        ).hexdigest()
        evidence_by_key[key] = evidence
    result.mapping_evidence = list(evidence_by_key.values())
    if result.clarifications:
        # Blocking facts must be answered before a trustworthy preview can be
        # rendered.  Defer non-blocking candidate/recipient decisions until
        # the blocking items are resolved and the whole batch can produce one
        # preview.  They are deterministically recomputed on the next run, so
        # the user sees one focused mid-flow question instead of a mixed list
        # of blockers and preview-only confirmations.
        result.decisions = []
        return result

    resolved_order_items_by_order = {
        id(order): [
            ResolvedOrderItem(
                quantity=line_item.quantity,
                product_code=line_item.product.product_code,
                spec_code=line_item.product.spec_code,
                product_name=line_item.product.product_name,
                spec_name=line_item.product.spec_name,
            )
            for line_item in resolved_lines_by_order[id(order)]
        ]
        for order in orders
    }
    (
        buyer_member_sequences,
        frozen_business_datetime,
        sequence_clarifications,
    ) = _allocate_sam_buyer_member_sequences(
        orders,
        resolved_order_items_by_order,
        rules,
        business_datetime=business_datetime,
        state_dir=buyer_member_sequence_state_dir,
        source_company_id=source_company_id,
        gateway=buyer_member_sequence_gateway,
    )
    if sequence_clarifications:
        result.clarifications.extend(sequence_clarifications)
        result.clarifications = deduplicate_clarifications(result.clarifications)
        return result

    columns = platform_profile["columns"]
    product_name_strategy = _field_strategy(
        rules, "product_name", "catalog_product_name"
    )
    spec_name_strategy = _field_strategy(rules, "spec_name", "catalog_spec_name")
    order_created_at_strategy = _field_strategy(
        rules, "order_created_at", "source_ordered_at"
    )
    contact_strategy = _field_strategy(
        rules, "contact_columns", "mobile_to_both_phone_preserved"
    )
    recipient_region_strategy = _field_strategy(
        rules, "recipient_region", "source_components"
    )
    seen_platform_numbers: dict[str, str] = {}
    for order in orders:
        platform_no = _platform_order_no(order, rules, platform_profile)
        previous = seen_platform_numbers.get(platform_no)
        if previous and previous != order.record_key:
            result.clarifications.append(
                ClarificationRequest(
                    code="DUPLICATE_GUANYI_PLATFORM_ORDER_NO",
                    scope="batch",
                    field="平台单号",
                    question=f"平台单号“{platform_no}”对应多张订单，请确认单号生成规则。",
                    reason="管易平台单号必须唯一标识一张订单。",
                    answer_type="text",
                )
            )
            continue
        seen_platform_numbers[platform_no] = order.record_key
        shipment = order.shipment_facts
        resolved_line_items = resolved_lines_by_order[id(order)]
        resolved_order_items = resolved_order_items_by_order[id(order)]
        try:
            store_assignment = (
                resolve_store_assignment(
                    resolved_order_items,
                    rules,
                    source_extensions=order.source_extensions,
                    source_company_id=source_company_id,
                    source_profile_id=order.source.profile_id,
                    manual_store_routes=catalog.manual_store_routes,
                )
                if rules.get("store_assignment_policy")
                else None
            )
            store_name = (
                store_assignment.store
                if store_assignment is not None
                else _rule_value(rules, "store")
            )
            if store_assignment is not None:
                result.store_assignment_notices.extend(
                    _store_notice_record(
                        notice,
                        platform_order_no=platform_no,
                    )
                    for notice in store_assignment.notices
                )
        except GuanyiStoreAssignmentError as exc:
            result.clarifications.append(
                ClarificationRequest(
                    code="CONFIRM_GUANYI_STORE_ASSIGNMENT",
                    scope="order",
                    source_ref=f"订单:{platform_no}",
                    field="店铺",
                    question="当前商品没有唯一、已确认的管易店铺，请确认应使用哪个店铺。",
                    reason=str(exc),
                    answer_type="text",
                )
            )
            continue
        try:
            carrier_name = select_logistics_carrier(order, resolved_order_items, rules)
        except GuanyiOrderPolicyError as exc:
            result.clarifications.append(
                ClarificationRequest(
                    code="GUANYI_LOGISTICS_POLICY_INVALID",
                    scope="batch",
                    source_ref=f"订单:{platform_no}",
                    field="物流公司",
                    question="当前物流规则无法安全执行，请先修正规则配置。",
                    reason=str(exc),
                    answer_type="text",
                )
            )
            continue
        try:
            generated_member = buyer_member_sequences.get(id(order))
            buyer_member = select_buyer_member(
                order,
                rules,
                resolved_order_items,
                business_date=business_date,
                business_datetime=frozen_business_datetime,
                generated_stamp=(generated_member[0] if generated_member else None),
                generated_sequence=(generated_member[1] if generated_member else None),
                generated_separator=(generated_member[2] if generated_member else None),
            )
        except GuanyiOrderPolicyError as exc:
            result.clarifications.append(
                ClarificationRequest(
                    code="GUANYI_BUYER_MEMBER_POLICY_INVALID",
                    scope="order",
                    source_ref=f"订单:{platform_no}",
                    field="买家会员",
                    question="当前买家会员规则无法安全执行，请先修正规则配置。",
                    reason=str(exc),
                    answer_type="text",
                )
            )
            continue
        note_routing_policy = rules.get("note_routing_policy") or {}
        explicit_extension_field = str(
            note_routing_policy.get("explicit_extension_field")
            or "delivery_instruction"
        )
        try:
            routed_notes = route_order_notes(
                order.recipient.full_address,
                order.source_note,
                note_routing_policy,
                explicit_delivery_instruction=order.source_extensions.get(
                    explicit_extension_field
                ),
            )
        except GuanyiOrderPolicyError as exc:
            result.clarifications.append(
                ClarificationRequest(
                    code="GUANYI_NOTE_ROUTING_POLICY_INVALID",
                    scope="batch",
                    source_ref=f"订单:{platform_no}",
                    field="收货地址/卖家备注",
                    question="当前备注分流规则无法安全执行，请先修正规则配置。",
                    reason=str(exc),
                    answer_type="text",
                )
            )
            continue
        seller_policy = rules.get("seller_remark_policy") or {}
        seller_strategy = str(seller_policy.get("strategy") or "")
        if seller_strategy != "source_note_then_system_notes_deduplicated":
            result.clarifications.append(
                ClarificationRequest(
                    code="GUANYI_SELLER_REMARK_POLICY_INVALID",
                    scope="batch",
                    source_ref=f"订单:{platform_no}",
                    field="卖家备注",
                    question="当前卖家备注规则无法安全执行，请先修正规则配置。",
                    reason=f"未知卖家备注策略: {seller_strategy or '未配置'}",
                    answer_type="text",
                )
            )
            continue
        separator = str(seller_policy.get("separator") or "；")
        system_notes: list[str] = []
        if white_label_library and white_label_library.matches_any(
            resolved_order_items
        ):
            system_notes.append(white_label_library.remark_token)
        seller_remark = compose_seller_remark(
            routed_notes.seller_source_note,
            system_notes,
            separator=separator,
        )
        if contact_strategy == "mobile_to_both_phone_preserved":
            contact_phone = order.recipient.mobile or order.recipient.phone
            contact_mobile = order.recipient.mobile
        elif contact_strategy == "source_split":
            contact_phone = order.recipient.phone
            contact_mobile = order.recipient.mobile
        else:
            raise GuanyiCustomImportError(
                f"未知联系电话/联系手机输出策略: {contact_strategy}"
            )
        try:
            resolved_contacts = apply_contact_suffix_policy(
                order,
                contact_phone,
                contact_mobile,
                rules.get("contact_suffix_policy"),
            )
        except GuanyiOrderPolicyError as exc:
            result.clarifications.append(
                ClarificationRequest(
                    code="GUANYI_CONTACT_SUFFIX_POLICY_INVALID",
                    scope="order",
                    source_ref=f"订单:{platform_no}",
                    field="联系电话/联系手机",
                    question="请确认京东中通订单姓名后的四位码，并确保姓名、地址和已有号码尾码一致。",
                    reason=str(exc),
                    answer_type="text",
                )
            )
            continue
        contact_phone = resolved_contacts.phone
        contact_mobile = resolved_contacts.mobile
        if recipient_region_strategy == "source_components":
            province = order.recipient.province
            city = order.recipient.city
            district = order.recipient.district
        elif recipient_region_strategy == "blank":
            province = city = district = ""
        else:
            raise GuanyiCustomImportError(
                f"未知省/市/区输出策略: {recipient_region_strategy}"
            )
        for line_item in resolved_line_items:
            item = line_item.source_item
            product = line_item.product
            line = {column: "" for column in columns}
            line.update(
                {
                    "店铺": store_name,
                    "平台单号": platform_no,
                    "买家会员": buyer_member,
                    "支付金额": _rule_value(rules, "payment_amount"),
                    "商品名称": _product_name_output_value(
                        product_name_strategy,
                        line_item,
                    ),
                    "商品代码": product.product_code,
                    "规格代码": product.spec_code,
                    "规格名称": _output_value(
                        spec_name_strategy,
                        product.spec_name,
                        "catalog_spec_name",
                        "规格名称",
                    ),
                    "是否赠品": 0,
                    "数量": line_item.quantity,
                    "价格": (
                        item.unit_price
                        if item.unit_price is not None
                        and not line_item.expansion_rule_id
                        else _rule_value(rules, "unit_price")
                    ),
                    "运费": (
                        order.amounts.freight
                        if order.amounts and order.amounts.freight is not None
                        else _rule_value(rules, "freight")
                    ),
                    "买家留言": order.buyer_message or "",
                    "收货人": order.recipient.name,
                    "联系电话": contact_phone,
                    "联系手机": contact_mobile,
                    "收货地址": routed_notes.recipient_address,
                    "省": province,
                    "市": city,
                    "区": district,
                    "订单创建时间": _output_value(
                        order_created_at_strategy,
                        order.ordered_at,
                        "source_ordered_at",
                        "订单创建时间",
                    ),
                    "订单付款时间": order.paid_at or "",
                    "物流单号": shipment.tracking_no if shipment else "",
                    "物流公司": carrier_name,
                    "卖家备注": seller_remark,
                    "是否手机订单": _rule_value(rules, "is_mobile_order"),
                    "是否货到付款": _rule_value(rules, "is_cod"),
                    "支付方式": _rule_value(rules, "payment_method"),
                    "仓库名称": _rule_value(rules, "warehouse_name"),
                    "订单类型": _rule_value(rules, "order_type"),
                    "是否分销商订单": _rule_value(rules, "is_distributor_order"),
                }
            )
            result.lines.append(line)

    result.clarifications = deduplicate_clarifications(result.clarifications)
    result.store_assignment_notices = _deduplicate_store_assignment_notices(
        result.store_assignment_notices
    )
    if result.clarifications:
        result.lines = []
    return result


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def render_custom_import(
    template_path: str | Path,
    output_path: str | Path,
    lines: list[dict[str, Any]],
    platform_profile: dict[str, Any],
) -> Path:
    if not lines:
        raise GuanyiCustomImportError("没有可导出的管易订单行")

    template = Path(template_path)
    expected_hash = platform_profile.get("template_sha256")
    if expected_hash and _sha256(template) != expected_hash:
        raise GuanyiCustomImportError("管易模板文件与已确认版本不一致")

    workbook = load_workbook(template, read_only=False, data_only=False)
    sheet_name = platform_profile["sheet_name"]
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise GuanyiCustomImportError(f"管易模板缺少工作表: {sheet_name}")
    sheet = workbook[sheet_name]
    header_row = int(platform_profile["header_row"])
    headers = [cell.value for cell in sheet[header_row] if cell.value is not None]
    if headers != platform_profile["columns"]:
        workbook.close()
        raise GuanyiCustomImportError("管易模板表头与平台配置不一致")

    donor_row = header_row + 1
    donor_styles = [copy(sheet.cell(donor_row, col)._style) for col in range(1, len(headers) + 1)]
    donor_height = sheet.row_dimensions[donor_row].height
    if sheet.max_row > header_row:
        sheet.delete_rows(header_row + 1, sheet.max_row - header_row)

    identifier_columns = set(platform_profile.get("identifier_columns", []))
    for row_offset, line in enumerate(lines, start=1):
        row_no = header_row + row_offset
        if donor_height is not None:
            sheet.row_dimensions[row_no].height = donor_height
        for col_no, column in enumerate(headers, start=1):
            cell = sheet.cell(row_no, col_no)
            cell._style = copy(donor_styles[col_no - 1])
            value = line.get(column, "")
            if column in identifier_columns and value not in (None, ""):
                value = str(value)
                cell.number_format = "@"
            cell.value = value

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()

    verification = load_workbook(output, read_only=False, data_only=False)
    try:
        sheet = verification[sheet_name]
        saved_headers = [cell.value for cell in sheet[header_row] if cell.value is not None]
        if saved_headers != headers:
            raise GuanyiCustomImportError("保存后的管易模板表头发生变化")
        if sheet.max_row != header_row + len(lines):
            raise GuanyiCustomImportError("保存后的管易订单行数不正确")
    finally:
        verification.close()
    return output
