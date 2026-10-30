from __future__ import annotations

import hashlib
import math
import uuid
import warnings
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from auto_shipped.domain import (
    ClarificationRequest,
    OrderIdentityError,
    ParsedItem,
    ParsedOrder,
    ParsedRecipient,
    ParsedShipmentFacts,
    ParsedSource,
    deduplicate_clarifications,
    resolve_order_identity,
)

from .tiantian_parsed import ParseResult
from .value_validation import identifier_requires_text, normalize_contact


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _quantity(value: Any) -> float | None:
    text = _text(value)
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return number if math.isfinite(number) and number > 0 else None


def _clean_source_note(value: Any, policy: dict[str, Any]) -> str:
    text = _text(value)
    if not policy:
        return text
    strategy = str(policy.get("strategy") or "").strip()
    if strategy != "remove_exact_phrases_preserve_remainder":
        raise ValueError(f"不支持的来源备注清洗策略：{strategy or '空'}")
    for raw_phrase in policy.get("remove_exact_phrases") or []:
        phrase = _text(raw_phrase)
        if phrase:
            text = text.replace(phrase, "")
    return text.strip(" \t\r\n;；,，/|")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


class ConfiguredExcelParsedAdapter:
    """Parse row-based order workbooks through a versioned field map."""

    adapter_id = "configured_excel_order_parsed"

    def parse(self, source_path: str | Path, profile: dict[str, Any]) -> ParseResult:
        path = Path(source_path)
        detection = profile["detection"]
        sheet_name = str(detection["sheet_name"])
        header_row = int(detection.get("header_row", 1))
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Workbook contains no default style.*")
            workbook = load_workbook(path, read_only=False, data_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                return ParseResult(
                    clarifications=[
                        ClarificationRequest(
                            code="MISSING_SOURCE_SHEET",
                            scope="file",
                            source_ref=path.name,
                            question=f"文件中缺少工作表“{sheet_name}”，请确认是否选择了正确文件。",
                            reason="来源适配器无法找到已配置工作表。",
                            answer_type="file",
                        )
                    ]
                )
            sheet = workbook[sheet_name]
            headers = [_text(cell.value) for cell in sheet[header_row]]
            missing = [
                name for name in detection["required_headers"] if name not in headers
            ]
            if missing:
                return ParseResult(
                    clarifications=[
                        ClarificationRequest(
                            code="MISSING_SOURCE_HEADERS",
                            scope="file",
                            source_ref=path.name,
                            question="来源Excel缺少必要列，请重新导出正确版本或确认是否需要新增格式。",
                            reason=f"缺少列：{', '.join(missing)}",
                            answer_type="file",
                        )
                    ]
                )
            rows: list[tuple[int, dict[str, Any]]] = []
            for row_no, cells in enumerate(
                sheet.iter_rows(min_row=header_row + 1, values_only=True),
                start=header_row + 1,
            ):
                row = {
                    headers[index]: value
                    for index, value in enumerate(cells)
                    if index < len(headers)
                }
                if any(value not in (None, "") for value in row.values()):
                    rows.append((row_no, row))
        finally:
            workbook.close()
        return self.parse_rows(rows, path, profile)

    def parse_rows(
        self,
        rows: Iterable[tuple[int, dict[str, Any]]],
        source_path: Path,
        profile: dict[str, Any],
    ) -> ParseResult:
        result = ParseResult()
        fields = profile.get("fields") or {}
        sheet_name = str(profile["detection"]["sheet_name"])

        def source_field(logical_name: str) -> str:
            return str(fields.get(logical_name) or "").strip()

        source_note_field = source_field("source_note")
        source_note_cleanup = profile.get("source_note_cleanup") or {}
        if source_note_field and source_note_cleanup:
            cleaned_rows: list[tuple[int, dict[str, Any]]] = []
            for row_no, row in rows:
                cleaned_row = dict(row)
                cleaned_row[source_note_field] = _clean_source_note(
                    row.get(source_note_field), source_note_cleanup
                )
                cleaned_rows.append((row_no, cleaned_row))
            rows = cleaned_rows

        order_no_field = source_field("order_no")
        source_order_no_field = source_field("source_order_no")
        if not order_no_field and not source_order_no_field:
            raise ValueError(
                "配置化Excel来源必须至少配置fields.order_no或fields.source_order_no。"
            )

        grouped: dict[str, list[tuple[int, dict[str, Any]]]] = {}
        for row_no, row in rows:
            raw_order_no = row.get(order_no_field) if order_no_field else None
            raw_source_order_no = (
                row.get(source_order_no_field) if source_order_no_field else None
            )
            invalid_identifier_field = next(
                (
                    field_name
                    for field_name, value in (
                        (order_no_field, raw_order_no),
                        (source_order_no_field, raw_source_order_no),
                    )
                    if field_name and identifier_requires_text(value)
                ),
                None,
            )
            if invalid_identifier_field:
                result.clarifications.append(
                    ClarificationRequest(
                        code="SOURCE_ORDER_ID_TEXT_REQUIRED",
                        scope="order",
                        source_ref=f"{sheet_name}!{row_no}:{row_no}",
                        field=invalid_identifier_field,
                        question=(
                            f'“{invalid_identifier_field}”是超长数值或科学计数法，'
                            "请在来源Excel中将该列改为文本后重新导出。"
                        ),
                        reason="Excel数值最多可靠保存15位有效数字，继续处理可能改变订单标识。",
                        answer_type="file",
                    )
                )
                continue
            order_no = _text(raw_order_no)
            source_order_no = _text(raw_source_order_no)
            try:
                identity = resolve_order_identity(
                    order_no=order_no,
                    source_order_no=source_order_no,
                )
            except OrderIdentityError:
                result.clarifications.append(
                    ClarificationRequest(
                        code="MISSING_SOURCE_ORDER_IDENTITY",
                        scope="order",
                        source_ref=f"{sheet_name}!{row_no}:{row_no}",
                        field="订单编号/来源单号",
                        question="这一行同时缺少订单编号和来源单号，请补充后重新提供文件。",
                        reason="无法判断该商品行属于哪张订单。",
                        answer_type="file",
                    )
                )
                continue
            grouped.setdefault(identity.value, []).append((row_no, row))

        fingerprint = _sha256(source_path) if source_path.exists() else "in-memory"
        extension_fields = {
            str(key): str(value)
            for key, value in (profile.get("order_extension_fields") or {}).items()
            if str(key).strip() and str(value).strip()
        }
        consistency_logical_fields = profile.get("consistency_fields") or [
            "recipient_name",
            "contact",
            "province",
            "city",
            "district",
            "address",
            "source_note",
            "buyer_message",
            "ordered_at",
            "carrier",
            "tracking_no",
        ]

        for order_identity, group_rows in grouped.items():
            row_numbers = tuple(row_no for row_no, _ in group_rows)
            row_ref = f"{sheet_name}!{row_numbers[0]}:{row_numbers[-1]}"
            group_clarifications: list[ClarificationRequest] = []
            first = group_rows[0][1]

            consistency_columns = [
                source_field(name) for name in consistency_logical_fields
            ] + list(extension_fields.values())
            for column in dict.fromkeys(value for value in consistency_columns if value):
                values = {_text(row.get(column)) for _, row in group_rows}
                if len(values) > 1:
                    group_clarifications.append(
                        ClarificationRequest(
                            code="CONFLICTING_ORDER_FIELD",
                            scope="order",
                            source_ref=row_ref,
                            field=column,
                            question=f"同一订单的“{column}”不一致，请确认或修正文件。",
                            reason="订单级字段不能在商品行之间变化。",
                            answer_type="file",
                        )
                    )

            for logical_name, reason in {
                "recipient_name": "缺少收件人",
                "contact": "缺少联系电话或手机",
                "address": "缺少收货地址",
            }.items():
                column = source_field(logical_name)
                if not column or not _text(first.get(column)):
                    group_clarifications.append(
                        ClarificationRequest(
                            code="MISSING_RECIPIENT_FIELD",
                            scope="order",
                            source_ref=row_ref,
                            field=column or logical_name,
                            question=f"订单缺少“{column or logical_name}”，请补充后重新提供文件。",
                            reason=reason,
                            answer_type="file",
                        )
                    )

            items: list[ParsedItem] = []
            for row_no, row in group_rows:
                product_code = _text(row.get(source_field("product_code")))
                product_name = _text(row.get(source_field("product_name")))
                source_spec = _text(row.get(source_field("source_spec")))
                source_display_spec = _text(
                    row.get(source_field("source_display_spec"))
                )
                quantity = _quantity(row.get(source_field("quantity")))
                if not product_code and not product_name:
                    group_clarifications.append(
                        ClarificationRequest(
                            code="MISSING_SOURCE_PRODUCT",
                            scope="field",
                            source_ref=f"{sheet_name}!{row_no}:{row_no}",
                            field="商品标识",
                            question="商品行缺少商品编码和商品名称，请补充后重新提供文件。",
                            reason="无法进行管易商品匹配。",
                            answer_type="file",
                        )
                    )
                if quantity is None:
                    group_clarifications.append(
                        ClarificationRequest(
                            code="INVALID_SOURCE_QUANTITY",
                            scope="field",
                            source_ref=f"{sheet_name}!{row_no}:{row_no}",
                            field=source_field("quantity") or "数量",
                            question="商品数量缺失或不是大于0的数字，请修正后重新提供文件。",
                            reason="管易导入要求每个商品行有有效数量。",
                            answer_type="file",
                        )
                    )
                if product_code or product_name:
                    items.append(
                        ParsedItem(
                            source_product_code=product_code,
                            source_product_name=product_name,
                            source_spec=source_spec,
                            source_display_spec=source_display_spec,
                            quantity=quantity or 0,
                            source_line_no=row_no,
                        )
                    )

            if group_clarifications:
                result.clarifications.extend(group_clarifications)
                continue

            contact_column = source_field("contact")
            contact = _text(first.get(contact_column))
            normalized_contact = normalize_contact(contact)
            if normalized_contact is None:
                result.clarifications.append(
                    ClarificationRequest(
                        code="INVALID_RECIPIENT_CONTACT",
                        scope="order",
                        source_ref=row_ref,
                        field=contact_column or "联系方式",
                        question="收件联系方式格式无法识别，请修正为11位大陆手机号或常用座机号码。",
                        reason="无效联系方式不能进入管易上传Excel。",
                        answer_type="file",
                    )
                )
                continue
            mobile, phone = normalized_contact
            carrier = _text(first.get(source_field("carrier")))
            tracking_no = _text(first.get(source_field("tracking_no")))
            source_extensions = {
                key: _text(first.get(column))
                for key, column in extension_fields.items()
            }
            order_numbers = sorted(
                {
                    _text(row.get(order_no_field))
                    for _, row in group_rows
                    if order_no_field and _text(row.get(order_no_field))
                }
            )
            source_order_numbers = sorted(
                {
                    _text(row.get(source_order_no_field))
                    for _, row in group_rows
                    if source_order_no_field
                    and _text(row.get(source_order_no_field))
                }
            )
            order_no = order_numbers[0] if len(order_numbers) == 1 else None
            source_order_no = (
                source_order_numbers[0]
                if len(source_order_numbers) == 1
                else None
            )
            source_extensions["source_order_numbers"] = source_order_numbers
            source_extensions["selected_order_identity_field"] = (
                "order_no" if order_no else "source_order_no"
            )
            source_order_code_field = source_field("source_order_code")
            if source_order_code_field:
                source_extensions["source_order_codes"] = sorted(
                    {
                        _text(row.get(source_order_code_field))
                        for _, row in group_rows
                        if _text(row.get(source_order_code_field))
                    }
                )
            record_key = str(
                uuid.uuid5(
                    uuid.NAMESPACE_URL,
                    f"{profile['profile_id']}:{fingerprint}:{order_identity}",
                )
            )
            result.orders.append(
                ParsedOrder(
                    record_key=record_key,
                    source=ParsedSource(
                        profile_id=profile["profile_id"],
                        source_label=profile.get("source_label") or profile["profile_id"],
                        order_type=profile.get("order_type", "standard_order"),
                        order_no=order_no,
                        source_order_no=source_order_no,
                        file_name=source_path.name,
                        file_fingerprint=f"sha256:{fingerprint}",
                        sheet_name=sheet_name,
                        row_numbers=row_numbers,
                    ),
                    recipient=ParsedRecipient(
                        name=_text(first.get(source_field("recipient_name"))),
                        mobile=mobile,
                        phone=phone,
                        province=_text(first.get(source_field("province"))),
                        city=_text(first.get(source_field("city"))),
                        district=_text(first.get(source_field("district"))),
                        raw_address=_text(first.get(source_field("address"))),
                    ),
                    items=items,
                    ordered_at=_text(first.get(source_field("ordered_at"))) or None,
                    buyer_message=_text(first.get(source_field("buyer_message"))) or None,
                    source_note=_text(first.get(source_field("source_note"))) or None,
                    shipment_facts=(
                        ParsedShipmentFacts(
                            carrier_name=carrier,
                            tracking_no=tracking_no,
                        )
                        if carrier or tracking_no
                        else None
                    ),
                    source_extensions=source_extensions,
                )
            )

        result.clarifications = deduplicate_clarifications(result.clarifications)
        return result
