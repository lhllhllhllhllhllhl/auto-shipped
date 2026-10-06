from __future__ import annotations

import re
import unicodedata
import warnings
from dataclasses import replace
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from auto_shipped.domain import ClarificationRequest, deduplicate_clarifications

from .configured_excel import ConfiguredExcelParsedAdapter
from .tiantian_parsed import ParseResult


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalized(value: Any) -> str:
    return unicodedata.normalize("NFKC", _text(value)).replace(" ", "")


def _positive_number(value: Any) -> float | None:
    try:
        number = float(_text(value))
    except (TypeError, ValueError):
        return None
    return number if number > 0 else None


def _split_region(raw_address: str) -> tuple[str, str, str, str] | None:
    parts = re.split(r"\s+", raw_address.strip(), maxsplit=3)
    if len(parts) != 4 or not all(parts):
        return None
    province, city, district, detail = parts
    if not province.endswith(("省", "市", "自治区")):
        return None
    if not city.endswith(("市", "州", "地区", "盟")):
        return None
    if not district.endswith(("区", "县", "市", "旗")):
        return None
    return province, city, district, detail


class NddGiftOrderParsedAdapter:
    """Parse NDD gift orders and guard the bundle definition embedded below them."""

    adapter_id = "ndd_gift_order_parsed"

    def parse(self, source_path: str | Path, profile: dict[str, Any]) -> ParseResult:
        path = Path(source_path)
        detection = profile["detection"]
        sheet_name = str(detection["sheet_name"])
        header_row = int(detection.get("header_row", 1))
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Workbook contains no default style.*")
            workbook = load_workbook(path, read_only=False, data_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                return ParseResult(
                    clarifications=[
                        ClarificationRequest(
                            code="MISSING_SOURCE_SHEET",
                            scope="file",
                            source_ref=path.name,
                            question=f"文件中缺少工作表“{sheet_name}”，请确认是否选择了正确的NDD文件。",
                            reason="NDD来源适配器无法找到已登记的订单工作表。",
                            answer_type="file",
                        )
                    ]
                )
            sheet = workbook[sheet_name]
            headers = [_text(cell.value) for cell in sheet[header_row]]
            missing = [
                name for name in detection["required_headers"] if name not in headers
            ]
            if missing:
                return ParseResult(
                    clarifications=[
                        ClarificationRequest(
                            code="MISSING_SOURCE_HEADERS",
                            scope="file",
                            source_ref=path.name,
                            question="NDD Excel缺少必要列，请重新导出正确版本或确认是否出现了新格式。",
                            reason=f"缺少列：{', '.join(missing)}",
                            answer_type="file",
                        )
                    ]
                )

            order_key = str(profile["fields"]["source_order_no"])
            order_key_index = headers.index(order_key)
            rows: list[tuple[int, dict[str, Any]]] = []
            for row_no, cells in enumerate(
                sheet.iter_rows(min_row=header_row + 1, values_only=True),
                start=header_row + 1,
            ):
                if order_key_index >= len(cells) or not _text(cells[order_key_index]):
                    continue
                rows.append(
                    (
                        row_no,
                        {
                            headers[index]: value
                            for index, value in enumerate(cells)
                            if index < len(headers) and headers[index]
                        },
                    )
                )

            clarifications = self._validate_order_rows(rows, profile, sheet_name)
            clarifications.extend(self._validate_bundle_guard(sheet, profile, path.name))
        finally:
            workbook.close()

        if clarifications:
            return ParseResult(clarifications=deduplicate_clarifications(clarifications))

        parsed = ConfiguredExcelParsedAdapter().parse_rows(rows, path, profile)
        if parsed.clarifications:
            return parsed

        for order in parsed.orders:
            components = _split_region(order.recipient.raw_address)
            if components is None:
                parsed.clarifications.append(
                    ClarificationRequest(
                        code="NDD_ADDRESS_COMPONENTS_UNRESOLVED",
                        scope="order",
                        source_ref=(
                            f"{order.source.sheet_name}!"
                            f"{order.source.row_numbers[0]}:{order.source.row_numbers[-1]}"
                        ),
                        field="完整地址",
                        question="这张NDD订单的完整地址无法可靠拆分省、市、区，请修正格式或明确地址。",
                        reason="当前NDD格式应以空格分隔省、市、区和详细地址。",
                        answer_type="file",
                    )
                )
                continue
            province, city, district, detail = components
            order.recipient = replace(
                order.recipient,
                province=province,
                city=city,
                district=district,
                detail=detail,
            )

        parsed.clarifications = deduplicate_clarifications(parsed.clarifications)
        return parsed

    @staticmethod
    def _validate_order_rows(
        rows: list[tuple[int, dict[str, Any]]],
        profile: dict[str, Any],
        sheet_name: str,
    ) -> list[ClarificationRequest]:
        requests: list[ClarificationRequest] = []
        selection = profile.get("order_row_policy") or {}
        product_field = str(profile["fields"].get("product_name") or "")
        expected_product_names = {
            _normalized(value) for value in selection.get("allowed_product_names") or []
        }
        status_field = str(selection.get("status_field") or "")
        allowed_statuses = {
            _normalized(value) for value in selection.get("allowed_statuses") or []
        }
        delivery_field = str(selection.get("delivery_field") or "")
        allowed_delivery = {
            _normalized(value) for value in selection.get("allowed_delivery") or []
        }
        if not rows:
            return [
                ClarificationRequest(
                    code="NDD_NO_ORDER_ROWS",
                    scope="file",
                    question="NDD文件中没有找到带订单编号的订单行，请确认导出范围。",
                    reason="当前文件只有表头或套餐说明，无法生成发货订单。",
                    answer_type="file",
                )
            ]
        for row_no, row in rows:
            if expected_product_names and _normalized(row.get(product_field)) not in expected_product_names:
                requests.append(
                    ClarificationRequest(
                        code="NDD_UNREGISTERED_GIFT_PRODUCT",
                        scope="field",
                        source_ref=f"{sheet_name}!{row_no}:{row_no}",
                        field=product_field,
                        question="NDD出现了尚未登记的礼品名称，请先确认套餐组成和管易商品编码。",
                        reason="不能把新的福利套餐沿用为现有套餐四。",
                        answer_type="text",
                    )
                )
            if allowed_statuses and _normalized(row.get(status_field)) not in allowed_statuses:
                requests.append(
                    ClarificationRequest(
                        code="NDD_ORDER_STATUS_NOT_SHIPPABLE",
                        scope="field",
                        source_ref=f"{sheet_name}!{row_no}:{row_no}",
                        field=status_field,
                        question="NDD订单状态不是已登记的待发货状态，请确认是否仍需导入管易。",
                        reason="系统只自动处理待发货订单。",
                        answer_type="confirmation",
                    )
                )
            if allowed_delivery and _normalized(row.get(delivery_field)) not in allowed_delivery:
                requests.append(
                    ClarificationRequest(
                        code="NDD_DELIVERY_MODE_UNSUPPORTED",
                        scope="field",
                        source_ref=f"{sheet_name}!{row_no}:{row_no}",
                        field=delivery_field,
                        question="NDD订单不是邮寄配送，请确认是否应退出发货导入流程。",
                        reason="自提等配送方式不能按邮寄订单生成管易发货单。",
                        answer_type="confirmation",
                    )
                )
        return requests

    @staticmethod
    def _validate_bundle_guard(sheet, profile: dict[str, Any], source_ref: str) -> list[ClarificationRequest]:
        guard = profile.get("bundle_identity_guard") or {}
        if not guard:
            return []
        all_values: list[tuple[int, int, str]] = []
        for row in sheet.iter_rows():
            for cell in row:
                text = _normalized(cell.value)
                if text:
                    all_values.append((cell.row, cell.column, text))

        expected_title = _normalized(guard.get("bundle_title"))
        title_matches = [item for item in all_values if item[2] == expected_title]
        failures: list[str] = []
        if len(title_matches) != 1:
            failures.append("套餐标题缺失或不唯一")

        for component in guard.get("components") or []:
            expected_name = _normalized(component.get("source_component_name"))
            expected_quantity = float(component.get("quantity_per_bundle") or 0)
            matches = [item for item in all_values if item[2] == expected_name]
            if len(matches) != 1:
                failures.append(f"套餐商品“{component.get('source_component_name')}”缺失或不唯一")
                continue
            row_no, column_no, _ = matches[0]
            actual_quantity = _positive_number(sheet.cell(row_no, column_no + 1).value)
            if actual_quantity != expected_quantity:
                failures.append(f"套餐商品“{component.get('source_component_name')}”数量发生变化")

        if not failures:
            return []
        return [
            ClarificationRequest(
                code="NDD_BUNDLE_COMPOSITION_CHANGED",
                scope="batch",
                source_ref=source_ref,
                field="套餐配套商品信息",
                question="NDD套餐说明与已确认的福利套餐四不一致，请确认新的套餐组成和商品编码。",
                reason="；".join(failures),
                answer_type="text",
            )
        ]
