from __future__ import annotations

import hashlib
import math
import uuid
import warnings
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from auto_shipped.domain import (
    ClarificationRequest,
    ParsedItem,
    ParsedOrder,
    ParsedRecipient,
    ParsedShipmentFacts,
    ParsedSource,
    deduplicate_clarifications,
)

from .value_validation import identifier_requires_text, normalize_contact


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _quantity(value: Any) -> float | None:
    text = _text(value)
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return number if math.isfinite(number) and number > 0 else None


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


@dataclass(slots=True)
class ParseResult:
    orders: list[ParsedOrder] = field(default_factory=list)
    clarifications: list[ClarificationRequest] = field(default_factory=list)

    @property
    def status(self) -> str:
        return "needs_input" if self.clarifications else "parsed"


class TiantianWarehouseParsedAdapter:
    adapter_id = "tiantian_warehouse_parsed"

    def parse(self, source_path: str | Path, profile: dict[str, Any]) -> ParseResult:
        path = Path(source_path)
        detection = profile["detection"]
        sheet_name = detection["sheet_name"]
        header_row = int(detection.get("header_row", 1))
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Workbook contains no default style.*")
            workbook = load_workbook(path, read_only=False, data_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                return ParseResult(
                    clarifications=[
                        ClarificationRequest(
                            code="MISSING_SOURCE_SHEET",
                            scope="file",
                            source_ref=path.name,
                            question=f"文件中缺少工作表“{sheet_name}”，请确认是否选择了正确文件。",
                            reason="来源适配器无法找到已配置工作表。",
                            answer_type="file",
                        )
                    ]
                )
            sheet = workbook[sheet_name]
            headers = [_text(cell.value) for cell in sheet[header_row]]
            missing = [name for name in detection["required_headers"] if name not in headers]
            if missing:
                return ParseResult(
                    clarifications=[
                        ClarificationRequest(
                            code="MISSING_SOURCE_HEADERS",
                            scope="file",
                            source_ref=path.name,
                            question="来源Excel缺少必要列，请重新导出正确版本或确认是否需要新增格式。",
                            reason=f"缺少列：{', '.join(missing)}",
                            answer_type="file",
                        )
                    ]
                )
            rows: list[tuple[int, dict[str, Any]]] = []
            for row_no, cells in enumerate(
                sheet.iter_rows(min_row=header_row + 1, values_only=True),
                start=header_row + 1,
            ):
                row = {headers[index]: value for index, value in enumerate(cells) if index < len(headers)}
                if any(value not in (None, "") for value in row.values()):
                    rows.append((row_no, row))
        finally:
            workbook.close()
        return self.parse_rows(rows, path, profile)

    def parse_rows(
        self,
        rows: Iterable[tuple[int, dict[str, Any]]],
        source_path: Path,
        profile: dict[str, Any],
    ) -> ParseResult:
        result = ParseResult()
        grouped: dict[str, list[tuple[int, dict[str, Any]]]] = {}
        for row_no, row in rows:
            raw_source_order_no = row.get("来源单号")
            if identifier_requires_text(raw_source_order_no):
                result.clarifications.append(
                    ClarificationRequest(
                        code="SOURCE_ORDER_NO_TEXT_REQUIRED",
                        scope="order",
                        source_ref=f"仓库订单!{row_no}:{row_no}",
                        field="来源单号",
                        question="来源单号是超长数值或科学计数法，请在来源Excel中将该列改为文本后重新导出。",
                        reason="Excel数值最多可靠保存15位有效数字，继续处理可能改变订单号。",
                        answer_type="file",
                    )
                )
                continue
            source_order_no = _text(raw_source_order_no)
            if not source_order_no:
                result.clarifications.append(
                    ClarificationRequest(
                        code="MISSING_SOURCE_ORDER_NO",
                        scope="order",
                        source_ref=f"仓库订单!{row_no}:{row_no}",
                        field="来源单号",
                        question="这一行缺少来源单号，请补充后重新提供文件。",
                        reason="无法判断该商品行属于哪张订单。",
                        answer_type="file",
                    )
                )
                continue
            grouped.setdefault(source_order_no, []).append((row_no, row))

        fingerprint = _sha256(source_path) if source_path.exists() else "in-memory"
        order_fields = [
            "收件人",
            "联系方式",
            "省",
            "市",
            "区",
            "省市区详细地址",
            "客户备注",
            "渠道名称",
            "下单时间",
            "物流公司",
            "物流单号",
        ]

        for source_order_no, group_rows in grouped.items():
            row_numbers = tuple(row_no for row_no, _ in group_rows)
            row_ref = f"仓库订单!{row_numbers[0]}:{row_numbers[-1]}"
            group_clarifications: list[ClarificationRequest] = []
            first = group_rows[0][1]

            for field_name in order_fields:
                values = {_text(row.get(field_name)) for _, row in group_rows}
                if len(values) > 1:
                    group_clarifications.append(
                        ClarificationRequest(
                            code="CONFLICTING_ORDER_FIELD",
                            scope="order",
                            source_ref=row_ref,
                            field=field_name,
                            question=f"同一来源单号的“{field_name}”不一致，请确认或修正文件。",
                            reason="订单级字段不能在商品行之间变化。",
                            answer_type="file",
                        )
                    )

            required_order_fields = {
                "收件人": "缺少收件人",
                "联系方式": "缺少联系电话或手机",
                "省市区详细地址": "缺少收货地址",
            }
            for field_name, reason in required_order_fields.items():
                if not _text(first.get(field_name)):
                    group_clarifications.append(
                        ClarificationRequest(
                            code="MISSING_RECIPIENT_FIELD",
                            scope="order",
                            source_ref=row_ref,
                            field=field_name,
                            question=f"订单缺少“{field_name}”，请补充后重新提供文件。",
                            reason=reason,
                            answer_type="file",
                        )
                    )

            items: list[ParsedItem] = []
            for row_no, row in group_rows:
                product_code = _text(row.get("商品编码"))
                product_name = _text(row.get("商品名称"))
                quantity = _quantity(row.get("数量"))
                if not product_code and not product_name:
                    group_clarifications.append(
                        ClarificationRequest(
                            code="MISSING_SOURCE_PRODUCT",
                            scope="field",
                            source_ref=f"仓库订单!{row_no}:{row_no}",
                            field="商品编码/商品名称",
                            question="商品行缺少商品编码和商品名称，请补充后重新提供文件。",
                            reason="无法进行管易商品匹配。",
                            answer_type="file",
                        )
                    )
                if quantity is None:
                    group_clarifications.append(
                        ClarificationRequest(
                            code="INVALID_SOURCE_QUANTITY",
                            scope="field",
                            source_ref=f"仓库订单!{row_no}:{row_no}",
                            field="数量",
                            question="商品数量缺失或不是大于0的数字，请修正后重新提供文件。",
                            reason="管易导入要求每个商品行有有效数量。",
                            answer_type="file",
                        )
                    )
                if product_code or product_name:
                    items.append(
                        ParsedItem(
                            source_product_code=product_code,
                            source_product_name=product_name,
                            quantity=quantity or 0,
                            source_line_no=row_no,
                        )
                    )

            if group_clarifications:
                result.clarifications.extend(group_clarifications)
                continue

            contact = _text(first.get("联系方式"))
            normalized_contact = normalize_contact(contact)
            if normalized_contact is None:
                result.clarifications.append(
                    ClarificationRequest(
                        code="INVALID_RECIPIENT_CONTACT",
                        scope="order",
                        source_ref=row_ref,
                        field="联系方式",
                        question="收件联系方式格式无法识别，请修正为11位大陆手机号或常用座机号码。",
                        reason="无效联系方式不能进入管易上传Excel。",
                        answer_type="file",
                    )
                )
                continue
            mobile, phone = normalized_contact
            record_key = str(
                uuid.uuid5(
                    uuid.NAMESPACE_URL,
                    f"{profile['profile_id']}:{fingerprint}:{source_order_no}",
                )
            )
            carrier = _text(first.get("物流公司"))
            tracking_no = _text(first.get("物流单号"))
            result.orders.append(
                ParsedOrder(
                    record_key=record_key,
                    source=ParsedSource(
                        profile_id=profile["profile_id"],
                        source_label=profile.get("source_label") or profile.get("company", ""),
                        order_type=profile.get("order_type", "warehouse_order"),
                        source_order_no=source_order_no,
                        file_name=source_path.name,
                        file_fingerprint=f"sha256:{fingerprint}",
                        sheet_name="仓库订单",
                        row_numbers=row_numbers,
                    ),
                    recipient=ParsedRecipient(
                        name=_text(first.get("收件人")),
                        mobile=mobile,
                        phone=phone,
                        province=_text(first.get("省")),
                        city=_text(first.get("市")),
                        district=_text(first.get("区")),
                        raw_address=_text(first.get("省市区详细地址")),
                    ),
                    items=items,
                    ordered_at=_text(first.get("下单时间")) or None,
                    source_note=_text(first.get("客户备注")) or None,
                    shipment_facts=(
                        ParsedShipmentFacts(carrier_name=carrier, tracking_no=tracking_no)
                        if carrier or tracking_no
                        else None
                    ),
                    source_extensions={
                        "source_channel": _text(first.get("渠道名称")),
                        "source_order_codes": sorted(
                            {
                                _text(row.get("订单编号"))
                                for _, row in group_rows
                                if _text(row.get("订单编号"))
                            }
                        ),
                    },
                )
            )

        result.clarifications = deduplicate_clarifications(result.clarifications)
        return result
