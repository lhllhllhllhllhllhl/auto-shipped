from __future__ import annotations

import math
import uuid
import warnings
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from auto_shipped.catalog import ProductCatalog, clean_identifier
from auto_shipped.domain import CanonicalOrder, Issue, OrderItem, Recipient, SourceRef
from auto_shipped.validation.order_validator import validate_order


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _quantity(value: Any) -> float:
    if value is None or _text(value) == "":
        return 0.0
    number = float(value)
    if not math.isfinite(number):
        return 0.0
    return number


class TiantianWarehouseAdapter:
    adapter_id = "tiantian_warehouse"

    def ingest(
        self,
        source_path: str | Path,
        profile: dict,
        catalog: ProductCatalog,
    ) -> list[CanonicalOrder]:
        path = Path(source_path)
        sheet_name = profile["detection"]["sheet_name"]
        # Some warehouse exports omit a valid worksheet dimension record.
        # openpyxl's read-only mode then exposes only A1, so compatibility
        # mode is required for these small source files.
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore",
                message="Workbook contains no default style.*",
                category=UserWarning,
            )
            workbook = load_workbook(path, read_only=False, data_only=True)
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"缺少工作表: {sheet_name}")
        sheet = workbook[sheet_name]
        header = [_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        missing = [name for name in profile["detection"]["required_headers"] if name not in header]
        if missing:
            raise ValueError(f"来源文件缺少必要列: {', '.join(missing)}")
        rows = []
        for row_no, cells in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            row = {header[index]: value for index, value in enumerate(cells) if index < len(header)}
            if any(value not in (None, "") for value in row.values()):
                rows.append((row_no, row))
        workbook.close()
        return self.adapt_rows(rows, path.name, profile, catalog)

    def adapt_rows(
        self,
        rows: Iterable[tuple[int, dict[str, Any]]],
        source_file: str,
        profile: dict,
        catalog: ProductCatalog,
    ) -> list[CanonicalOrder]:
        defaults = profile["defaults"]
        orders: list[CanonicalOrder] = []
        grouped: dict[str, list[tuple[int, dict[str, Any]]]] = {}
        for row_no, row in rows:
            source_order_id = _text(row.get("来源单号"))
            group_key = source_order_id or f"__missing_source_order_row_{row_no}"
            grouped.setdefault(group_key, []).append((row_no, row))

        for group_rows in grouped.values():
            first_row_no, first = group_rows[0]
            last_row_no = group_rows[-1][0]
            row_ref = f"仓库订单!A{first_row_no}:P{last_row_no}"
            source_order_id = _text(first.get("来源单号"))
            platform_order_no = source_order_id
            issues: list[Issue] = []

            if not profile["platform_order_number"].get("confirmed", False):
                issues.append(
                    Issue(
                        code="UNCONFIRMED_PLATFORM_NUMBER_STRATEGY",
                        severity="error",
                        message="平台单号暂按来源单号生成，等待业务确认",
                        field="platform_order_no",
                        source_ref=row_ref,
                    )
                )

            header_fields = [
                "收件人",
                "联系方式",
                "省",
                "市",
                "区",
                "省市区详细地址",
                "客户备注",
                "渠道名称",
            ]
            for header_field in header_fields:
                values = {_text(row.get(header_field)) for _, row in group_rows}
                if len(values) > 1:
                    issues.append(
                        Issue(
                            code="CONFLICTING_ORDER_HEADER",
                            severity="error",
                            message=f"同一来源单号的 {header_field} 不一致",
                            field=header_field,
                            source_ref=row_ref,
                        )
                    )

            items: list[OrderItem] = []
            for item_index, (row_no, row) in enumerate(group_rows):
                external_sku = clean_identifier(row.get("商品编码"))
                resolution = catalog.resolve(external_sku)
                product = resolution.product
                if resolution.status != "confirmed":
                    issues.append(
                        Issue(
                            code=(
                                "UNCONFIRMED_SKU_MAPPING"
                                if resolution.status == "unconfirmed"
                                else "SKU_MAPPING_FAILED"
                            ),
                            severity="error",
                            message=f"外部商品码 {external_sku}：{resolution.reason}",
                            field=f"items[{item_index}]",
                            source_ref=f"仓库订单!A{row_no}:P{row_no}",
                        )
                    )
                items.append(
                    OrderItem(
                        external_sku=external_sku,
                        product_code=product.product_code if product else None,
                        spec_code=product.spec_code if product else None,
                        spec_name=product.spec_name if product else None,
                        title=product.product_name if product else _text(row.get("商品名称")),
                        quantity=_quantity(row.get("数量")),
                        unit_price=float(defaults.get("unit_price", 0)),
                        weight_grams=product.weight_grams if product else None,
                        mapping_status=resolution.status,
                    )
                )

            deterministic_key = f"{profile['profile_id']}:{source_order_id or first_row_no}"
            order = CanonicalOrder(
                order_id=str(uuid.uuid5(uuid.NAMESPACE_URL, deterministic_key)),
                source=SourceRef(
                    profile_id=profile["profile_id"],
                    source_file=source_file,
                    source_order_id=source_order_id,
                    row_ref=row_ref,
                    company=profile.get("company", ""),
                    created_at=_text(first.get("下单时间")) or None,
                ),
                platform_order_no=platform_order_no,
                buyer_member=_text(defaults.get("buyer_member")),
                store_name=_text(defaults.get("store_name")),
                warehouse_name=_text(defaults.get("warehouse_name")),
                recipient=Recipient(
                    name=_text(first.get("收件人")),
                    mobile=_text(first.get("联系方式")),
                    province=_text(first.get("省")),
                    city=_text(first.get("市")),
                    district=_text(first.get("区")),
                    detail=_text(first.get("省市区详细地址")),
                ),
                items=items,
                order_status=_text(defaults.get("order_status")),
                order_type=_text(defaults.get("order_type")),
                payment_method=_text(defaults.get("payment_method")),
                payment_amount=float(defaults.get("payment_amount", 0)),
                freight=float(defaults.get("freight", 0)),
                seller_remark=_text(first.get("客户备注")),
                logistics_policy=_text(profile.get("logistics", {}).get("strategy")),
                carrier_name=_text(first.get("物流公司")),
                tracking_no=_text(first.get("物流单号")),
                issues=issues,
                metadata={
                    "source_order_codes": sorted(
                        {_text(row.get("订单编号")) for _, row in group_rows if _text(row.get("订单编号"))}
                    ),
                    "source_channel": _text(first.get("渠道名称")),
                    "source_product_names": [_text(row.get("商品名称")) for _, row in group_rows],
                },
            )
            order.issues.extend(validate_order(order))
            orders.append(order)
        return orders
