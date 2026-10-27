from __future__ import annotations

import hashlib
import json
import re
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from auto_shipped.domain import (
    DecisionBundle,
    DecisionItem,
    MappingReuseBundle,
    MappingReuseCandidate,
    deduplicate_decisions,
    deduplicate_mapping_reuse_candidates,
)
from auto_shipped.platforms.guanyi import (
    GuanyiCustomImportError,
    preflight_custom_import,
    render_custom_import,
)
from auto_shipped.services.convert import PROJECT_ROOT, convert_order_file
from auto_shipped.services.conversion_workflow import (
    ConversionWorkflowError,
    build_batch_workflow,
    complete_conversion_workflow,
    verify_rendered_output,
)


@dataclass(slots=True)
class BatchSourceSummary:
    source_file: str
    source_sha256: str
    status: str
    source_profile_id: str | None = None
    company_id: str | None = None
    route_id: str | None = None
    target_platform: str | None = None
    target_profile_id: str | None = None
    parsed_order_count: int = 0
    output_row_count: int = 0
    schema_warnings: list[dict[str, Any]] = field(default_factory=list)
    clarifications: list[dict[str, Any]] = field(default_factory=list)
    decisions: list[dict[str, Any]] = field(default_factory=list)
    mapping_identity_notices: list[dict[str, Any]] = field(default_factory=list)
    store_assignment_notices: list[dict[str, Any]] = field(default_factory=list)
    package_semantics_notices: list[dict[str, Any]] = field(default_factory=list)
    workflow_execution: dict[str, Any] | None = None
    errors: list[dict[str, Any]] = field(default_factory=list)

    def to_public_dict(self) -> dict[str, Any]:
        return {
            "source_file": self.source_file,
            "source_sha256": self.source_sha256,
            "status": self.status,
            "source_profile_id": self.source_profile_id,
            "company_id": self.company_id,
            "route_id": self.route_id,
            "target_platform": self.target_platform,
            "target_profile_id": self.target_profile_id,
            "parsed_order_count": self.parsed_order_count,
            "output_row_count": self.output_row_count,
            "schema_warnings": self.schema_warnings,
            "clarifications": self.clarifications,
            "decisions": self.decisions,
            "mapping_identity_notices": self.mapping_identity_notices,
            "store_assignment_notices": self.store_assignment_notices,
            "package_semantics_notices": self.package_semantics_notices,
            "workflow_execution": self.workflow_execution,
            "errors": self.errors,
        }


@dataclass(slots=True)
class BatchConversionResult:
    status: str
    source_count: int
    order_count: int = 0
    item_row_count: int = 0
    target_platform: str | None = None
    target_profile_id: str | None = None
    sources: list[BatchSourceSummary] = field(default_factory=list)
    outputs: list[str] = field(default_factory=list)
    upload_manifest: str | None = None
    clarifications: list[dict[str, Any]] = field(default_factory=list)
    decisions: list[DecisionItem] = field(default_factory=list)
    mapping_reuse_candidates: list[MappingReuseCandidate] = field(default_factory=list)
    mapping_evidence: list[dict[str, Any]] = field(default_factory=list)
    mapping_identity_notices: list[dict[str, Any]] = field(default_factory=list)
    store_assignment_notices: list[dict[str, Any]] = field(default_factory=list)
    package_semantics_notices: list[dict[str, Any]] = field(default_factory=list)
    workflow_execution: dict[str, Any] | None = None
    errors: list[dict[str, Any]] = field(default_factory=list)

    def to_public_dict(self) -> dict[str, Any]:
        return {
            "status": self.status,
            "source_count": self.source_count,
            "order_count": self.order_count,
            "item_row_count": self.item_row_count,
            "target_platform": self.target_platform,
            "target_profile_id": self.target_profile_id,
            "sources": [item.to_public_dict() for item in self.sources],
            "outputs": self.outputs,
            "upload_manifest": self.upload_manifest,
            "clarifications": self.clarifications,
            "decision_bundle": DecisionBundle(self.decisions).to_dict(),
            "mapping_reuse_bundle": MappingReuseBundle(
                deduplicate_mapping_reuse_candidates(self.mapping_reuse_candidates)
                if self.status == "ready"
                else []
            ).to_dict(),
            "mapping_evidence": self.mapping_evidence,
            "mapping_identity_notices": self.mapping_identity_notices,
            "store_assignment_notices": self.store_assignment_notices,
            "package_semantics_notices": self.package_semantics_notices,
            "workflow_execution": self.workflow_execution,
            "errors": self.errors,
        }


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _safe_stem(value: str) -> str:
    cleaned = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff._-]+", "_", value).strip("._")
    return cleaned or "批量订单"


def _load_profile(directory: str | Path, profile_id: str) -> dict[str, Any] | None:
    for path in sorted(Path(directory).rglob("*.json")):
        data = json.loads(path.read_text(encoding="utf-8"))
        if data.get("profile_id") == profile_id:
            return data
    return None


def _hint_for_source(
    source: Path,
    source_profile_hints: dict[str, str],
) -> str | None:
    candidates = (str(source), str(source.resolve()), source.name)
    for candidate in candidates:
        if candidate in source_profile_hints:
            return source_profile_hints[candidate]
    return None


def _summarize_conversion(
    source: Path,
    source_sha256: str,
    result: Any,
) -> BatchSourceSummary:
    detection = result.detection
    route = result.route
    warnings: list[dict[str, Any]] = []
    if detection is not None:
        warnings = list(getattr(detection, "schema_warnings", []) or [])
    return BatchSourceSummary(
        source_file=source.name,
        source_sha256=source_sha256,
        status=result.status,
        source_profile_id=(
            detection.source_profile_id if detection is not None else None
        ),
        company_id=detection.company_id if detection is not None else None,
        route_id=route.route_id if route is not None else None,
        target_platform=route.target_platform if route is not None else None,
        target_profile_id=route.target_profile_id if route is not None else None,
        parsed_order_count=result.parsed_order_count,
        output_row_count=result.output_row_count,
        clarifications=[item.to_dict() for item in result.clarifications],
        decisions=[item.to_dict() for item in result.decisions],
        mapping_identity_notices=list(result.mapping_identity_notices),
        store_assignment_notices=list(result.store_assignment_notices),
        package_semantics_notices=list(result.package_semantics_notices),
        workflow_execution=result.workflow_execution,
        errors=list(result.errors),
        schema_warnings=warnings,
    )


def _find_rendered_workbook(outputs: list[str]) -> Path | None:
    candidates = [Path(value) for value in outputs if Path(value).suffix.lower() == ".xlsx"]
    return candidates[0] if len(candidates) == 1 else None


def _read_rendered_lines(
    workbook_path: Path,
    platform_profile: dict[str, Any],
) -> list[dict[str, Any]]:
    workbook = load_workbook(
        workbook_path,
        read_only=False,
        data_only=False,
        keep_links=False,
    )
    try:
        sheet = workbook[str(platform_profile["sheet_name"])]
        header_row = int(platform_profile["header_row"])
        columns = list(platform_profile["columns"])
        headers = [
            sheet.cell(header_row, column_no).value
            for column_no in range(1, len(columns) + 1)
        ]
        if headers != columns:
            raise GuanyiCustomImportError("暂存文件表头与批次目标模板不一致")
        lines: list[dict[str, Any]] = []
        for row_no in range(header_row + 1, sheet.max_row + 1):
            values = [
                sheet.cell(row_no, column_no).value
                for column_no in range(1, len(columns) + 1)
            ]
            if not any(value not in (None, "") for value in values):
                continue
            lines.append(dict(zip(columns, values, strict=True)))
        return lines
    finally:
        workbook.close()


def _with_source_file(
    source_file: str,
    records: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    return [{"source_file": source_file, **record} for record in records]


def convert_order_batch(
    source_paths: list[str | Path] | tuple[str | Path, ...],
    catalog_csv: str | Path,
    output_dir: str | Path,
    *,
    preferred_catalog_paths: tuple[str | Path, ...] = (),
    source_profiles_dir: str | Path = PROJECT_ROOT / "config/source_profiles",
    company_registry_path: str | Path = PROJECT_ROOT
    / "config/companies/company_registry_v1.json",
    routing_path: str | Path = PROJECT_ROOT / "config/routing/order_routes_v1.json",
    platform_profiles_dir: str | Path = PROJECT_ROOT / "config/platform_profiles",
    platform_rules_dir: str | Path = PROJECT_ROOT / "config/platform_rules",
    business_rule_catalog_path: str | Path = PROJECT_ROOT
    / "config/business_rules/shangyu_sop_rule_catalog_v1.json",
    mappings_path: str | Path = PROJECT_ROOT
    / "config/catalog/external_sku_mappings.json",
    guanyi_template_path: str | Path = PROJECT_ROOT
    / "assets/templates/guanyi/自定义订单导入模板.xlsx",
    source_profile_hints: dict[str, str] | None = None,
    company_hints: dict[str, str] | None = None,
    adaptive_plan_hints: dict[str, str] | None = None,
    mapping_overlay_paths: tuple[str | Path, ...] = (),
    mapping_replacement_paths: tuple[str | Path, ...] = (),
    official_mapping_ready: bool = True,
    official_mapping_error: str = "",
    batch_name: str = "批量订单",
) -> BatchConversionResult:
    sources = [Path(value).expanduser().resolve() for value in source_paths]
    result = BatchConversionResult(status="blocked", source_count=len(sources))
    hints = source_profile_hints or {}
    companies = company_hints or {}
    adaptive_plans = adaptive_plan_hints or {}

    if len(sources) < 2:
        result.status = "needs_input"
        result.clarifications.append(
            {
                "source_file": sources[0].name if sources else None,
                "code": "BATCH_REQUIRES_MULTIPLE_SOURCES",
                "scope": "batch",
                "question": "批次合并至少需要两个原始订单文件；单个文件请使用单文件转换入口。",
                "reason": f"source_count={len(sources)}",
                "answer_type": "file",
                "blocking": True,
            }
        )
        return result

    missing = [source.name for source in sources if not source.is_file()]
    if missing:
        result.errors.append(
            {
                "code": "BATCH_SOURCE_NOT_FOUND",
                "message": "以下原始订单文件不存在：" + "、".join(missing),
            }
        )
        return result

    hashes = [_sha256(source) for source in sources]
    duplicate_groups: dict[str, list[str]] = {}
    for source, digest in zip(sources, hashes, strict=True):
        duplicate_groups.setdefault(digest, []).append(source.name)
    duplicates = [names for names in duplicate_groups.values() if len(names) > 1]
    if duplicates:
        result.errors.append(
            {
                "code": "BATCH_DUPLICATE_SOURCE_FILE",
                "message": "批次中包含内容完全相同的重复文件，请去重后重新运行。",
                "duplicate_file_groups": duplicates,
            }
        )
        return result

    staged_workbooks: list[tuple[int, Path]] = []
    with tempfile.TemporaryDirectory(prefix="shangyu-batch-") as temporary:
        staging_root = Path(temporary)
        for index, (source, digest) in enumerate(zip(sources, hashes, strict=True)):
            stage_output = staging_root / f"source-{index + 1:03d}"
            conversion = convert_order_file(
                source_path=source,
                catalog_csv=catalog_csv,
                output_dir=stage_output,
                preferred_catalog_paths=preferred_catalog_paths,
                source_profiles_dir=source_profiles_dir,
                company_registry_path=company_registry_path,
                routing_path=routing_path,
                platform_profiles_dir=platform_profiles_dir,
                platform_rules_dir=platform_rules_dir,
                business_rule_catalog_path=business_rule_catalog_path,
                mappings_path=mappings_path,
                guanyi_template_path=guanyi_template_path,
                source_profile_hint=_hint_for_source(source, hints),
                company_hint=_hint_for_source(source, companies),
                adaptive_plan_path=_hint_for_source(source, adaptive_plans),
                mapping_overlay_paths=mapping_overlay_paths,
                mapping_replacement_paths=mapping_replacement_paths,
                official_mapping_ready=official_mapping_ready,
                official_mapping_error=official_mapping_error,
            )
            summary = _summarize_conversion(source, digest, conversion)
            result.sources.append(summary)
            result.order_count += conversion.parsed_order_count
            result.item_row_count += conversion.output_row_count
            result.decisions.extend(conversion.decisions)
            result.mapping_reuse_candidates.extend(
                conversion.mapping_reuse_candidates
            )
            result.mapping_evidence.extend(conversion.mapping_evidence)
            result.mapping_identity_notices.extend(
                _with_source_file(
                    source.name,
                    conversion.mapping_identity_notices,
                )
            )
            result.store_assignment_notices.extend(
                _with_source_file(
                    source.name,
                    conversion.store_assignment_notices,
                )
            )
            result.package_semantics_notices.extend(
                _with_source_file(
                    source.name,
                    conversion.package_semantics_notices,
                )
            )
            if conversion.status == "needs_input":
                result.clarifications.extend(
                    _with_source_file(source.name, summary.clarifications)
                )
                continue
            if conversion.status not in {"ready", "preview_ready"}:
                result.errors.extend(_with_source_file(source.name, summary.errors))
                continue

            workbook_path = _find_rendered_workbook(conversion.outputs)
            if workbook_path is None:
                result.errors.append(
                    {
                        "source_file": source.name,
                        "code": "BATCH_STAGED_ARTIFACT_MISSING",
                        "message": "单文件转换通过，但暂存Excel缺失。",
                    }
                )
                continue
            if conversion.status == "preview_ready":
                if conversion.upload_manifest is not None or not conversion.decisions:
                    result.errors.append(
                        {
                            "source_file": source.name,
                            "code": "BATCH_PREVIEW_SAFETY_INVALID",
                            "message": "待确认预览错误地携带上传清单，或缺少决策记录。",
                        }
                    )
                    continue
                staged_workbooks.append((index, workbook_path))
                continue
            if not conversion.upload_manifest:
                result.errors.append(
                    {
                        "source_file": source.name,
                        "code": "BATCH_STAGED_ARTIFACT_MISSING",
                        "message": "正式暂存结果缺少上传预检清单。",
                    }
                )
                continue
            manifest_path = Path(conversion.upload_manifest)
            try:
                manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError) as exc:
                result.errors.append(
                    {
                        "source_file": source.name,
                        "code": "BATCH_STAGED_MANIFEST_INVALID",
                        "message": str(exc),
                    }
                )
                continue
            if (
                manifest.get("status") != "ready"
                or manifest.get("artifact", {}).get("sha256") != _sha256(workbook_path)
            ):
                result.errors.append(
                    {
                        "source_file": source.name,
                        "code": "BATCH_STAGED_ARTIFACT_UNVERIFIED",
                        "message": "暂存Excel未通过哈希一致性校验。",
                    }
                )
                continue
            staged_workbooks.append((index, workbook_path))

        if result.errors:
            result.status = "blocked"
            result.outputs = []
            result.upload_manifest = None
            return result
        if result.clarifications:
            result.status = "needs_input"
            result.outputs = []
            result.upload_manifest = None
            return result
        if len(staged_workbooks) != len(sources):
            result.errors.append(
                {
                    "code": "BATCH_STAGING_INCOMPLETE",
                    "message": "并非所有原始文件都形成了可验证的暂存结果。",
                }
            )
            return result
        result.decisions = deduplicate_decisions(result.decisions)
        result.mapping_reuse_candidates = deduplicate_mapping_reuse_candidates(
            result.mapping_reuse_candidates
        )
        unique_mapping_evidence: dict[str, dict[str, Any]] = {}
        for evidence in result.mapping_evidence:
            identity = evidence.get("canonical_identity") or {}
            target = evidence.get("target") or {}
            key = json.dumps(
                {
                    "source_item_text": evidence.get("source_item_text"),
                    "initial_split": evidence.get("initial_split"),
                    "canonical_identity": identity,
                    "mapping_ids": evidence.get("mapping_ids"),
                    "match_method": evidence.get("match_method"),
                    "package_semantics": evidence.get("package_semantics"),
                    "target": target,
                },
                ensure_ascii=False,
                sort_keys=True,
            )
            unique_mapping_evidence[key] = evidence
        result.mapping_evidence = list(unique_mapping_evidence.values())
        unique_mapping_identity_notices: dict[str, dict[str, Any]] = {}
        for notice in result.mapping_identity_notices:
            key = json.dumps(notice, ensure_ascii=False, sort_keys=True)
            unique_mapping_identity_notices[key] = notice
        result.mapping_identity_notices = list(
            unique_mapping_identity_notices.values()
        )
        unique_store_notices: dict[str, dict[str, Any]] = {}
        for notice in result.store_assignment_notices:
            key = json.dumps(notice, ensure_ascii=False, sort_keys=True)
            unique_store_notices[key] = notice
        result.store_assignment_notices = list(unique_store_notices.values())
        unique_package_notices: dict[str, dict[str, Any]] = {}
        for notice in result.package_semantics_notices:
            key = json.dumps(notice, ensure_ascii=False, sort_keys=True)
            unique_package_notices[key] = notice
        result.package_semantics_notices = list(unique_package_notices.values())

        route_keys = {
            (item.target_platform, item.target_profile_id) for item in result.sources
        }
        if len(route_keys) != 1:
            result.status = "needs_input"
            result.clarifications.append(
                {
                    "source_file": None,
                    "code": "BATCH_TARGET_PROFILE_MISMATCH",
                    "scope": "batch",
                    "question": "这些文件对应不同平台或不同上传模板，不能合并为一个Excel。是否按目标平台拆分批次？",
                    "reason": repr(sorted(route_keys)),
                    "answer_type": "confirmation",
                    "blocking": True,
                }
            )
            return result

        target_platform, target_profile_id = next(iter(route_keys))
        result.target_platform = target_platform
        result.target_profile_id = target_profile_id
        if target_platform != "guanyi" or not target_profile_id:
            result.errors.append(
                {
                    "code": "BATCH_TARGET_NOT_IMPLEMENTED",
                    "message": "批次合并目前只支持管易自定义订单导入。",
                }
            )
            return result
        platform_profile = _load_profile(platform_profiles_dir, target_profile_id)
        if platform_profile is None:
            result.errors.append(
                {
                    "code": "BATCH_TARGET_PROFILE_NOT_FOUND",
                    "message": "无法加载批次目标平台配置。",
                }
            )
            return result

        merged_lines: list[dict[str, Any]] = []
        order_sources: dict[str, set[int]] = {}
        order_key = str(platform_profile.get("order_group_key", "平台单号"))
        for source_index, workbook_path in staged_workbooks:
            try:
                lines = _read_rendered_lines(workbook_path, platform_profile)
            except (OSError, KeyError, ValueError, GuanyiCustomImportError) as exc:
                result.errors.append(
                    {
                        "source_file": sources[source_index].name,
                        "code": "BATCH_STAGED_WORKBOOK_INVALID",
                        "message": str(exc),
                    }
                )
                continue
            for line in lines:
                platform_order_no = str(line.get(order_key, "")).strip()
                if platform_order_no:
                    order_sources.setdefault(platform_order_no, set()).add(source_index)
            merged_lines.extend(lines)

        if result.errors:
            return result
        collisions = sorted(
            order_no
            for order_no, source_indexes in order_sources.items()
            if len(source_indexes) > 1
        )
        if collisions:
            result.status = "blocked"
            result.errors.append(
                {
                    "code": "BATCH_PLATFORM_ORDER_NUMBER_COLLISION",
                    "message": "不同原始文件生成了相同的平台单号，整批停止。",
                    "platform_order_numbers": collisions,
                }
            )
            return result

        output_root = Path(output_dir).expanduser().resolve()
        output_stem = _safe_stem(batch_name)
        is_preview = bool(result.decisions)
        output_path = output_root / (
            f"{output_stem}_待确认预览.xlsx"
            if is_preview
            else f"{output_stem}_管易自定义订单导入.xlsx"
        )
        manifest_path = output_root / f"{output_stem}_上传预检.json"
        report_path = output_root / (
            f"{output_stem}_待确认决策.json"
            if is_preview
            else f"{output_stem}_处理报告.json"
        )
        source_workflows = [
            summary.workflow_execution or {}
            for summary in result.sources
        ]
        try:
            batch_workflow = build_batch_workflow(
                source_workflows=source_workflows,
                expected_lines=merged_lines,
            )
            rendered = render_custom_import(
                guanyi_template_path,
                output_path,
                merged_lines,
                platform_profile,
            )
            readback = verify_rendered_output(
                rendered,
                expected_lines=merged_lines,
                platform_profile=platform_profile,
            )
            preflight = preflight_custom_import(rendered, platform_profile)
            if preflight.status != "ready":
                rendered.unlink(missing_ok=True)
                result.errors.extend(issue.to_dict() for issue in preflight.issues)
                return result
            batch_workflow = complete_conversion_workflow(
                batch_workflow,
                readback=readback,
                preflight_checks=preflight.checks,
            )
            if not is_preview:
                preflight.write_manifest(manifest_path)
        except (
            ConversionWorkflowError,
            GuanyiCustomImportError,
            OSError,
            KeyError,
            ValueError,
        ) as exc:
            output_path.unlink(missing_ok=True)
            manifest_path.unlink(missing_ok=True)
            result.errors.append(
                exc.to_dict()
                if isinstance(exc, ConversionWorkflowError)
                else {
                    "code": "BATCH_RENDER_FAILED",
                    "message": str(exc),
                }
            )
            return result

        result.status = "preview_ready" if is_preview else "ready"
        result.order_count = preflight.order_count
        result.item_row_count = preflight.item_row_count
        result.upload_manifest = None if is_preview else str(manifest_path)
        result.workflow_execution = batch_workflow
        result.outputs = [str(rendered)]
        report_path.write_text(
            json.dumps(result.to_public_dict(), ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        result.outputs.append(str(report_path))
        return result
