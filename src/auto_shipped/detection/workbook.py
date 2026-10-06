from __future__ import annotations

import json
import warnings
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from auto_shipped.domain import ClarificationRequest


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


@dataclass(slots=True)
class DetectionResult:
    status: str
    confidence: str
    company_id: str | None = None
    source_profile_id: str | None = None
    source_label: str | None = None
    order_type: str | None = None
    matched_features: list[str] = field(default_factory=list)
    candidate_profile_ids: list[str] = field(default_factory=list)
    schema_warnings: list[dict[str, Any]] = field(default_factory=list)
    clarifications: list[ClarificationRequest] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        data = asdict(self)
        data["clarifications"] = [item.to_dict() for item in self.clarifications]
        return data


def load_source_profiles(directory: str | Path) -> list[dict[str, Any]]:
    root = Path(directory)
    profiles: list[dict[str, Any]] = []
    for path in sorted(root.glob("*.json")):
        profile = json.loads(path.read_text(encoding="utf-8"))
        if profile.get("enabled", True):
            profiles.append(profile)
    return profiles


def _headers_for_profile(
    workbook,
    profile: dict[str, Any],
) -> tuple[list[str], list[str], list[dict[str, Any]]] | None:
    detection = profile.get("detection", {})
    sheet_name = detection.get("sheet_name")
    if not sheet_name or sheet_name not in workbook.sheetnames:
        return None
    sheet = workbook[sheet_name]
    header_row = int(detection.get("header_row", 1))
    headers = [_text(cell.value) for cell in sheet[header_row]]
    header_set = {header for header in headers if header}
    required = [str(value) for value in detection.get("required_headers", [])]
    if not required or not set(required).issubset(header_set):
        return None
    matched = [f"sheet:{sheet_name}", f"header_row:{header_row}"]
    matched.extend(f"header:{header}" for header in required)
    warnings: list[dict[str, Any]] = []
    known_headers = {
        str(value) for value in detection.get("known_headers", required) if str(value)
    }
    unknown_headers = sorted(header_set - known_headers)
    if unknown_headers and detection.get("unknown_header_policy", "warn") == "warn":
        warnings.append(
            {
                "code": "SOURCE_SCHEMA_NEW_HEADERS",
                "message": "来源Excel出现尚未登记的新列，请确认是否包含新的业务信息。",
                "headers": unknown_headers,
            }
        )
    return headers, matched, warnings


def _value_matches(value: str, rule: dict[str, Any]) -> bool:
    expected = tuple(_text(item) for item in rule.get("values", ()))
    if not value or not expected:
        return False
    strategy = str(rule.get("match") or "exact")
    if strategy == "exact":
        return value in expected
    if strategy == "contains":
        return any(token in value for token in expected)
    raise ValueError(f"未知公司证据匹配策略: {strategy}")


def _company_identity_features(
    workbook,
    profile: dict[str, Any],
    headers: list[str],
) -> list[str]:
    detection = profile.get("detection", {})
    identity = detection.get("company_identity") or {}
    rules = identity.get("rules") or []
    if not rules:
        return []

    default_sheet_name = detection.get("sheet_name")
    default_header_row = int(detection.get("header_row", 1))
    features: list[str] = []
    for rule in rules:
        kind = str(rule.get("kind") or "")
        sheet_name = str(rule.get("sheet_name") or default_sheet_name or "")
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        if kind == "fixed_cell":
            cell = str(rule.get("cell") or "")
            if cell and _value_matches(_text(sheet[cell].value), rule):
                features.append(f"company_cell:{sheet_name}!{cell}")
            continue
        if kind == "column_value":
            header = str(rule.get("header") or "")
            header_row = int(rule.get("header_row", default_header_row))
            row_headers = (
                headers
                if header_row == default_header_row and sheet_name == default_sheet_name
                else [_text(cell.value) for cell in sheet[header_row]]
            )
            if header not in row_headers:
                continue
            column_no = row_headers.index(header) + 1
            scan_rows = max(1, int(rule.get("scan_rows", 200)))
            end_row = min(sheet.max_row, header_row + scan_rows)
            for row_no in range(header_row + 1, end_row + 1):
                if _value_matches(_text(sheet.cell(row_no, column_no).value), rule):
                    features.append(f"company_column:{sheet_name}!{header}")
                    break
            continue
        raise ValueError(f"未知公司证据规则类型: {kind or '未配置'}")
    return features


def detect_source(
    source_path: str | Path,
    profiles: list[dict[str, Any]],
    source_profile_hint: str | None = None,
) -> DetectionResult:
    path = Path(source_path)
    if not path.exists():
        return DetectionResult(
            status="needs_input",
            confidence="low",
            clarifications=[
                ClarificationRequest(
                    code="SOURCE_FILE_NOT_FOUND",
                    scope="file",
                    question="请重新提供需要处理的Excel文件。",
                    reason=f"找不到文件：{path.name}",
                    answer_type="file",
                )
            ],
        )
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return DetectionResult(
            status="needs_input",
            confidence="low",
            clarifications=[
                ClarificationRequest(
                    code="UNSUPPORTED_SOURCE_FILE_TYPE",
                    scope="file",
                    question="请提供Excel .xlsx文件，或确认是否需要新增这种文件格式的来源适配器。",
                    reason=f"当前识别器不支持 {path.suffix or '无扩展名'} 文件。",
                    answer_type="file",
                    source_ref=path.name,
                )
            ],
        )

    try:
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Workbook contains no default style.*")
            workbook = load_workbook(path, read_only=False, data_only=True)
    except Exception as exc:
        return DetectionResult(
            status="needs_input",
            confidence="low",
            clarifications=[
                ClarificationRequest(
                    code="UNREADABLE_WORKBOOK",
                    scope="file",
                    source_ref=path.name,
                    question="这个Excel无法读取，请重新导出或提供未损坏的文件。",
                    reason=str(exc),
                    answer_type="file",
                )
            ],
        )

    format_matches: list[
        tuple[dict[str, Any], list[str], list[dict[str, Any]], list[str]]
    ] = []
    try:
        for profile in profiles:
            inspected = _headers_for_profile(workbook, profile)
            if inspected is None:
                continue
            headers, matched, schema_warnings = inspected
            filename_contains = profile.get("detection", {}).get("filename_contains", [])
            if filename_contains:
                matching_tokens = [token for token in filename_contains if token in path.name]
                if not matching_tokens:
                    continue
                matched.extend(f"filename:{token}" for token in matching_tokens)
            identity_features = _company_identity_features(workbook, profile, headers)
            format_matches.append((profile, matched, schema_warnings, identity_features))
    finally:
        workbook.close()

    if source_profile_hint:
        hinted = next(
            (
                item
                for item in format_matches
                if item[0].get("profile_id") == source_profile_hint
            ),
            None,
        )
        if hinted is not None:
            profile, matched, schema_warnings, identity_features = hinted
            return DetectionResult(
                status="matched",
                confidence="high",
                company_id=profile.get("company_id"),
                source_profile_id=profile["profile_id"],
                source_label=profile.get("source_label")
                or profile.get("company")
                or profile["profile_id"],
                order_type=profile.get("order_type") or "unknown",
                matched_features=[
                    *matched,
                    *identity_features,
                    f"explicit_source_profile:{source_profile_hint}",
                ],
                candidate_profile_ids=[profile["profile_id"]],
                schema_warnings=schema_warnings,
            )
        known_ids = {str(profile.get("profile_id")) for profile in profiles}
        code = (
            "SOURCE_PROFILE_HINT_FORMAT_MISMATCH"
            if source_profile_hint in known_ids
            else "SOURCE_PROFILE_HINT_UNKNOWN"
        )
        return DetectionResult(
            status="needs_input",
            confidence="low",
            candidate_profile_ids=[source_profile_hint],
            clarifications=[
                ClarificationRequest(
                    code=code,
                    scope="file",
                    source_ref=path.name,
                    question="用户指定的来源公司与当前Excel格式不一致，请确认公司或重新提供文件。",
                    reason=f"source_profile_hint={source_profile_hint}",
                    answer_type="text",
                )
            ],
        )

    matches = []
    unidentified_format_matches = []
    for item in format_matches:
        profile, matched, schema_warnings, identity_features = item
        identity = profile.get("detection", {}).get("company_identity") or {}
        if identity.get("required", False) and not identity_features:
            unidentified_format_matches.append(item)
            continue
        matches.append(item)

    if len(matches) == 1:
        profile, matched, schema_warnings, identity_features = matches[0]
        return DetectionResult(
            status="matched",
            confidence="high",
            company_id=profile.get("company_id"),
            source_profile_id=profile["profile_id"],
            source_label=profile.get("source_label") or profile.get("company") or profile["profile_id"],
            order_type=profile.get("order_type") or "unknown",
            matched_features=[*matched, *identity_features],
            candidate_profile_ids=[profile["profile_id"]],
            schema_warnings=schema_warnings,
        )

    if len(matches) > 1:
        candidate_ids = [profile["profile_id"] for profile, *_ in matches]
        return DetectionResult(
            status="needs_input",
            confidence="medium",
            candidate_profile_ids=candidate_ids,
            clarifications=[
                ClarificationRequest(
                    code="AMBIGUOUS_SOURCE_PROFILE",
                    scope="file",
                    source_ref=path.name,
                    question="这个文件同时匹配多个来源格式，请确认它属于哪一种。",
                    reason=f"候选来源配置：{', '.join(candidate_ids)}",
                    answer_type="single_choice",
                    choices=tuple(candidate_ids),
                )
            ],
        )

    if unidentified_format_matches:
        candidate_ids = [profile["profile_id"] for profile, *_ in unidentified_format_matches]
        accumulated_schema_warnings = [
            warning
            for _, _, profile_warnings, _ in unidentified_format_matches
            for warning in profile_warnings
        ]
        return DetectionResult(
            status="needs_input",
            confidence="medium",
            candidate_profile_ids=candidate_ids,
            schema_warnings=accumulated_schema_warnings,
            clarifications=[
                ClarificationRequest(
                    code="COMPANY_IDENTITY_UNCONFIRMED",
                    scope="file",
                    source_ref=path.name,
                    question="这个Excel符合已知订单格式，但缺少唯一的公司级证据。请确认它来自哪个公司。",
                    reason=f"仅匹配格式，候选来源：{', '.join(candidate_ids)}",
                    answer_type="single_choice" if len(candidate_ids) > 1 else "confirmation",
                    choices=tuple(candidate_ids),
                    next_action="用户确认后使用显式来源配置重新运行，不把本次回答自动写成永久规则。",
                )
            ],
        )

    return DetectionResult(
        status="needs_input",
        confidence="low",
        clarifications=[
            ClarificationRequest(
                code="UNKNOWN_SOURCE_PROFILE",
                scope="file",
                source_ref=path.name,
                question="暂时无法判断这个Excel属于哪个公司或哪种订单，请告诉我来源公司和文件用途。",
                reason="工作表和表头没有匹配任何已配置来源。",
                answer_type="text",
                next_action="根据用户确认新增或修订来源配置后重新运行。",
            )
        ],
    )
