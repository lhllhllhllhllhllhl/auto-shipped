from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from auto_shipped.detection import detect_source, load_source_profiles
from auto_shipped.routing import load_routing, resolve_route


PROJECT_ROOT = Path(__file__).resolve().parents[1]


TIANTIAN_HEADERS = [
    "订单编号",
    "来源单号",
    "下单时间",
    "收件人",
    "联系方式",
    "省",
    "市",
    "区",
    "省市区详细地址",
    "商品编码",
    "商品名称",
    "数量",
    "客户备注",
    "物流公司",
    "物流单号",
    "渠道名称",
]


class DetectionRoutingTests(unittest.TestCase):
    def test_tiantian_workbook_detects_and_routes_to_guanyi(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "orders.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "仓库订单"
            sheet.append(TIANTIAN_HEADERS)
            sheet.append(
                ["W1", "S1", "", "测试", "13800138000", "上海", "上海市", "宝山区", "测试地址", "P1", "测试商品", 1, "", "", "", "GW_12233787_代发仓-恬田"]
            )
            workbook.save(path)

            profiles = load_source_profiles(PROJECT_ROOT / "config/source_profiles")
            detected = detect_source(path, profiles)
            self.assertEqual(detected.status, "matched")
            self.assertEqual(detected.confidence, "high")
            self.assertEqual(detected.source_profile_id, "tiantian_warehouse_v2")

            routing = load_routing(PROJECT_ROOT / "config/routing/order_routes_v1.json")
            routed = resolve_route(routing, detected.source_profile_id, detected.order_type)
            self.assertEqual(routed.status, "resolved")
            self.assertEqual(routed.decision.target_platform, "guanyi")
            self.assertEqual(
                routed.decision.platform_rules_profile_id,
                "guanyi_tiantian_warehouse_v2",
            )
            self.assertEqual(
                routed.decision.business_rule_scope_id,
                "tiantian_to_guanyi_v1",
            )

    def test_generic_warehouse_headers_do_not_prove_tiantian_company(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "orders.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "仓库订单"
            sheet.append(TIANTIAN_HEADERS)
            sheet.append(
                ["W1", "S1", "", "测试", "13800138000", "上海", "上海市", "宝山区", "测试地址", "P1", "测试商品", 1, "", "", "", "其他渠道"]
            )
            workbook.save(path)

            profiles = load_source_profiles(PROJECT_ROOT / "config/source_profiles")
            detected = detect_source(path, profiles)
            self.assertEqual(detected.status, "needs_input")
            self.assertEqual(detected.confidence, "medium")
            self.assertEqual(detected.clarifications[0].code, "COMPANY_IDENTITY_UNCONFIRMED")

    def test_explicit_company_hint_still_requires_matching_format(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "orders.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "仓库订单"
            sheet.append(TIANTIAN_HEADERS)
            workbook.save(path)

            profiles = load_source_profiles(PROJECT_ROOT / "config/source_profiles")
            detected = detect_source(
                path,
                profiles,
                source_profile_hint="tiantian_warehouse_v2",
            )
            self.assertEqual(detected.status, "matched")
            self.assertIn(
                "explicit_source_profile:tiantian_warehouse_v2",
                detected.matched_features,
            )

    def test_unknown_source_header_is_reported_without_changing_column_order(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "orders.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "仓库订单"
            sheet.append([*reversed(TIANTIAN_HEADERS), "新增业务列"])
            values = {header: "" for header in TIANTIAN_HEADERS}
            values.update(
                {
                    "订单编号": "W1",
                    "来源单号": "S1",
                    "收件人": "测试",
                    "联系方式": "13800138000",
                    "省": "上海",
                    "市": "上海市",
                    "区": "宝山区",
                    "省市区详细地址": "测试地址",
                    "商品编码": "P1",
                    "商品名称": "测试商品",
                    "数量": 1,
                    "渠道名称": "GW_12233787_代发仓-恬田",
                }
            )
            sheet.append([*(values[header] for header in reversed(TIANTIAN_HEADERS)), "待判断"])
            workbook.save(path)

            profiles = load_source_profiles(PROJECT_ROOT / "config/source_profiles")
            detected = detect_source(path, profiles)
            self.assertEqual(detected.status, "matched")
            self.assertEqual(detected.schema_warnings[0]["code"], "SOURCE_SCHEMA_NEW_HEADERS")
            self.assertEqual(detected.schema_warnings[0]["headers"], ["新增业务列"])

    def test_inbound_order_routes_to_agent_clarification(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "inbound.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Sheet1"
            sheet.append([""])
            sheet.append([""])
            sheet.append(["产品名称", "订货数量（根/袋）", "金额"])
            workbook.save(path)

            profiles = load_source_profiles(PROJECT_ROOT / "config/source_profiles")
            detected = detect_source(path, profiles)
            self.assertEqual(detected.source_profile_id, "corn_inbound_order_v1")

            routing = load_routing(PROJECT_ROOT / "config/routing/order_routes_v1.json")
            routed = resolve_route(routing, detected.source_profile_id, detected.order_type)
            self.assertEqual(routed.status, "needs_input")
            self.assertEqual(routed.clarifications[0].code, "ORDER_PURPOSE_UNCLEAR")

    def test_bailihui_backend_orders_route_to_guanyi(self) -> None:
        routing = load_routing(PROJECT_ROOT / "config/routing/order_routes_v1.json")
        routed = resolve_route(routing, "bailihui_dropship_v1", "dropship_order")

        self.assertEqual(routed.status, "resolved")
        self.assertEqual(routed.decision.target_platform, "guanyi")
        self.assertEqual(routed.decision.target_profile_id, "guanyi_order_import_v1")
        self.assertEqual(
            routed.decision.post_fulfillment_profile_id,
            "bailihui_shipping_backfill_v1",
        )


if __name__ == "__main__":
    unittest.main()
