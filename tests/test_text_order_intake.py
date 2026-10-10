from __future__ import annotations

import json
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest.mock import patch
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from auto_shipped.services.text_intake import (
    prepare_text_order_draft,
    validate_text_order_draft,
)
from auto_shipped.services.convert import convert_order_file
from auto_shipped.source_adapters import AdaptiveExcelParsedAdapter, load_adaptive_plan


PROJECT_ROOT = Path(__file__).resolve().parents[1]
REGISTRY = PROJECT_ROOT / "config/companies/company_registry_v1.json"


def _draft(*, company_id: str | None = "tiantian") -> dict:
    return {
        "schema_version": "1.0",
        "draft_type": "agent_text_order_draft",
        "draft_id": "text-20261008-001",
        "source": {
            "kind": "image_ocr",
            "sha256": "a" * 64,
            "business_date": "2026-10-08",
        },
        "company_id": company_id,
        "rule_source_profile_id": None,
        "order_number_policy": {"strategy": "source"},
        "orders": [
            {
                "order_ref": "order-1",
                "source_channel": "天猫",
                "source_order_no": "3316469738341068554",
                "recipient": {
                    "name": "测试收件人",
                    "contact": "13800138000",
                    "province": "上海市",
                    "city": "上海市",
                    "district": "徐汇区",
                    "address": "上海市上海市徐汇区测试路1号",
                },
                "items": [
                    {
                        "source_product_code": "",
                        "source_product_name": "测试商品甲",
                        "source_spec": "黄色",
                        "quantity": 2,
                        "unit": "根",
                    },
                    {
                        "source_product_code": "",
                        "source_product_name": "测试商品乙",
                        "source_spec": "袋装",
                        "quantity": 1,
                        "unit": "袋",
                    },
                ],
                "source_carrier": "",
                "source_note": "",
                "ordered_at": None,
                "source_extensions": {},
                "field_provenance": {
                    "source_channel": "explicit",
                    "source_order_no": "explicit",
                    "recipient_name": "explicit",
                    "contact": "explicit",
                    "address": "explicit",
                    "items": "explicit",
                },
                "uncertainties": [],
            }
        ],
        "confirmation": {"token": "CONFIRM_TEXT_ORDER_DRAFT"},
    }


def _write(path: Path, payload: dict) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


class TextOrderIntakeTests(unittest.TestCase):
    def test_tmall_channel_without_company_context_still_asks_company(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "draft.json"
            payload = _draft(company_id=None)
            payload["orders"][0]["source_channel"] = "天猫"
            _write(path, payload)

            result, materialized = validate_text_order_draft(
                path,
                company_registry_path=REGISTRY,
            )

            self.assertEqual(result.status, "needs_input")
            self.assertIsNone(materialized)
            self.assertIn(
                "TEXT_ORDER_COMPANY_REQUIRED",
                {item.code for item in result.clarifications},
            )

    def test_missing_company_and_inferred_fields_return_masked_questions(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "draft.json"
            payload = _draft(company_id=None)
            payload["orders"][0]["field_provenance"]["contact"] = "inferred"
            payload["orders"][0]["uncertainties"] = [
                {
                    "field": "contact",
                    "reason": "两串数字相连",
                    "question": "请确认联系方式。",
                }
            ]
            payload["confirmation"]["token"] = None
            _write(path, payload)

            result, materialized = validate_text_order_draft(
                path,
                company_registry_path=REGISTRY,
            )

            self.assertEqual(result.status, "needs_input")
            self.assertIsNone(materialized)
            codes = {item.code for item in result.clarifications}
            self.assertIn("TEXT_ORDER_COMPANY_REQUIRED", codes)
            self.assertIn("TEXT_ORDER_CRITICAL_FIELD_UNCONFIRMED", codes)
            self.assertIn("TEXT_ORDER_UNCERTAINTY_REQUIRES_CONFIRMATION", codes)
            public = json.dumps(result.to_public_dict(), ensure_ascii=False)
            self.assertNotIn("测试收件人", public)
            self.assertNotIn("13800138000", public)
            self.assertNotIn("测试路1号", public)
            self.assertIn("138****8000", public)

    def test_confirmed_draft_prepares_excel_and_multi_item_plan(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "draft.json"
            _write(path, _draft())

            result = prepare_text_order_draft(
                path,
                root / "output",
                company_registry_path=REGISTRY,
            )

            self.assertEqual(result.status, "ready")
            self.assertEqual(result.rule_source_profile_id, "tiantian_warehouse_v2")
            self.assertEqual(len(result.outputs), 2)
            workbook_path = Path(result.outputs[0])
            plan_path = Path(result.outputs[1])
            workbook = load_workbook(workbook_path, data_only=False)
            try:
                sheet = workbook["TextOrders"]
                self.assertEqual(sheet.max_row, 3)
                self.assertEqual(sheet.cell(2, 1).data_type, "s")
                self.assertEqual(sheet.cell(2, 1).number_format, "@")
            finally:
                workbook.close()

            plan = load_adaptive_plan(plan_path)
            profile = json.loads(
                (
                    PROJECT_ROOT
                    / "config/source_profiles/tiantian_warehouse_v2.json"
                ).read_text(encoding="utf-8")
            )
            parsed = AdaptiveExcelParsedAdapter().parse(workbook_path, profile, plan)
            self.assertFalse(parsed.clarifications)
            self.assertEqual(len(parsed.orders), 1)
            self.assertEqual(len(parsed.orders[0].items), 2)
            self.assertEqual([item.unit for item in parsed.orders[0].items], ["根", "袋"])
            self.assertEqual(parsed.orders[0].source_extensions["source_channel"], "天猫")

    def test_confirmed_date_sequence_supplies_missing_source_number(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "draft.json"
            payload = _draft(company_id="ndd")
            payload["orders"][0]["source_order_no"] = None
            payload["orders"][0]["field_provenance"]["source_order_no"] = "missing"
            payload["order_number_policy"] = {
                "strategy": "date_sequence",
                "business_date": "2026-10-08",
                "prefix": "TXT",
                "sequence_start": 1,
                "sequence_width": 3,
                "confirmation_token": "CONFIRM_TEXT_ORDER_NUMBER_GENERATION",
            }
            _write(path, payload)

            result = prepare_text_order_draft(
                path,
                root / "output",
                company_registry_path=REGISTRY,
            )

            self.assertEqual(result.status, "ready")
            workbook = load_workbook(result.outputs[0], data_only=False)
            try:
                self.assertEqual(
                    workbook["TextOrders"].cell(2, 1).value,
                    "TXT261008001",
                )
            finally:
                workbook.close()

    def test_confirmed_known_company_text_order_runs_existing_rules(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "draft.json"
            payload = _draft(company_id="ndd")
            payload["orders"][0]["source_channel"] = "NDD文字单"
            payload["orders"][0]["source_order_no"] = "2104856428830863362"
            payload["orders"][0]["recipient"]["address"] = (
                "上海市 上海市 浦东新区 测试路1号"
            )
            payload["orders"][0]["items"] = [
                {
                    "source_product_code": "",
                    "source_product_name": "光明 福利套餐四",
                    "source_spec": "黄糯8根装+白糯8根装",
                    "quantity": 1,
                    "unit": "份",
                }
            ]
            _write(path, payload)

            intake = prepare_text_order_draft(
                path,
                root / "prepared",
                company_registry_path=REGISTRY,
            )
            self.assertEqual(intake.status, "ready")
            conversion = convert_order_file(
                intake.outputs[0],
                PROJECT_ROOT / "assets/catalog/current_product_catalog.csv",
                root / "converted",
                company_hint="ndd",
                adaptive_plan_path=intake.outputs[1],
            )

            self.assertEqual(conversion.status, "ready")
            self.assertEqual(conversion.parsed_order_count, 1)
            self.assertEqual(conversion.output_row_count, 2)
            workbook = load_workbook(conversion.outputs[0], data_only=False)
            try:
                sheet = workbook["Sheet1"]
                self.assertEqual(sheet.cell(2, 6).value, "JTW8E1")
                self.assertEqual(sheet.cell(3, 6).value, "JTBN1E1-EH")
            finally:
                workbook.close()

    def test_self_operated_tmall_text_order_uses_channel_store_and_hidden_mobile(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "draft.json"
            payload = _draft(company_id="self_operated")
            payload["rule_source_profile_id"] = "self_operated_text_v1"
            order = payload["orders"][0]
            order["source_channel"] = "天猫"
            order["source_order_no"] = "3316469738341068554"
            order["recipient"]["contact"] = "13800138000-3940"
            order["items"] = [
                {
                    "source_product_code": "",
                    "source_product_name": "花糯",
                    "source_spec": "",
                    "quantity": 1,
                    "unit": "袋",
                }
            ]
            _write(path, payload)

            intake = prepare_text_order_draft(
                path,
                root / "prepared",
                company_registry_path=REGISTRY,
            )
            self.assertEqual(intake.status, "ready")
            conversion = convert_order_file(
                intake.outputs[0],
                PROJECT_ROOT / "assets/catalog/current_product_catalog.csv",
                root / "converted",
                company_hint="self_operated",
                adaptive_plan_path=intake.outputs[1],
            )

            self.assertEqual(conversion.status, "ready")
            workbook = load_workbook(conversion.outputs[0], data_only=False)
            try:
                sheet = workbook["Sheet1"]
                self.assertEqual(sheet.cell(2, 1).value, "MANYANGCHI旗舰店（天猫）")
                self.assertEqual(sheet.cell(2, 2).value, "3316469738341068554A")
                self.assertEqual(sheet.cell(2, 6).value, "JTTZE1")
                self.assertEqual(sheet.cell(2, 7).value, "6974768564835")
                self.assertEqual(sheet.cell(2, 16).value, "13800138000-3940")
                self.assertEqual(sheet.cell(2, 17).value, "13800138000-3940")
            finally:
                workbook.close()

    def test_self_operated_jd_bag_uses_family_pack_and_literal_jd_zhongtong(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "draft.json"
            payload = _draft(company_id="self_operated")
            payload["rule_source_profile_id"] = "self_operated_text_v1"
            order = payload["orders"][0]
            order["source_channel"] = "京东"
            order["source_carrier"] = "京东中通"
            order["source_order_no"] = "3644496001658889"
            order["recipient"]["name"] = "测试收件人[0534]"
            order["recipient"]["contact"] = "13800138000"
            order["recipient"]["address"] = "上海市上海市徐汇区测试路1号[0534]"
            order["items"] = [
                {
                    "source_product_code": "",
                    "source_product_name": "黄糯",
                    "source_spec": "",
                    "quantity": 1,
                    "unit": "袋",
                }
            ]
            _write(path, payload)

            intake = prepare_text_order_draft(
                path,
                root / "prepared",
                company_registry_path=REGISTRY,
            )
            self.assertEqual(intake.status, "ready")
            conversion = convert_order_file(
                intake.outputs[0],
                PROJECT_ROOT / "assets/catalog/current_product_catalog.csv",
                root / "converted",
                company_hint="self_operated",
                adaptive_plan_path=intake.outputs[1],
            )

            self.assertEqual(conversion.status, "ready")
            workbook = load_workbook(conversion.outputs[0], data_only=False)
            try:
                sheet = workbook["Sheet1"]
                self.assertEqual(sheet.cell(2, 1).value, "MAN YUAN QI旗舰店（京东）")
                self.assertEqual(sheet.cell(2, 2).value, "3644496001658889A")
                self.assertEqual(sheet.cell(2, 6).value, "JTW8E1")
                self.assertEqual(sheet.cell(2, 7).value, "6974768564811")
                self.assertEqual(sheet.cell(2, 16).value, "13800138000-0534")
                self.assertEqual(sheet.cell(2, 17).value, "13800138000-0534")
                self.assertEqual(sheet.cell(2, 27).value, "京东中通")
            finally:
                workbook.close()

    def test_taojuzi_text_profiles_use_minute_numbers_store_and_confirmed_aliases(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            overlay = root / "overlay.json"
            overlay.write_text(
                json.dumps(
                    {
                        "schema_version": "2.0",
                        "mappings": [
                            {
                                "source_profile_id": "taojuzi_taobao_text_v1",
                                "identifier_type": "product_name",
                                "source_value": "搓澡巾",
                                "source_spec": "蓝色",
                                "product_code": "KKAW0030",
                                "spec_code": "4548404200030",
                                "confirmed": True,
                            },
                            {
                                "source_profile_id": "taojuzi_kqyd_text_v1",
                                "identifier_type": "product_name",
                                "source_value": "Revo海绵",
                                "source_spec": "绿色",
                                "product_code": "KKRevo1894",
                                "spec_code": "4548404101894",
                                "confirmed": True,
                            },
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            cases = (
                (
                    "taojuzi_taobao_text_v1",
                    "淘宝",
                    "TB20261009154201A",
                    "搓澡巾",
                    "蓝色",
                    "KKAW0030",
                    "4548404200030",
                    "",
                ),
                (
                    "taojuzi_kqyd_text_v1",
                    "KQYD",
                    "KQYD20261009154201A",
                    "Revo海绵",
                    "绿色",
                    "KKRevo1894",
                    "4548404101894",
                    "送货上门",
                ),
            )
            for index, (
                profile_id,
                channel,
                expected_number,
                source_name,
                source_spec,
                product_code,
                spec_code,
                note,
            ) in enumerate(cases, start=1):
                path = root / f"draft-{index}.json"
                payload = _draft(company_id="self_operated")
                payload["draft_id"] = f"taojuzi-{index}"
                payload["source"]["sha256"] = str(index) * 64
                payload["rule_source_profile_id"] = profile_id
                payload.pop("order_number_policy")
                order = payload["orders"][0]
                order["source_channel"] = channel
                order["source_order_no"] = None
                order["field_provenance"]["source_order_no"] = "missing"
                order["recipient"]["contact"] = "13800138000-8337"
                order["recipient"]["address"] = "测试路1号15507405579"
                order["items"] = [
                    {
                        "source_product_code": "",
                        "source_product_name": source_name,
                        "source_spec": source_spec,
                        "quantity": 1,
                        "unit": "个",
                    }
                ]
                order["source_note"] = note
                _write(path, payload)

                with patch(
                    "auto_shipped.services.text_intake._now_in_timezone",
                    return_value=datetime(
                        2026, 10, 9, 15, 42, tzinfo=ZoneInfo("Asia/Shanghai")
                    ),
                ):
                    intake = prepare_text_order_draft(
                        path,
                        root / f"prepared-{index}",
                        company_registry_path=REGISTRY,
                        order_number_state_dir=root / "state",
                    )
                self.assertEqual(intake.status, "ready")
                conversion = convert_order_file(
                    intake.outputs[0],
                    PROJECT_ROOT / "assets/catalog/current_product_catalog.csv",
                    root / f"converted-{index}",
                    company_hint="self_operated",
                    adaptive_plan_path=intake.outputs[1],
                    mapping_overlay_paths=(overlay,),
                )
                self.assertEqual(conversion.status, "ready", conversion.to_public_dict())
                workbook = load_workbook(conversion.outputs[0], data_only=False)
                try:
                    sheet = workbook["Sheet1"]
                    self.assertEqual(sheet.cell(2, 1).value, "淘橘子")
                    self.assertEqual(sheet.cell(2, 2).value, expected_number)
                    self.assertEqual(sheet.cell(2, 6).value, product_code)
                    self.assertEqual(sheet.cell(2, 7).value, spec_code)
                    self.assertEqual(sheet.cell(2, 16).value, "13800138000-8337")
                    self.assertIn("15507405579", sheet.cell(2, 18).value)
                    if note:
                        self.assertTrue(sheet.cell(2, 18).value.endswith(f"（{note}）"))
                    self.assertEqual(sheet.cell(2, 27).value, "韵达快递")
                    self.assertEqual(sheet.cell(2, 28).value or "", "白标商品")
                finally:
                    workbook.close()

    def test_rongzhida_uses_batch_minute_elastic_sequence_and_corn_store(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "draft.json"
            payload = _draft(company_id="RZD")
            payload.pop("order_number_policy")
            order = payload["orders"][0]
            order["source_channel"] = "RZD"
            order["source_order_no"] = None
            order["field_provenance"]["source_order_no"] = "missing"
            order["items"] = [
                {
                    "source_product_code": "",
                    "source_product_name": "黄糯",
                    "source_spec": "",
                    "quantity": 2,
                    "unit": "根",
                }
            ]
            _write(path, payload)

            with patch(
                "auto_shipped.services.text_intake._now_in_timezone",
                return_value=datetime(
                    2026, 10, 8, 14, 23, tzinfo=ZoneInfo("Asia/Shanghai")
                ),
            ):
                intake = prepare_text_order_draft(
                    path,
                    root / "prepared",
                    company_registry_path=REGISTRY,
                    order_number_state_dir=root / "state",
                )

            self.assertEqual(intake.status, "ready")
            conversion = convert_order_file(
                intake.outputs[0],
                PROJECT_ROOT / "assets/catalog/current_product_catalog.csv",
                root / "converted",
                company_hint="rongzhida",
                adaptive_plan_path=intake.outputs[1],
            )
            self.assertEqual(conversion.status, "ready")
            workbook = load_workbook(conversion.outputs[0], data_only=False)
            try:
                sheet = workbook["Sheet1"]
                self.assertEqual(sheet.cell(2, 1).value, "光明满元气")
                self.assertEqual(sheet.cell(2, 2).value, "RZD20261008142301A")
                self.assertEqual(sheet.cell(2, 6).value, "GMHN-1")
                self.assertEqual(sheet.cell(2, 7).value, "6974768564729")
            finally:
                workbook.close()

            self.assertEqual(
                intake.order_number_allocation["batch_minute"],
                "202610081423",
            )
            self.assertEqual(intake.order_number_allocation["sequence_start"], 1)
            self.assertEqual(intake.order_number_allocation["sequence_end"], 1)
            self.assertFalse(intake.order_number_allocation["reused"])

    def test_rongzhida_sequence_minimum_width_expands_after_99(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "draft.json"
            payload = _draft(company_id="rongzhida")
            payload.pop("order_number_policy")
            template = payload["orders"][0]
            template["source_channel"] = "RZD"
            template["source_order_no"] = None
            template["field_provenance"]["source_order_no"] = "missing"
            template["items"] = [template["items"][0]]
            orders = []
            for index in range(100):
                order = json.loads(json.dumps(template, ensure_ascii=False))
                order["order_ref"] = f"order-{index + 1}"
                orders.append(order)
            payload["orders"] = orders
            _write(path, payload)

            with patch(
                "auto_shipped.services.text_intake._now_in_timezone",
                return_value=datetime(
                    2026, 10, 8, 14, 23, tzinfo=ZoneInfo("Asia/Shanghai")
                ),
            ):
                result = prepare_text_order_draft(
                    path,
                    root / "prepared",
                    company_registry_path=REGISTRY,
                    order_number_state_dir=root / "state",
                )

            self.assertEqual(result.status, "ready")
            workbook = load_workbook(result.outputs[0], data_only=False)
            try:
                sheet = workbook["TextOrders"]
                self.assertEqual(sheet.cell(2, 1).value, "RZD20261008142301")
                self.assertEqual(sheet.cell(100, 1).value, "RZD20261008142399")
                self.assertEqual(sheet.cell(101, 1).value, "RZD202610081423100")
            finally:
                workbook.close()

    def test_rongzhida_same_minute_continues_and_same_source_reuses(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)

            def make_payload(source_sha: str, draft_id: str) -> dict:
                payload = _draft(company_id="rongzhida")
                payload.pop("order_number_policy")
                payload["draft_id"] = draft_id
                payload["source"]["sha256"] = source_sha
                order = payload["orders"][0]
                order["source_channel"] = "RZD"
                order["source_order_no"] = None
                order["field_provenance"]["source_order_no"] = "missing"
                order["items"] = [order["items"][0]]
                return payload

            first_path = root / "first.json"
            second_path = root / "second.json"
            _write(first_path, make_payload("a" * 64, "first"))
            _write(second_path, make_payload("b" * 64, "second"))
            fixed_minute = datetime(
                2026, 10, 8, 14, 23, tzinfo=ZoneInfo("Asia/Shanghai")
            )
            with patch(
                "auto_shipped.services.text_intake._now_in_timezone",
                return_value=fixed_minute,
            ):
                first = prepare_text_order_draft(
                    first_path,
                    root / "first-output",
                    company_registry_path=REGISTRY,
                    order_number_state_dir=root / "state",
                )
                second = prepare_text_order_draft(
                    second_path,
                    root / "second-output",
                    company_registry_path=REGISTRY,
                    order_number_state_dir=root / "state",
                )

            with patch(
                "auto_shipped.services.text_intake._now_in_timezone",
                return_value=datetime(
                    2026, 10, 8, 14, 30, tzinfo=ZoneInfo("Asia/Shanghai")
                ),
            ):
                repeated = prepare_text_order_draft(
                    first_path,
                    root / "repeated-output",
                    company_registry_path=REGISTRY,
                    order_number_state_dir=root / "state",
                )

            self.assertEqual(first.order_number_allocation["sequence_start"], 1)
            self.assertEqual(second.order_number_allocation["sequence_start"], 2)
            self.assertEqual(repeated.order_number_allocation["sequence_start"], 1)
            self.assertEqual(
                repeated.order_number_allocation["batch_minute"],
                "202610081423",
            )
            self.assertTrue(repeated.order_number_allocation["reused"])

            for result, expected in (
                (first, "RZD20261008142301"),
                (second, "RZD20261008142302"),
                (repeated, "RZD20261008142301"),
            ):
                workbook = load_workbook(result.outputs[0], data_only=False)
                try:
                    self.assertEqual(workbook["TextOrders"].cell(2, 1).value, expected)
                finally:
                    workbook.close()


if __name__ == "__main__":
    unittest.main()
