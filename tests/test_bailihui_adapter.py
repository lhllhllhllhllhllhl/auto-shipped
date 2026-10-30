from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from auto_shipped.detection import detect_source, load_source_profiles
from auto_shipped.routing import load_routing, resolve_route
from auto_shipped.services.convert import convert_order_file
from auto_shipped.source_adapters import BailihuiDropshipParsedAdapter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
BAILIHUI_BOILERPLATE_NOTE = (
    "【请在24小时内使用快递发货，发货时请保证货品清洁，"
    "邮寄时请用邮购箱发货。】"
)
PROFILE = json.loads(
    (PROJECT_ROOT / "config/source_profiles/bailihui_dropship_v1.json").read_text(
        encoding="utf-8"
    )
)


HEADERS = [
    "采购单号",
    "采购类型",
    "采购状态",
    "产品名称",
    "SKU码",
    "SKU",
    "69码",
    "数量",
    "单价",
    "商品金额",
    "标准运费",
    "收货人",
    "收货电话",
    "收货地址",
    "下单日期",
    "备注",
    "物流公司",
    "物流单号",
    "实际运费",
]


def source_row(
    *,
    order_no: str = "AP-TEST-001",
    purchase_type: str = "代发采购",
    status: str = "待发货",
    product_code: str = "SWSH8072",
    spec_code: str = "4901792038072",
    product_name: str = "SHOWA尚和 丝滑中厚型手套/珍珠粉(S) SWSH8072",
    note: str = "",
) -> list[object]:
    return [
        order_no,
        purchase_type,
        status,
        product_name,
        "2206566",
        product_code,
        spec_code,
        1,
        1,
        1,
        0,
        "测试收件人",
        "13800138000",
        "上海市测试区测试路1号",
        "2026-10-30 09:30:00",
        note,
        "",
        "",
        0,
    ]


def write_source(path: Path, rows: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "代发订单信息"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    workbook.save(path)


class BailihuiAdapterTests(unittest.TestCase):
    def test_fixed_export_detects_and_routes_to_bailihui(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "代发订单数据.xlsx"
            write_source(path, [source_row()])

            detection = detect_source(
                path,
                load_source_profiles(PROJECT_ROOT / "config/source_profiles"),
            )
            self.assertEqual(detection.status, "matched")
            self.assertEqual(detection.confidence, "high")
            self.assertEqual(detection.company_id, "bailihui")
            self.assertEqual(detection.source_profile_id, "bailihui_dropship_v1")

            route = resolve_route(
                load_routing(PROJECT_ROOT / "config/routing/order_routes_v1.json"),
                detection.source_profile_id,
                detection.order_type,
            )
            self.assertEqual(route.status, "resolved")
            self.assertEqual(
                route.decision.platform_rules_profile_id,
                "guanyi_bailihui_dropship_v1",
            )

    def test_parser_maps_sku_and_69_code_without_guessing(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.xlsx"
            write_source(path, [source_row()])

            parsed = BailihuiDropshipParsedAdapter().parse(path, PROFILE)
            self.assertEqual(parsed.status, "parsed")
            self.assertEqual(len(parsed.orders), 1)
            order = parsed.orders[0]
            self.assertEqual(order.source.order_no, "AP-TEST-001")
            self.assertEqual(order.items[0].source_product_code, "SWSH8072")
            self.assertEqual(order.items[0].source_spec, "4901792038072")
            self.assertEqual(order.items[0].quantity, 1)

    def test_non_dropship_or_non_pending_rows_are_blocked(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.xlsx"
            write_source(
                path,
                [
                    source_row(purchase_type="普通采购"),
                    source_row(order_no="AP-TEST-002", status="已发货"),
                ],
            )

            parsed = BailihuiDropshipParsedAdapter().parse(path, PROFILE)
            codes = {item.code for item in parsed.clarifications}
            self.assertIn("BAILIHUI_PURCHASE_TYPE_NOT_DROPSHIP", codes)
            self.assertIn("BAILIHUI_ORDER_STATUS_NOT_PENDING_SHIPMENT", codes)
            self.assertEqual(parsed.orders, [])

    def test_boilerplate_note_is_removed_and_other_note_is_preserved(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.xlsx"
            write_source(
                path,
                [
                    source_row(note=BAILIHUI_BOILERPLATE_NOTE),
                    source_row(
                        order_no="AP-TEST-002",
                        note=f"{BAILIHUI_BOILERPLATE_NOTE}；请周末发货。",
                    ),
                ],
            )

            parsed = BailihuiDropshipParsedAdapter().parse(path, PROFILE)
            self.assertEqual(parsed.status, "parsed")
            self.assertEqual(parsed.orders[0].source_note, None)
            self.assertEqual(parsed.orders[1].source_note, "请周末发货。")

    def test_project_pipeline_generates_guanyi_excel_and_manifest(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            output = root / "output"
            write_source(
                source,
                [
                    source_row(note=BAILIHUI_BOILERPLATE_NOTE),
                    source_row(
                        order_no="AP-TEST-002",
                        product_code="SWJH7662",
                        spec_code="4901792037662",
                        product_name=(
                            "SHOWA尚和 植绒加厚款手套/宝石粉(S) SWJH7662"
                        ),
                        note=f"请周末发货；{BAILIHUI_BOILERPLATE_NOTE}",
                    ),
                ],
            )

            result = convert_order_file(
                source,
                PROJECT_ROOT / "assets/catalog/current_product_catalog.csv",
                output,
                preferred_catalog_paths=(
                    PROJECT_ROOT / "assets/catalog/preferred_product_catalog.csv",
                ),
            )

            self.assertEqual(result.status, "ready")
            self.assertEqual(result.parsed_order_count, 2)
            self.assertEqual(result.output_row_count, 2)
            self.assertIsNotNone(result.upload_manifest)

            workbook = load_workbook(result.outputs[0], data_only=True, read_only=True)
            sheet = workbook["Sheet1"]
            headers = [cell.value for cell in sheet[1]]
            rendered = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
            workbook.close()

            self.assertEqual([row["店铺"] for row in rendered], ["百礼汇", "百礼汇"])
            self.assertEqual(
                [row["平台单号"] for row in rendered],
                ["AP-TEST-001", "AP-TEST-002"],
            )
            self.assertEqual(
                [(row["商品代码"], row["规格代码"]) for row in rendered],
                [
                    ("SWSH8072", "4901792038072"),
                    ("SWJH7662", "4901792037662"),
                ],
            )
            self.assertEqual([row["买家会员"] for row in rendered], ["张", "张"])
            self.assertEqual(
                [row["物流公司"] for row in rendered],
                ["韵达快递", "韵达快递"],
            )
            self.assertTrue(all("白标商品" in row["卖家备注"] for row in rendered))
            self.assertTrue(
                all(BAILIHUI_BOILERPLATE_NOTE not in row["卖家备注"] for row in rendered)
            )
            self.assertNotIn("请周末发货", rendered[0]["卖家备注"])
            self.assertIn("请周末发货", rendered[1]["卖家备注"])
            self.assertTrue(all(row["订单类型"] == "销售订单" for row in rendered))


if __name__ == "__main__":
    unittest.main()
