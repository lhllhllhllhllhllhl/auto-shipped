from __future__ import annotations

import hashlib
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from auto_shipped.integrations.feishu.daily_sequence_ledger import (
    DailySequenceLedgerRecord,
)
from auto_shipped.services.convert import convert_order_file
from auto_shipped.source_adapters import LabeledTableOrderParsedAdapter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
PROFILE = json.loads(
    (PROJECT_ROOT / "config/source_profiles/corn_inbound_order_v1.json").read_text(
        encoding="utf-8"
    )
)


class MemoryDailySequenceGateway:
    def __init__(self) -> None:
        self.records: list[DailySequenceLedgerRecord] = []

    def current_user_id(self) -> str:
        return "ou-test"

    def list_records(self) -> list[DailySequenceLedgerRecord]:
        return list(self.records)

    def append_record(self, record: DailySequenceLedgerRecord) -> None:
        self.records.append(record)


def build_fuze_table(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["B2"] = "上海孚泽经贸有限公司采购订单"
    headers = [
        "序号",
        "产品名称",
        "克重",
        "箱规",
        "结算价（个）",
        "订货数量（根/袋）",
        "金额",
    ]
    for column, value in enumerate(headers, start=2):
        sheet.cell(3, column).value = value
    values = [1, "黄糯玉米 单棒装", "230g/棒", "", 3.9, 40, 156]
    for column, value in enumerate(values, start=2):
        sheet.cell(4, column).value = value
    sheet["B7"] = "合计"
    sheet["B8"] = "收货地址"
    sheet["C8"] = "上海市杨浦区测试路1号"
    sheet["D8"] = "联系人"
    sheet["E8"] = "测试人 13800138000"
    workbook.save(path)


class LabeledTableOrderTests(unittest.TestCase):
    def test_adapter_reads_item_table_and_footer_recipient(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "玉米入仓订单.xlsx"
            build_fuze_table(source)

            parsed = LabeledTableOrderParsedAdapter().parse(source, PROFILE)

            self.assertEqual(parsed.status, "parsed")
            self.assertEqual(len(parsed.orders), 1)
            order = parsed.orders[0]
            self.assertIsNone(order.source.order_no)
            self.assertIsNone(order.source.source_order_no)
            self.assertEqual(order.recipient.name, "测试人")
            self.assertEqual(order.recipient.mobile, "13800138000")
            self.assertEqual(order.recipient.raw_address, "上海市杨浦区测试路1号")
            self.assertEqual(order.items[0].source_product_name, "黄糯玉米 单棒装")
            self.assertEqual(order.items[0].source_spec, "230g/棒")
            self.assertEqual(order.items[0].quantity, 40)

    def test_full_conversion_uses_company_daily_sequence_and_shared_corn_store(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "玉米入仓订单.xlsx"
            output = root / "output"
            overlay = root / "mapping.json"
            build_fuze_table(source)
            digest = hashlib.sha256(source.read_bytes()).hexdigest()
            overlay.write_text(
                json.dumps(
                    {
                        "schema_version": "2.0",
                        "mappings": [
                            {
                                "source_profile_id": "corn_inbound_order_v1",
                                "identifier_type": "product_name",
                                "source_value": "黄糯玉米 单棒装",
                                "source_spec": "230g/棒",
                                "product_code": "GMHN-1",
                                "spec_code": "6974768564729",
                                "confirmed": True,
                                "reuse_scope": "current_batch",
                                "mapping_authority": "user_batch_confirmation",
                                "source_fingerprint": f"sha256:{digest}",
                            }
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            ledger = MemoryDailySequenceGateway()

            result = convert_order_file(
                source,
                PROJECT_ROOT / "assets/catalog/current_product_catalog.csv",
                output,
                preferred_catalog_paths=(
                    PROJECT_ROOT / "assets/catalog/preferred_product_catalog.csv",
                ),
                mapping_overlay_paths=(overlay,),
                order_number_gateway=ledger,
                order_number_state_dir=root / "state",
            )

            self.assertEqual(result.status, "ready", result.to_public_dict())
            self.assertEqual(len(ledger.records), 1)
            workbook_path = next(Path(value) for value in result.outputs if value.endswith(".xlsx"))
            workbook = load_workbook(workbook_path, data_only=False)
            try:
                sheet = workbook["Sheet1"]
                self.assertEqual(sheet.cell(2, 1).value, "光明满元气")
                self.assertRegex(str(sheet.cell(2, 2).value), r"^FZ\d{8}\d{3}$")
                self.assertEqual(sheet.cell(2, 6).value, "GMHN-1")
                self.assertEqual(sheet.cell(2, 10).value, 40)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
