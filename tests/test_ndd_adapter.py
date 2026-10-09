from __future__ import annotations

import json
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from auto_shipped.catalog import ProductCatalog, ProductRecord
from auto_shipped.detection import detect_source, load_source_profiles
from auto_shipped.platforms.guanyi import (
    build_custom_import_lines,
    resolve_guanyi_policy_modules,
)
from auto_shipped.services.convert import convert_order_file
from auto_shipped.source_adapters import NddGiftOrderParsedAdapter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
PROFILE = json.loads(
    (PROJECT_ROOT / "config/source_profiles/ndd_order_v1.json").read_text(
        encoding="utf-8"
    )
)
RULES = resolve_guanyi_policy_modules(json.loads(
    (PROJECT_ROOT / "config/platform_rules/guanyi/ndd_order_v1.json").read_text(
        encoding="utf-8"
    )
), PROJECT_ROOT / "config/platform_rules")
GUANYI_PROFILE = json.loads(
    (PROJECT_ROOT / "config/platform_profiles/guanyi_order_import_v1.json").read_text(
        encoding="utf-8"
    )
)
STANDARD_TEMPLATE = (
    PROJECT_ROOT / "assets/templates/ndd/NDD完整字段标准模板_v1.2.xlsx"
)


HEADERS = [
    "商品小计",
    "订单编号",
    "礼品名称",
    "礼品描述",
    "礼品价格",
    "数量",
    "订单状态",
    "配送方式",
    "收件人姓名",
    "收件人手机",
    "完整地址",
    "快递公司",
    "快递单号",
    "兑换码",
    "订单生成时间",
    "自提点名称",
    "自提点地址",
    "提货人姓名",
    "提货人手机",
    "礼品规格",
    "",
    "套餐配套商品信息",
    "送货时间",
]


def build_sample(path: Path, *, white_quantity: int = 1) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(HEADERS)
    sheet.append(
        [
            500,
            "2104856428830863362",
            "光明 福利套餐四",
            "",
            "500.00",
            1,
            "待发货",
            "邮寄",
            "测试用户",
            "13800138000",
            "上海市 上海市 浦东新区 测试路1号",
            "",
            "",
            "CODE1",
            "2026-09-29 16:50:54",
            "",
            "",
            "",
            "",
            "",
            "",
            "无套餐配套商品",
            "抓紧发货。",
        ]
    )
    for _ in range(4):
        sheet.append([])
    sheet.append([None] * 6 + ["光明满元气鲜食玉米2袋组合", None, "黄糯玉米8根装", 1])
    sheet.append([None] * 8 + ["白糯玉米8根装", white_quantity])
    workbook.save(path)


def build_standard_sample(
    path: Path,
    *,
    bundle_summary: str = "黄糯玉米8根装×1；白糯玉米8根装×1",
    quantity: int = 1,
) -> None:
    shutil.copy2(STANDARD_TEMPLATE, path)
    workbook = load_workbook(path)
    try:
        sheet = workbook["Sheet1"]
        values = [
            "2104856428830863362",
            "光明 福利套餐四",
            "标准模板测试订单",
            500,
            quantity,
            "待发货",
            "邮寄",
            "测试用户",
            "13800138000",
            "上海市 上海市 浦东新区 测试路1号",
            "",
            "",
            "CODE1",
            "2026-10-08 14:00:00",
            "",
            "",
            "",
            "",
            "黄糯8根装+白糯8根装",
            "",
            bundle_summary,
            "抓紧发货。",
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(4, column).value = value
        workbook.save(path)
    finally:
        workbook.close()


class NddAdapterTests(unittest.TestCase):
    def test_requires_explicit_source_hint_but_matches_known_format(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "orders.xlsx"
            build_sample(path)
            profiles = load_source_profiles(PROJECT_ROOT / "config/source_profiles")

            unidentified = detect_source(path, profiles)
            self.assertEqual(unidentified.status, "needs_input")
            self.assertIn("ndd_order_v1", unidentified.candidate_profile_ids)
            self.assertEqual(
                unidentified.clarifications[0].code,
                "COMPANY_IDENTITY_UNCONFIRMED",
            )

            confirmed = detect_source(
                path,
                profiles,
                source_profile_hint="ndd_order_v1",
            )
            self.assertEqual(confirmed.status, "matched")
            self.assertEqual(confirmed.company_id, "ndd")
            self.assertIn(
                "explicit_source_profile:ndd_order_v1",
                confirmed.matched_features,
            )

    def test_standard_template_is_auto_detected_and_examples_are_ignored(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "NDD标准订单.xlsx"
            build_standard_sample(path)
            profiles = load_source_profiles(PROJECT_ROOT / "config/source_profiles")

            detected = detect_source(path, profiles)
            self.assertEqual(detected.status, "matched")
            self.assertEqual(detected.source_profile_id, "ndd_order_v1")
            self.assertIn("company_cell:Sheet1!A2", detected.matched_features)

            parsed = NddGiftOrderParsedAdapter().parse(path, PROFILE)
            self.assertEqual(parsed.status, "parsed")
            self.assertEqual(len(parsed.orders), 1)
            self.assertEqual(parsed.orders[0].source.row_numbers, (4,))
            self.assertEqual(
                parsed.orders[0].source.source_order_no,
                "2104856428830863362",
            )

    def test_standard_template_bundle_summary_change_is_blocked(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "NDD标准订单.xlsx"
            build_standard_sample(path, bundle_summary="黄糯玉米8根装×2")

            parsed = NddGiftOrderParsedAdapter().parse(path, PROFILE)

            self.assertIn(
                "NDD_BUNDLE_COMPOSITION_CHANGED",
                {request.code for request in parsed.clarifications},
            )

    def test_standard_template_runs_end_to_end_to_guanyi_excel(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "NDD标准订单.xlsx"
            build_standard_sample(source)

            result = convert_order_file(
                source,
                PROJECT_ROOT / "assets/catalog/current_product_catalog.csv",
                root / "output",
            )

            self.assertEqual(result.status, "ready")
            self.assertEqual(result.parsed_order_count, 1)
            self.assertEqual(result.output_row_count, 2)
            self.assertEqual(len(result.outputs), 2)
            self.assertTrue(result.upload_manifest)

            workbook = load_workbook(result.outputs[0], data_only=False)
            try:
                sheet = workbook["Sheet1"]
                self.assertEqual(sheet.cell(2, 2).value, "NDD2104856428830863362A")
                self.assertEqual(sheet.cell(3, 2).value, "NDD2104856428830863362A")
                self.assertEqual(sheet.cell(2, 6).value, "JTW8E1")
                self.assertEqual(sheet.cell(3, 6).value, "JTBN1E1-EH")
                self.assertEqual(sheet.cell(2, 10).value, 1)
                self.assertEqual(sheet.cell(3, 10).value, 1)
            finally:
                workbook.close()

    def test_standard_template_quantity_multiplies_each_bundle_component(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "NDD标准订单_数量2.xlsx"
            build_standard_sample(source, quantity=2)

            result = convert_order_file(
                source,
                PROJECT_ROOT / "assets/catalog/current_product_catalog.csv",
                root / "output",
            )

            self.assertEqual(result.status, "ready")
            workbook = load_workbook(result.outputs[0], data_only=False)
            try:
                sheet = workbook["Sheet1"]
                self.assertEqual(sheet.cell(2, 6).value, "JTW8E1")
                self.assertEqual(sheet.cell(3, 6).value, "JTBN1E1-EH")
                self.assertEqual(sheet.cell(2, 10).value, 2)
                self.assertEqual(sheet.cell(3, 10).value, 2)
            finally:
                workbook.close()

    def test_parses_orders_and_splits_region_after_bundle_guard(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "orders.xlsx"
            build_sample(path)
            parsed = NddGiftOrderParsedAdapter().parse(path, PROFILE)

            self.assertEqual(parsed.status, "parsed")
            self.assertEqual(len(parsed.orders), 1)
            order = parsed.orders[0]
            self.assertEqual(order.source.source_order_no, "2104856428830863362")
            self.assertEqual(order.recipient.province, "上海市")
            self.assertEqual(order.recipient.city, "上海市")
            self.assertEqual(order.recipient.district, "浦东新区")
            self.assertEqual(order.source_note, "抓紧发货。")
            self.assertEqual(order.items[0].source_product_name, "光明 福利套餐四")

    def test_changed_bundle_composition_is_blocked(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "orders.xlsx"
            build_sample(path, white_quantity=2)
            parsed = NddGiftOrderParsedAdapter().parse(path, PROFILE)

            self.assertIn(
                "NDD_BUNDLE_COMPOSITION_CHANGED",
                {request.code for request in parsed.clarifications},
            )

    def test_build_matches_manual_business_fields_and_fixes_missing_values(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "orders.xlsx"
            build_sample(path)
            parsed = NddGiftOrderParsedAdapter().parse(path, PROFILE)
            catalog = ProductCatalog(
                [
                    ProductRecord(
                        "JTW8E1",
                        "光明满元气 黄糯玉米8棒家庭装",
                        "6974768564811",
                        "8袋/箱",
                        1210,
                        "上海尚舆商贸有限公司",
                    ),
                    ProductRecord(
                        "JTBN1E1-EH",
                        "光明满元气 白糯玉米8棒家庭装",
                        "6974768564842",
                        "8袋/箱",
                        1210,
                        "上海尚舆商贸有限公司",
                    ),
                ]
            )

            built = build_custom_import_lines(
                parsed.orders,
                catalog,
                RULES,
                GUANYI_PROFILE,
            )

            self.assertEqual(built.status, "ready")
            self.assertEqual(len(built.lines), 2)
            self.assertEqual(
                [line["商品代码"] for line in built.lines],
                ["JTW8E1", "JTBN1E1-EH"],
            )
            self.assertEqual(
                [line["商品名称"] for line in built.lines],
                ["黄糯玉米8根装", "白糯玉米8根装"],
            )
            self.assertEqual(
                {line["平台单号"] for line in built.lines},
                {"NDD2104856428830863362A"},
            )
            self.assertEqual(
                {line["物流公司"] for line in built.lines},
                {"中通快递（重货）"},
            )
            self.assertEqual(
                {line["卖家备注"] for line in built.lines},
                {"抓紧发货。"},
            )
            self.assertEqual(
                {line["联系手机"] for line in built.lines},
                {"13800138000"},
            )


if __name__ == "__main__":
    unittest.main()
