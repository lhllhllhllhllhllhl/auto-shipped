from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from auto_shipped.services.batch_convert import convert_order_batch


PROJECT_ROOT = Path(__file__).resolve().parents[1]

HEADERS = [
    "商品小计",
    "订单编号",
    "礼品名称",
    "礼品描述",
    "礼品价格",
    "数量",
    "订单状态",
    "配送方式",
    "收件人姓名",
    "收件人手机",
    "完整地址",
    "快递公司",
    "快递单号",
    "兑换码",
    "订单生成时间",
    "自提点名称",
    "自提点地址",
    "提货人姓名",
    "提货人手机",
    "礼品规格",
    "",
    "套餐配套商品信息",
    "送货时间",
]


def build_ndd_source(
    path: Path,
    *,
    order_no: str,
    note: str = "尽快发货",
    white_quantity: int = 1,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(HEADERS)
    sheet.append(
        [
            500,
            order_no,
            "光明 福利套餐四",
            "",
            "500.00",
            1,
            "待发货",
            "邮寄",
            "测试用户",
            "13800138000",
            "上海市 上海市 浦东新区 测试路1号",
            "",
            "",
            f"CODE-{order_no}",
            "2026-09-30 10:30:00",
            "",
            "",
            "",
            "",
            "",
            "",
            "无套餐配套商品",
            note,
        ]
    )
    for _ in range(4):
        sheet.append([])
    sheet.append([None] * 6 + ["光明满元气鲜食玉米2袋组合", None, "黄糯玉米8根装", 1])
    sheet.append([None] * 8 + ["白糯玉米8根装", white_quantity])
    workbook.save(path)


def build_catalog(path: Path) -> None:
    path.write_text(
        "商品代码,商品名称,规格代码,规格名称,重量,默认仓库\n"
        "JTW8E1,光明满元气 黄糯玉米8棒家庭装,6974768564811,8袋/箱,1210,上海尚舆商贸有限公司\n"
        "JTBN1E1-EH,光明满元气 白糯玉米8棒家庭装,6974768564842,8袋/箱,1210,上海尚舆商贸有限公司\n",
        encoding="gb18030",
    )


class BatchConvertServiceTests(unittest.TestCase):
    def test_merges_ready_sources_in_input_order(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            first = root / "first.xlsx"
            second = root / "second.xlsx"
            catalog = root / "catalog.csv"
            build_ndd_source(first, order_no="2104856428830863362")
            build_ndd_source(second, order_no="2104856428830863363")
            build_catalog(catalog)

            result = convert_order_batch(
                [first, second],
                catalog,
                root / "output",
                source_profile_hints={
                    first.name: "ndd_order_v1",
                    second.name: "ndd_order_v1",
                },
                batch_name="两份NDD",
            )

            self.assertEqual(result.status, "ready")
            self.assertEqual(result.source_count, 2)
            self.assertEqual(result.order_count, 2)
            self.assertEqual(result.item_row_count, 4)
            self.assertEqual(result.target_platform, "guanyi")
            self.assertEqual(result.target_profile_id, "guanyi_order_import_v1")
            self.assertTrue(result.upload_manifest)
            self.assertEqual(len(result.outputs), 2)

            workbook = load_workbook(result.outputs[0], data_only=False)
            try:
                sheet = workbook["Sheet1"]
                platform_numbers = [sheet.cell(row, 2).value for row in range(2, 6)]
            finally:
                workbook.close()
            self.assertEqual(
                platform_numbers,
                [
                    "NDD2104856428830863362A",
                    "NDD2104856428830863362A",
                    "NDD2104856428830863363A",
                    "NDD2104856428830863363A",
                ],
            )

    def test_duplicate_source_content_blocks_before_conversion(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            first = root / "first.xlsx"
            second = root / "second.xlsx"
            catalog = root / "catalog.csv"
            build_ndd_source(first, order_no="2104856428830863362")
            second.write_bytes(first.read_bytes())
            build_catalog(catalog)

            result = convert_order_batch([first, second], catalog, root / "output")

            self.assertEqual(result.status, "blocked")
            self.assertEqual(result.outputs, [])
            self.assertEqual(result.errors[0]["code"], "BATCH_DUPLICATE_SOURCE_FILE")

    def test_cross_file_platform_order_collision_blocks_whole_batch(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            first = root / "first.xlsx"
            second = root / "second.xlsx"
            catalog = root / "catalog.csv"
            build_ndd_source(first, order_no="2104856428830863362", note="备注一")
            build_ndd_source(second, order_no="2104856428830863362", note="备注二")
            build_catalog(catalog)

            result = convert_order_batch(
                [first, second],
                catalog,
                root / "output",
                source_profile_hints={
                    first.name: "ndd_order_v1",
                    second.name: "ndd_order_v1",
                },
            )

            self.assertEqual(result.status, "blocked")
            self.assertEqual(result.outputs, [])
            self.assertEqual(
                result.errors[0]["code"],
                "BATCH_PLATFORM_ORDER_NUMBER_COLLISION",
            )

    def test_one_uncertain_source_prevents_partial_batch_output(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            ready = root / "ready.xlsx"
            uncertain = root / "uncertain.xlsx"
            catalog = root / "catalog.csv"
            build_ndd_source(ready, order_no="2104856428830863362")
            build_ndd_source(
                uncertain,
                order_no="2104856428830863363",
                white_quantity=2,
            )
            build_catalog(catalog)

            result = convert_order_batch(
                [ready, uncertain],
                catalog,
                root / "output",
                source_profile_hints={
                    ready.name: "ndd_order_v1",
                    uncertain.name: "ndd_order_v1",
                },
            )

            self.assertEqual(result.status, "needs_input")
            self.assertEqual(result.outputs, [])
            self.assertIsNone(result.upload_manifest)
            self.assertIn(
                "NDD_BUNDLE_COMPOSITION_CHANGED",
                {item["code"] for item in result.clarifications},
            )


if __name__ == "__main__":
    unittest.main()
