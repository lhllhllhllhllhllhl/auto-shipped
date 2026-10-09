from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from auto_shipped.services.convert import convert_order_file


PROJECT_ROOT = Path(__file__).resolve().parents[1]


class ConvertServiceTests(unittest.TestCase):
    def test_confirmed_tiantian_pipeline_creates_single_guanyi_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_profiles = root / "source_profiles"
            platform_profiles = root / "platform_profiles"
            platform_rules = root / "platform_rules"
            output_dir = root / "output"
            source_profiles.mkdir()
            platform_profiles.mkdir()
            platform_rules.mkdir()

            source_profile = {
                "profile_id": "tiantian_warehouse_v2",
                "company_id": "tiantian",
                "source_label": "恬田代发仓",
                "order_type": "warehouse_order",
                "adapter": "tiantian_warehouse_parsed",
                "enabled": True,
                "detection": {
                    "sheet_name": "仓库订单",
                    "header_row": 1,
                    "required_headers": [
                        "订单编号",
                        "来源单号",
                        "收件人",
                        "联系方式",
                        "省",
                        "市",
                        "区",
                        "省市区详细地址",
                        "商品编码",
                        "商品名称",
                        "数量",
                    ],
                },
            }
            (source_profiles / "tiantian.json").write_text(
                json.dumps(source_profile, ensure_ascii=False),
                encoding="utf-8",
            )

            routing = {
                "routes": [
                    {
                        "route_id": "tiantian_to_guanyi",
                        "source_profile_id": "tiantian_warehouse_v2",
                        "order_type": "warehouse_order",
                        "target_platform": "guanyi",
                        "target_profile_id": "guanyi_order_import_v1",
                        "platform_rules_profile_id": "guanyi_tiantian_test",
                        "business_rule_scope_id": "test_scope",
                    }
                ]
            }
            routing_path = root / "routing.json"
            routing_path.write_text(json.dumps(routing, ensure_ascii=False), encoding="utf-8")

            company_registry_path = root / "companies.json"
            company_registry_path.write_text(
                json.dumps(
                    {
                        "schema_version": "1.0",
                        "registry_id": "test_company_registry",
                        "companies": [
                            {
                                "company_id": "tiantian",
                                "display_name": "恬田",
                                "status": "active",
                                "enabled": True,
                                "source_profiles": [
                                    {
                                        "source_profile_id": "tiantian_warehouse_v2",
                                        "status": "active",
                                        "enabled": True,
                                    }
                                ],
                                "workflows": [
                                    {
                                        "workflow_id": "tiantian_test_workflow",
                                        "source_profile_id": "tiantian_warehouse_v2",
                                        "route_id": "tiantian_to_guanyi",
                                        "target_platform": "guanyi",
                                        "target_profile_id": "guanyi_order_import_v1",
                                        "platform_rules_profile_id": "guanyi_tiantian_test",
                                        "runtime_enabled": True,
                                    }
                                ],
                            }
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            guanyi_profile = json.loads(
                (PROJECT_ROOT / "config/platform_profiles/guanyi_order_import_v1.json").read_text(
                    encoding="utf-8"
                )
            )
            (platform_profiles / "guanyi.json").write_text(
                json.dumps(guanyi_profile, ensure_ascii=False),
                encoding="utf-8",
            )

            rules = {
                "profile_id": "guanyi_tiantian_test",
                "platform_order_number": {
                    "strategy": "source_order_no",
                    "confirmed": True,
                },
                "buyer_member_policy": {
                    "strategy": "fixed",
                    "value": "张",
                    "confirmed": True,
                },
                "defaults": {
                    "store": {"value": "光明满元气", "confirmed": True},
                    "payment_amount": {"value": 0, "confirmed": True},
                    "unit_price": {"value": 0, "confirmed": True},
                    "freight": {"value": 0, "confirmed": True},
                    "warehouse_name": {"value": "上海尚舆商贸有限公司", "confirmed": True},
                    "payment_method": {"value": "网银在线", "confirmed": True},
                    "order_type": {"value": "销售订单", "confirmed": True},
                    "is_mobile_order": {"value": 0, "confirmed": True},
                    "is_cod": {"value": 0, "confirmed": True},
                    "is_distributor_order": {"value": 0, "confirmed": True},
                },
                "logistics_policy": {
                    "strategy": "source_override_then_confirmed_packaging_rules",
                    "source_carrier_priority": True,
                    "confirmed": True,
                    "default_carrier": "韵达快递",
                    "heavy_carrier": "中通快递（重货）",
                    "heavy_conditions": [
                        {
                            "match": {"product_codes": ["P1"]},
                            "minimum_quantity": 2,
                        }
                    ],
                },
                "note_routing_policy": {
                    "strategy": "courier_instructions_to_address_other_notes_to_seller",
                    "explicit_extension_field": "delivery_instruction",
                    "courier_keywords": ["送货上门", "放门口"],
                    "split_pattern": r"[；;，,\n\r]+",
                    "address_note_open": "（",
                    "address_note_close": "）",
                    "separator": "；",
                    "confirmed": True,
                },
                "seller_remark_policy": {
                    "strategy": "source_note_then_system_notes_deduplicated",
                    "separator": "；",
                    "confirmed": True,
                },
                "business_rule_coverage": {
                    "catalog_id": "test_catalog",
                    "scope_id": "test_scope",
                    "rules": {
                        "TEST-CONVERSION-RULE": {
                            "status": "implemented",
                            "confirmed": True,
                            "implementation_ref": "test.impl",
                        }
                    },
                },
            }
            (platform_rules / "rules.json").write_text(
                json.dumps(rules, ensure_ascii=False),
                encoding="utf-8",
            )

            business_rule_catalog_path = root / "business_rules.json"
            business_rule_catalog_path.write_text(
                json.dumps(
                    {
                        "schema_version": "1.0",
                        "catalog_id": "test_catalog",
                        "rules": {
                            "TEST-CONVERSION-RULE": {
                                "title": "测试转换规则",
                                "module": "test",
                                "phase": "conversion",
                                "source_refs": ["test"],
                                "allowed_implementation_refs": ["test.impl"],
                            }
                        },
                        "scopes": {
                            "test_scope": {
                                "required_rules": {
                                    "conversion": ["TEST-CONVERSION-RULE"]
                                }
                            }
                        },
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            catalog_path = root / "catalog.csv"
            catalog_path.write_text(
                "商品代码,商品名称,规格代码,规格名称,重量,默认仓库\n"
                "P1,管易测试商品,S1,测试规格,100,上海仓\n",
                encoding="gb18030",
            )
            mappings_path = root / "mappings.json"
            mappings_path.write_text(
                json.dumps(
                    {
                        "mappings": {
                            "EXT-1": {
                                "product_code": "P1",
                                "spec_code": "S1",
                                "confirmed": True,
                            }
                        }
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            source_path = root / "source.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "仓库订单"
            sheet.append(
                [
                    "订单编号",
                    "来源单号",
                    "下单时间",
                    "收件人",
                    "联系方式",
                    "省",
                    "市",
                    "区",
                    "省市区详细地址",
                    "商品编码",
                    "商品名称",
                    "数量",
                    "客户备注",
                    "物流公司",
                    "物流单号",
                    "渠道名称",
                ]
            )
            sheet.append(
                [
                    "W1",
                    "SRC-1",
                    "2026-09-20 14:09:15",
                    "测试用户",
                    "13800138000",
                    "上海",
                    "上海市",
                    "宝山区",
                    "上海市宝山区测试路1号",
                    "EXT-1",
                    "测试商品",
                    2,
                    "",
                    "",
                    "",
                    "测试渠道",
                ]
            )
            workbook.save(source_path)

            result = convert_order_file(
                source_path,
                catalog_path,
                output_dir,
                source_profiles_dir=source_profiles,
                company_registry_path=company_registry_path,
                routing_path=routing_path,
                platform_profiles_dir=platform_profiles,
                platform_rules_dir=platform_rules,
                business_rule_catalog_path=business_rule_catalog_path,
                mappings_path=mappings_path,
                guanyi_template_path=PROJECT_ROOT
                / "assets/templates/guanyi/自定义订单导入模板.xlsx",
            )

            self.assertEqual(result.status, "ready")
            self.assertEqual(result.parsed_order_count, 1)
            self.assertEqual(result.output_row_count, 1)
            self.assertEqual(len(result.outputs), 2)
            self.assertIsNotNone(result.upload_manifest)
            manifest = json.loads(Path(result.upload_manifest).read_text(encoding="utf-8"))
            self.assertEqual(manifest["status"], "ready")
            self.assertEqual(manifest["counts"], {"orders": 1, "item_rows": 1})
            self.assertFalse(manifest["safety"]["upload_authorized"])

            output_xlsx = Path(result.outputs[0])
            rendered = load_workbook(output_xlsx, read_only=False, data_only=True)
            output_sheet = rendered["Sheet1"]
            self.assertEqual(output_sheet.max_row, 2)
            self.assertEqual(output_sheet["B2"].value, "SRC-1A")
            self.assertEqual(output_sheet["F2"].value, "P1")
            self.assertEqual(output_sheet["P2"].value, "13800138000")
            self.assertEqual(output_sheet["Q2"].value, "13800138000")
            self.assertEqual(output_sheet["AA2"].value, "中通快递（重货）")
            rendered.close()


if __name__ == "__main__":
    unittest.main()
