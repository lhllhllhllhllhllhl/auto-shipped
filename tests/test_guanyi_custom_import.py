from __future__ import annotations

import json
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from auto_shipped.catalog import WhiteLabelLibrary
from auto_shipped.catalog.product_catalog import ProductCatalog, ProductRecord
from auto_shipped.integrations.feishu.daily_sequence_ledger import (
    DailySequenceLedgerRecord,
)
from auto_shipped.domain import (
    ParsedItem,
    ParsedOrder,
    ParsedRecipient,
    ParsedShipmentFacts,
    ParsedSource,
)
from auto_shipped.platforms.guanyi import build_custom_import_lines, render_custom_import


PROJECT_ROOT = Path(__file__).resolve().parents[1]


class _MemoryDailySequenceGateway:
    def __init__(self) -> None:
        self.records: list[DailySequenceLedgerRecord] = []

    def current_user_id(self) -> str:
        return "ou_sam_test"

    def list_records(self) -> list[DailySequenceLedgerRecord]:
        return list(self.records)

    def append_record(self, record: DailySequenceLedgerRecord) -> None:
        self.records.append(record)


def confirmed_rules() -> dict:
    return {
        "platform_order_number": {"strategy": "preferred_order_identity", "confirmed": True},
        "store_assignment_policy": {
            "strategy": "all_items_same_store_by_product",
            "confirmed": True,
            "rules": [
                {
                    "store": "光明满元气",
                    "match": {"product_codes": ["P1"], "spec_codes": ["S1"]},
                }
            ],
        },
        "field_output_policy": {
            "product_name": {
                "strategy": "catalog_product_name",
                "confirmed": True,
            },
            "spec_name": {
                "strategy": "catalog_spec_name",
                "confirmed": True,
            },
            "order_created_at": {
                "strategy": "source_ordered_at",
                "confirmed": True,
            },
            "contact_columns": {
                "strategy": "mobile_to_both_phone_preserved",
                "confirmed": True,
            },
            "recipient_region": {
                "strategy": "source_components",
                "confirmed": True,
            },
        },
        "buyer_member_policy": {
            "strategy": "fixed",
            "value": "张",
            "confirmed": True,
        },
        "defaults": {
            "payment_amount": {"value": 0, "confirmed": True},
            "unit_price": {"value": 0, "confirmed": True},
            "freight": {"value": 0, "confirmed": True},
            "warehouse_name": {"value": "上海尚舆商贸有限公司", "confirmed": True},
            "payment_method": {"value": "网银在线", "confirmed": True},
            "order_type": {"value": "销售订单", "confirmed": True},
            "is_mobile_order": {"value": 0, "confirmed": True},
            "is_cod": {"value": 0, "confirmed": True},
            "is_distributor_order": {"value": 0, "confirmed": True},
        },
        "logistics_policy": {
            "strategy": "source_override_then_confirmed_packaging_rules",
            "source_carrier_priority": True,
            "source_carrier_aliases": {
                "顺丰速运": "顺丰速运",
            },
            "confirmed": True,
            "default_carrier": "韵达快递",
            "heavy_carrier": "中通快递（重货）",
            "heavy_conditions": [
                {
                    "match": {"product_codes": ["P1"]},
                    "minimum_quantity": 2,
                }
            ],
        },
        "note_routing_policy": {
            "strategy": "courier_instructions_to_address_other_notes_to_seller",
            "explicit_extension_field": "delivery_instruction",
            "courier_keywords": [
                "送货上门",
                "送达",
                "派送",
                "配送",
                "放门口",
                "驿站",
                "快递柜",
                "丰巢",
                "电话联系",
                "来电",
                "联系收件人",
                "敲门",
                "门铃",
                "轻放",
            ],
            "split_pattern": r"[；;，,\n\r]+",
            "address_note_open": "（",
            "address_note_close": "）",
            "separator": "；",
            "confirmed": True,
        },
        "seller_remark_policy": {
            "strategy": "source_note_then_system_notes_deduplicated",
            "separator": "；",
            "confirmed": True,
        },
    }


def parsed_order() -> ParsedOrder:
    return ParsedOrder(
        record_key="order-1",
        source=ParsedSource(
            profile_id="tiantian_warehouse_v2",
            source_label="恬田代发仓",
            order_type="warehouse_order",
            source_order_no="SRC-001",
            file_name="sample.xlsx",
            file_fingerprint="sha256:test",
            sheet_name="仓库订单",
            row_numbers=(2,),
        ),
        recipient=ParsedRecipient(
            name="测试用户",
            mobile="13800138000",
            province="上海",
            city="上海市",
            district="宝山区",
            raw_address="上海市宝山区测试路1号",
        ),
        items=[
            ParsedItem(
                source_product_code="EXT-1",
                source_product_name="测试商品",
                quantity=2,
                source_line_no=2,
            )
        ],
        ordered_at="2026-09-20 14:09:15",
    )


class GuanyiCustomImportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.catalog = ProductCatalog(
            [ProductRecord("P1", "管易测试商品", "S1", "测试规格", 100, "上海仓")],
            [
                {
                    "source_profile_id": "*",
                    "identifier_type": "product_code",
                    "source_value": "EXT-1",
                    "product_code": "P1",
                    "spec_code": "S1",
                    "confirmed": True,
                },
                {
                    "source_profile_id": "*",
                    "identifier_type": "product_code",
                    "source_value": "EXT-1",
                    "source_spec": "SWSH8072",
                    "product_code": "P1",
                    "spec_code": "S1",
                    "confirmed": True,
                },
            ],
        )
        self.profile = json.loads(
            (PROJECT_ROOT / "config/platform_profiles/guanyi_order_import_v1.json").read_text(
                encoding="utf-8"
            )
        )

    def test_confirmed_rules_build_one_row_per_item(self) -> None:
        result = build_custom_import_lines(
            [parsed_order()],
            self.catalog,
            confirmed_rules(),
            self.profile,
        )
        self.assertEqual(result.status, "ready")
        self.assertEqual(len(result.lines), 1)
        line = result.lines[0]
        self.assertEqual(line["店铺"], "光明满元气")
        self.assertEqual(line["平台单号"], "SRC-001")
        self.assertEqual(line["商品代码"], "P1")
        self.assertEqual(line["规格代码"], "S1")
        self.assertEqual(line["商品名称"], "管易测试商品")
        self.assertEqual(line["规格名称"], "测试规格")
        self.assertEqual(line["数量"], 2)
        self.assertEqual(line["收货地址"], "上海市宝山区测试路1号")
        self.assertEqual(line["联系电话"], "13800138000")
        self.assertEqual(line["联系手机"], "13800138000")
        self.assertEqual(line["省"], "上海")
        self.assertEqual(line["市"], "上海市")
        self.assertEqual(line["区"], "宝山区")
        self.assertEqual(line["订单创建时间"], "2026-09-20 14:09:15")
        self.assertEqual(line["物流公司"], "中通快递（重货）")

    def test_generic_fallback_reuses_declared_legacy_item_expansion(self) -> None:
        order = parsed_order()
        order.source = ParsedSource(
            profile_id="generic_company_manual_v1",
            source_label="已登记公司通用手工订单兜底",
            order_type="manual_channel_order",
            source_order_no="KLD20261019001",
            file_name="manual.txt",
            file_fingerprint="sha256:kld",
            sheet_name="",
            row_numbers=(1,),
        )
        order.items = [
            ParsedItem(
                source_product_name="花糯玉米",
                source_spec="9根装",
                quantity=1,
                source_line_no=1,
            )
        ]
        catalog = ProductCatalog(
            [ProductRecord("P1", "花糯玉米", "S1", "8棒家庭装", 100, "上海仓")],
            [],
            item_expansion_rules=[
                {
                    "rule_id": "kld-flower-9",
                    "company_id": "shanghai_shimao_plaza",
                    "source_profile_id": "kld_text_v1",
                    "target_platform": "guanyi",
                    "identifier_type": "product_name",
                    "source_value": "花糯玉米",
                    "source_spec": "9根装",
                    "reuse_scope": "future",
                    "confirmed": True,
                    "components": [
                        {
                            "product_code": "P1",
                            "spec_code": "S1",
                            "quantity_multiplier": 1,
                        }
                    ],
                }
            ],
        )
        rules = confirmed_rules()
        rules["item_expansion_policy"] = {
            "strategy": "confirmed_source_item_expansions",
            "on_unmatched": "standard_product_mapping",
            "confirmed": True,
        }

        result = build_custom_import_lines(
            [order],
            catalog,
            rules,
            self.profile,
            source_company_id="shanghai_shimao_plaza",
            source_profile_aliases=("kld_text_v1",),
        )

        self.assertEqual(result.status, "ready", result.clarifications)
        self.assertEqual(len(result.lines), 1)
        self.assertEqual(result.lines[0]["商品代码"], "P1")
        self.assertEqual(result.lines[0]["规格代码"], "S1")

    def test_generic_fallback_passes_company_scope_to_store_assignment(self) -> None:
        catalog = ProductCatalog(
            list(self.catalog.records),
            list(self.catalog.mapping_records),
            manual_store_routes=[
                {
                    "scope_id": "shanghai_shimao_plaza",
                    "product_code": "P1",
                    "spec_code": "S1",
                    "store": "光明满元气",
                    "status": "active",
                },
                {
                    "scope_id": "fuze",
                    "product_code": "P1",
                    "spec_code": "S1",
                    "store": "淘橘子",
                    "status": "active",
                },
            ],
        )
        rules = confirmed_rules()
        rules["store_assignment_policy"] = {
            "strategy": "manual_store_route_then_product",
            "confirmed": True,
            "include_company_scope": True,
            "include_source_profile_scope": False,
            "fallback_scope_ids": ["manual_all"],
            "allowed_stores": ["光明满元气", "淘橘子"],
        }

        result = build_custom_import_lines(
            [parsed_order()],
            catalog,
            rules,
            self.profile,
            source_company_id="shanghai_shimao_plaza",
        )

        self.assertEqual(result.status, "ready", result.clarifications)
        self.assertEqual(result.lines[0]["店铺"], "光明满元气")

    def test_non_corn_store_fallback_is_ready_and_reported_non_blocking(self) -> None:
        rules = confirmed_rules()
        rules["store_assignment_policy"] = {
            "strategy": "manual_store_route_then_product",
            "confirmed": True,
            "include_company_scope": True,
            "include_source_profile_scope": False,
            "fallback_scope_ids": ["manual_all"],
            "allowed_stores": ["光明满元气", "淘橘子"],
            "product_fallback_rules": [
                {
                    "rule_id": "confirmed_non_corn_to_taojuzi",
                    "store": "淘橘子",
                    "match": {"product_name_excludes_any": ["玉米"]},
                    "result_notice": {
                        "code": "NON_CORN_TO_TAOJUZI_FALLBACK",
                        "message": "非玉米兜底",
                    },
                }
            ],
        }

        result = build_custom_import_lines(
            [parsed_order()],
            self.catalog,
            rules,
            self.profile,
            source_company_id="shanghai_shimao_plaza",
        )

        self.assertEqual(result.status, "ready", result.clarifications)
        self.assertEqual(result.lines[0]["店铺"], "淘橘子")
        self.assertEqual(len(result.store_assignment_notices), 1)
        notice = result.store_assignment_notices[0]
        self.assertEqual(notice["code"], "NON_CORN_TO_TAOJUZI_FALLBACK")
        self.assertEqual(notice["store"], "淘橘子")
        self.assertFalse(notice["blocking"])
        self.assertNotIn("测试用户", json.dumps(notice, ensure_ascii=False))

    def test_platform_order_number_prefers_order_number_over_source_number(self) -> None:
        order = parsed_order()
        order.source = ParsedSource(
            profile_id=order.source.profile_id,
            source_label=order.source.source_label,
            order_type=order.source.order_type,
            order_no="ORDER-001",
            source_order_no="SOURCE-999",
            file_name=order.source.file_name,
            file_fingerprint=order.source.file_fingerprint,
            sheet_name=order.source.sheet_name,
            row_numbers=order.source.row_numbers,
        )
        result = build_custom_import_lines(
            [order],
            self.catalog,
            confirmed_rules(),
            self.profile,
        )
        self.assertEqual(result.status, "ready")
        self.assertEqual(result.lines[0]["平台单号"], "ORDER-001")

    def test_code_only_multi_spec_candidate_asks_scope_before_persistence(self) -> None:
        catalog = ProductCatalog(
            [
                ProductRecord("P1", "测试商品", "S1", "红色", 100, "上海仓"),
                ProductRecord("P1", "测试商品", "S2", "蓝色", 100, "上海仓"),
            ],
            [
                {
                    "source_profile_id": "tiantian_warehouse_v2",
                    "identifier_type": "product_code",
                    "source_value": "EXT-1",
                    "source_spec": "",
                    "product_code": "P1",
                    "spec_code": "S1",
                    "confirmed": True,
                }
            ],
        )

        result = build_custom_import_lines(
            [parsed_order()],
            catalog,
            confirmed_rules(),
            self.profile,
        )

        self.assertEqual(result.status, "needs_input")
        request = next(
            item
            for item in result.clarifications
            if item.code == "CONFIRM_GUANYI_PRODUCT_CANDIDATE"
        )
        self.assertIn("红色", request.question)
        self.assertIn("蓝色", request.question)
        self.assertIn("仅本次使用", request.question)
        self.assertIn("以后都沿用", request.question)

    def test_unique_unconfirmed_candidate_builds_preview_decision(self) -> None:
        catalog = ProductCatalog(
            [ProductRecord("P1", "测试商品", "S1", "唯一规格", 100, "上海仓")]
        )
        order = parsed_order()
        order.items[0] = ParsedItem(
            source_product_code="P1",
            source_product_name="测试商品",
            source_spec="",
            quantity=2,
            source_line_no=2,
        )

        result = build_custom_import_lines(
            [order],
            catalog,
            confirmed_rules(),
            self.profile,
        )

        self.assertEqual(result.status, "preview_ready")
        self.assertEqual(len(result.lines), 1)
        self.assertEqual(result.lines[0]["商品代码"], "P1")
        self.assertEqual(result.lines[0]["规格代码"], "S1")
        self.assertEqual(len(result.decisions), 1)
        decision = result.decisions[0]
        self.assertEqual(decision.code, "CONFIRM_GUANYI_PRODUCT_MAPPING")
        self.assertEqual(
            decision.persistence_options,
            ("current_batch", "future"),
        )
        self.assertIn("仅本次使用", decision.confirmation_question)
        self.assertIn("以后遇到相同来源时沿用", decision.confirmation_question)

    def test_confirmed_code_name_mismatch_builds_preview_and_mapping_evidence(self) -> None:
        order = parsed_order()
        order.items[0] = ParsedItem(
            source_product_code="EXT-1",
            source_item_text="完全不同的新商品1件",
            source_product_name="完全不同的新商品",
            quantity=1,
            unit="件",
            source_line_no=2,
        )
        catalog = ProductCatalog(
            [ProductRecord("P1", "管易测试商品", "S1", "测试规格", 100, "上海仓")],
            [
                {
                    "mapping_id": "official-ext-1",
                    "source_profile_id": "tiantian_warehouse_v2",
                    "identifier_type": "product_code",
                    "source_value": "EXT-1",
                    "source_spec": "",
                    "product_code": "P1",
                    "spec_code": "S1",
                    "confirmed": True,
                    "reuse_scope": "future",
                    "mapping_authority": "feishu_official",
                }
            ],
            {
                "tiantian_warehouse_v2": {
                    "source_identifier_type": "product_code",
                    "require_identity_guard_for_blank_source_spec": True,
                    "mapping_identity_guards": [
                        {
                            "identifier_type": "product_code",
                            "source_value": "EXT-1",
                            "allowed_product_names": ["已确认商品"],
                        }
                    ],
                }
            },
            mapping_revision="23",
            mapping_snapshot_sha256="sha256:test",
        )

        result = build_custom_import_lines(
            [order],
            catalog,
            confirmed_rules(),
            self.profile,
        )

        self.assertEqual(result.status, "preview_ready")
        self.assertEqual(len(result.lines), 1)
        self.assertEqual(result.lines[0]["商品代码"], "P1")
        self.assertEqual(result.lines[0]["规格代码"], "S1")
        self.assertEqual(len(result.decisions), 1)
        self.assertEqual(
            result.decisions[0].code,
            "SOURCE_NAME_DIFFERS_FROM_CONFIRMED_CODE_MAPPING",
        )
        self.assertEqual(
            result.decisions[0].persistence_options,
            ("current_batch",),
        )
        self.assertEqual(len(result.mapping_evidence), 1)
        evidence = result.mapping_evidence[0]
        self.assertEqual(evidence["mapping_revision"], "23")
        self.assertEqual(evidence["mapping_ids"], ["official-ext-1"])
        self.assertEqual(evidence["initial_split"]["unit"], "件")
        self.assertEqual(result.lines[0]["数量"], 1)
        self.assertNotIn("recipient", evidence)

    def test_current_batch_confirmation_suppresses_code_name_mismatch_warning(self) -> None:
        order = parsed_order()
        order.items[0] = ParsedItem(
            source_product_code="EXT-1",
            source_item_text="完全不同的新商品1件",
            source_product_name="完全不同的新商品",
            quantity=1,
            unit="件",
            source_line_no=2,
        )
        common = {
            "source_profile_id": "tiantian_warehouse_v2",
            "identifier_type": "product_code",
            "source_value": "EXT-1",
            "source_spec": "",
            "product_code": "P1",
            "spec_code": "S1",
            "confirmed": True,
        }
        catalog = ProductCatalog(
            [ProductRecord("P1", "管易测试商品", "S1", "测试规格", 100, "上海仓")],
            [
                {
                    **common,
                    "mapping_id": "official-ext-1",
                    "reuse_scope": "future",
                    "mapping_authority": "feishu_official",
                },
                {
                    **common,
                    "mapping_id": "batch-ext-1",
                    "reuse_scope": "current_batch",
                    "mapping_authority": "user_batch_confirmation",
                    "source_fingerprint": order.source.file_fingerprint,
                },
            ],
            {
                "tiantian_warehouse_v2": {
                    "source_identifier_type": "product_code",
                    "require_identity_guard_for_blank_source_spec": True,
                    "mapping_identity_guards": [
                        {
                            "identifier_type": "product_code",
                            "source_value": "EXT-1",
                            "allowed_product_names": ["已确认商品"],
                        }
                    ],
                }
            },
        )

        result = build_custom_import_lines(
            [order],
            catalog,
            confirmed_rules(),
            self.profile,
        )

        self.assertEqual(result.status, "ready")
        self.assertEqual(result.decisions, [])
        self.assertEqual(result.lines[0]["数量"], 1)

    def test_single_low_confidence_fuzzy_candidate_blocks_instead_of_previewing(
        self,
    ) -> None:
        catalog = ProductCatalog(
            [
                ProductRecord(
                    "J-579",
                    "和匠日本进口搓澡巾手套去鸡皮去角质搓泥双面两用沐浴花洗澡巾",
                    "J-579",
                    "双面两用",
                    100,
                    "上海仓",
                )
            ],
            source_policies={
                "tiantian_warehouse_v2": {
                    "source_identifier_type": "product_name",
                    "allow_exact_catalog_name": False,
                    "fuzzy_name_action": "candidate_requires_user_confirmation",
                }
            },
        )
        order = parsed_order()
        order.items[0] = ParsedItem(
            source_product_code="",
            source_product_name="搓澡巾",
            source_spec="超硬深蓝色",
            quantity=1,
            source_line_no=2,
        )

        result = build_custom_import_lines(
            [order],
            catalog,
            confirmed_rules(),
            self.profile,
        )

        self.assertEqual(result.status, "needs_input")
        self.assertEqual(result.lines, [])
        self.assertEqual(result.decisions, [])
        request = next(
            item
            for item in result.clarifications
            if item.code == "CONFIRM_LOW_CONFIDENCE_GUANYI_PRODUCT_CANDIDATE"
        )
        self.assertIn("J-579", request.question)
        self.assertIn("超硬深蓝色", request.question)
        self.assertIn("仅供查找，未作为默认值", request.question)
        self.assertIn("关键语义属性", request.reason)

    def test_blocking_product_hides_preview_only_decisions_until_rerun(self) -> None:
        catalog = ProductCatalog(
            [ProductRecord("P1", "测试商品", "S1", "唯一规格", 100, "上海仓")]
        )
        order = parsed_order()
        order.items = [
            ParsedItem(
                source_product_code="P1",
                source_product_name="测试商品",
                source_spec="",
                quantity=1,
                source_line_no=2,
            ),
            ParsedItem(
                source_product_code="UNKNOWN",
                source_product_name="未登记商品",
                source_spec="粉色",
                quantity=1,
                source_line_no=3,
            ),
        ]

        result = build_custom_import_lines(
            [order],
            catalog,
            confirmed_rules(),
            self.profile,
        )

        self.assertEqual(result.status, "needs_input")
        self.assertEqual(result.lines, [])
        self.assertEqual(result.decisions, [])
        self.assertTrue(
            any(
                request.code == "PROVIDE_GUANYI_PRODUCT_MAPPING"
                for request in result.clarifications
            )
        )

    def test_current_batch_mapping_emits_nonblocking_reuse_candidate(self) -> None:
        order = parsed_order()
        catalog = ProductCatalog(
            [ProductRecord("P1", "测试商品", "S1", "唯一规格", 100, "上海仓")],
            [
                {
                    "source_profile_id": "tiantian_warehouse_v2",
                    "identifier_type": "product_code",
                    "source_value": "EXT-1",
                    "source_spec": "",
                    "product_code": "P1",
                    "spec_code": "S1",
                    "confirmed": True,
                    "reuse_scope": "current_batch",
                    "mapping_authority": "user_batch_confirmation",
                    "source_fingerprint": order.source.file_fingerprint,
                }
            ],
        )

        result = build_custom_import_lines(
            [order],
            catalog,
            confirmed_rules(),
            self.profile,
            source_company_id="tiantian",
        )

        self.assertEqual(result.status, "ready")
        self.assertEqual(len(result.mapping_reuse_candidates), 1)
        candidate = result.mapping_reuse_candidates[0]
        self.assertEqual(candidate.source_value, "EXT-1")
        self.assertEqual(candidate.product_code, "P1")
        self.assertEqual(candidate.spec_code, "S1")
        self.assertEqual(candidate.future_scope_options, ("source",))

    def test_formal_mapping_does_not_emit_reuse_candidate(self) -> None:
        order = parsed_order()
        catalog = ProductCatalog(
            [ProductRecord("P1", "测试商品", "S1", "唯一规格", 100, "上海仓")],
            [
                {
                    "source_profile_id": "tiantian_warehouse_v2",
                    "identifier_type": "product_code",
                    "source_value": "EXT-1",
                    "source_spec": "",
                    "product_code": "P1",
                    "spec_code": "S1",
                    "confirmed": True,
                    "reuse_scope": "future",
                    "mapping_authority": "feishu_official",
                }
            ],
        )

        result = build_custom_import_lines(
            [order],
            catalog,
            confirmed_rules(),
            self.profile,
            source_company_id="tiantian",
        )

        self.assertEqual(result.status, "ready")
        self.assertEqual(result.mapping_reuse_candidates, [])

    def test_name_candidate_offers_source_or_global_future_scope(self) -> None:
        catalog = ProductCatalog(
            [ProductRecord("P1", "蓝色硬款搓澡巾", "S1", "硬款", 100, "上海仓")],
            source_policies={
                "tiantian_warehouse_v2": {
                    "source_identifier_type": "product_name",
                    "fuzzy_name_action": "candidate_requires_user_confirmation",
                }
            },
        )
        order = parsed_order()
        order.items[0] = ParsedItem(
            source_product_code="",
            source_product_name="蓝色硬款搓澡巾",
            source_spec="",
            quantity=1,
            source_line_no=2,
        )

        result = build_custom_import_lines(
            [order],
            catalog,
            confirmed_rules(),
            self.profile,
        )

        self.assertEqual(result.status, "preview_ready")
        decision = result.decisions[0]
        self.assertEqual(decision.future_scope_options, ("source", "global"))
        self.assertIn("以后所有来源沿用", decision.confirmation_question)
        self.assertEqual(
            decision.proposed_value["confidence_tier"],
            "preview_safe",
        )
        self.assertEqual(
            decision.proposed_value["candidate_evidence"][0]["confidence_tier"],
            "preview_safe",
        )

    def test_white_label_library_adds_one_order_level_system_remark(self) -> None:
        order = parsed_order()
        order.source_note = "周末发货"
        library = WhiteLabelLibrary(
            library_id="test_white_label",
            version=1,
            remark_token="白标商品",
            sku_keys=frozenset({("P1", "S1")}),
        )
        result = build_custom_import_lines(
            [order],
            self.catalog,
            confirmed_rules(),
            self.profile,
            white_label_library=library,
        )
        self.assertEqual(result.status, "ready")
        self.assertEqual(result.lines[0]["卖家备注"], "周末发货；白标商品")

    def test_confirmed_bare_stick_mapping_is_blocked_for_shipping(self) -> None:
        catalog = ProductCatalog(
            [ProductRecord("P1", "黄糯玉米（裸棒）", "S1", "40棒/箱", 100, "上海仓")],
            {
                "EXT-1": {
                    "product_code": "P1",
                    "spec_code": "S1",
                    "confirmed": True,
                }
            },
            package_semantics_registry={
                "forbidden_shipping_products": [
                    {
                        "product_code": "P1",
                        "spec_code": "S1",
                        "reason": "裸棒仅用于试吃",
                    }
                ]
            },
        )
        result = build_custom_import_lines(
            [parsed_order()],
            catalog,
            confirmed_rules(),
            self.profile,
        )
        self.assertEqual(result.status, "needs_input")
        self.assertEqual(result.lines, [])
        self.assertIn(
            "SHIPPING_PRODUCT_FORBIDDEN",
            {item.code for item in result.clarifications},
        )

    def test_registered_pack_content_count_converts_only_complete_packs(self) -> None:
        catalog = ProductCatalog(
            [ProductRecord("P1", "黄糯玉米8棒家庭装", "S1", "8袋/箱", 100, "上海仓")],
            package_semantics_registry={
                "product_family_rules": [
                    {
                        "source_name_contains_any": ["黄糯"],
                        "target_name_contains_all": ["黄糯", "玉米"],
                    }
                ],
                "tracked_source_units": ["根"],
                "source_semantics": {
                    "company_x_text_v1": {
                        "confirmed": True,
                        "rules": [
                            {
                                "rule_id": "company_x_roots_to_eight_pack_v1",
                                "source_spec_values": ["根"],
                                "target_name_contains_all": ["8棒家庭装"],
                                "quantity_strategy": "divide_by_sticks_per_target_unit",
                                "sticks_per_target_unit": 8,
                            }
                        ],
                    }
                },
            },
        )
        order = parsed_order()
        order.source = ParsedSource(
            profile_id="company_x_text_v1",
            source_label="公司X文字单",
            order_type="manual_channel_order",
            source_order_no="SRC-001",
            file_name="sample.xlsx",
            file_fingerprint="sha256:test",
            sheet_name="Sheet1",
            row_numbers=(2,),
        )
        order.items = [
            ParsedItem(
                source_product_name="黄糯",
                source_spec="根",
                quantity=16,
                source_line_no=2,
            )
        ]
        ready = build_custom_import_lines(
            [order],
            catalog,
            confirmed_rules(),
            self.profile,
        )
        self.assertEqual(ready.status, "ready")
        self.assertEqual(ready.lines[0]["数量"], 2)

        order.items = [
            ParsedItem(
                source_product_name="黄糯",
                source_spec="根",
                quantity=10,
                source_line_no=2,
            )
        ]
        blocked = build_custom_import_lines(
            [order],
            catalog,
            confirmed_rules(),
            self.profile,
        )
        self.assertEqual(blocked.status, "needs_input")
        self.assertIn(
            "CONFIRM_PACKAGE_QUANTITY_CONVERSION",
            {item.code for item in blocked.clarifications},
        )

    def test_platform_order_number_supports_company_prefix_without_global_suffix(self) -> None:
        rules = confirmed_rules()
        rules["platform_order_number"] = {
            "strategy": "preferred_order_identity_with_affixes",
            "prefix": "TT",
            "suffix": "",
            "confirmed": True,
        }
        result = build_custom_import_lines(
            [parsed_order()],
            self.catalog,
            rules,
            self.profile,
        )
        self.assertEqual(result.status, "ready")
        self.assertEqual(result.lines[0]["平台单号"], "TTSRC-001")

    def test_logistics_defaults_to_yunda_below_heavy_threshold(self) -> None:
        order = parsed_order()
        order.items[0] = ParsedItem(
            source_product_code="EXT-1",
            source_product_name="测试商品",
            quantity=1,
            source_line_no=2,
        )
        result = build_custom_import_lines(
            [order],
            self.catalog,
            confirmed_rules(),
            self.profile,
        )
        self.assertEqual(result.status, "ready")
        self.assertEqual(result.lines[0]["物流公司"], "韵达快递")

    def test_source_carrier_has_priority_over_packaging_rule(self) -> None:
        order = parsed_order()
        order.shipment_facts = ParsedShipmentFacts(carrier_name="顺丰速运")
        result = build_custom_import_lines(
            [order],
            self.catalog,
            confirmed_rules(),
            self.profile,
        )
        self.assertEqual(result.status, "ready")
        self.assertEqual(result.lines[0]["物流公司"], "顺丰速运")

    def test_sam_bundle_expands_to_eight_pack_and_single_item(self) -> None:
        order = parsed_order()
        order.items = [
            ParsedItem(
                source_product_code="276388650",
                source_product_name="光明满元气鲜糯玉米（加送1根）",
                source_spec="光9-JTBN1T1-EH",
                source_display_spec="白糯8根+1根",
                quantity=2,
                source_line_no=2,
            )
        ]
        order.source_extensions["sam_group_number"] = "848"
        catalog = ProductCatalog(
            [
                ProductRecord(
                    "JTBN1E1-EH",
                    "光明满元气 白糯玉米8棒家庭装",
                    "6974768564842",
                    "8袋/箱",
                    1210,
                    "上海尚舆商贸有限公司",
                ),
                ProductRecord(
                    "GMBN1",
                    "光明满元气 白糯玉米单根装",
                    "6974768564804",
                    "50根/箱",
                    180,
                    "上海尚舆商贸有限公司",
                ),
            ],
            [],
        )
        rules = confirmed_rules()
        rules.pop("store_assignment_policy")
        rules["defaults"]["store"] = {"value": "sam", "confirmed": True}
        rules["defaults"]["freight"] = {"value": "", "confirmed": True}
        rules["field_output_policy"]["product_name"]["strategy"] = (
            "source_name_with_display_spec_for_standard_items"
        )
        rules["buyer_member_policy"] = {
            "strategy": "sam_product_type_date_group",
            "platform_prefix": "sam",
            "group_number_field": "sam_group_number",
            "group_separator": "_",
            "date_format": "%y%m%d",
            "timezone": "Asia/Shanghai",
            "product_type_rules": [
                {
                    "type_code": "YM",
                    "match": {"product_name_contains_all": ["玉米"]},
                    "confirmed": True,
                }
            ],
            "confirmed": True,
        }
        rules["item_expansion_policy"] = {
            "strategy": "confirmed_source_item_expansions",
            "confirmed": True,
            "rules": [
                {
                    "rule_id": "sam_corn_white_8_plus_1",
                    "match": {
                        "source_product_code": "276388650",
                        "source_spec": "光9-JTBN1T1-EH",
                    },
                    "components": [
                        {
                            "product_code": "JTBN1E1-EH",
                            "spec_code": "6974768564842",
                            "quantity_multiplier": 1,
                        },
                        {
                            "product_code": "GMBN1",
                            "spec_code": "6974768564804",
                            "quantity_multiplier": 1,
                        },
                    ],
                    "confirmed": True,
                }
            ],
        }
        rules["logistics_policy"]["heavy_conditions"] = [
            {
                "match": {"product_codes": ["JTBN1E1-EH"]},
                "minimum_quantity": 2,
            }
        ]

        result = build_custom_import_lines(
            [order],
            catalog,
            rules,
            self.profile,
            business_date=date(2026, 9, 29),
            official_mapping_ready=True,
        )

        self.assertEqual(result.status, "ready")
        self.assertEqual(len(result.lines), 2)
        self.assertEqual(
            [(line["商品代码"], line["规格代码"], line["数量"]) for line in result.lines],
            [
                ("JTBN1E1-EH", "6974768564842", 2),
                ("GMBN1", "6974768564804", 2),
            ],
        )
        self.assertTrue(all(line["买家会员"] == "samYM260929_848" for line in result.lines))
        self.assertTrue(all(line["价格"] == 0 for line in result.lines))
        self.assertTrue(all(line["支付金额"] == 0 for line in result.lines))
        self.assertTrue(all(line["运费"] == "" for line in result.lines))
        self.assertTrue(all(line["物流公司"] == "中通快递（重货）" for line in result.lines))
        self.assertEqual(
            [line["商品名称"] for line in result.lines],
            ["光明满元气 白糯玉米8棒家庭装", "光明满元气 白糯玉米单根装"],
        )

        order.source_extensions.pop("sam_group_number")
        order.source = ParsedSource(
            profile_id=order.source.profile_id,
            source_label=order.source.source_label,
            order_type=order.source.order_type,
            source_order_no=order.source.source_order_no,
            file_name=order.source.file_name,
            file_fingerprint=f"sha256:{'a' * 64}",
            sheet_name=order.source.sheet_name,
            row_numbers=order.source.row_numbers,
        )
        rules["buyer_member_policy"].update(
            {
                "on_missing_group_number": "generate_daily_sequence",
                "fallback_stamp_format": "%y%m%d",
                "fallback_sequence_separator": "_",
                "fallback_sequence_start": 1,
                "fallback_sequence_width": 3,
            }
        )
        daily_gateway = _MemoryDailySequenceGateway()
        with tempfile.TemporaryDirectory() as state_dir:
            generated = build_custom_import_lines(
                [order],
                catalog,
                rules,
                self.profile,
                business_datetime=datetime(
                    2026,
                    10,
                    12,
                    14,
                    30,
                    tzinfo=ZoneInfo("Asia/Shanghai"),
                ),
                buyer_member_sequence_state_dir=state_dir,
                buyer_member_sequence_gateway=daily_gateway,
                official_mapping_ready=True,
            )
            repeated = build_custom_import_lines(
                [order],
                catalog,
                rules,
                self.profile,
                business_datetime=datetime(
                    2026,
                    10,
                    12,
                    15,
                    45,
                    tzinfo=ZoneInfo("Asia/Shanghai"),
                ),
                buyer_member_sequence_state_dir=state_dir,
                buyer_member_sequence_gateway=daily_gateway,
                official_mapping_ready=True,
            )
        self.assertEqual(generated.status, "ready")
        self.assertTrue(
            all(
                line["买家会员"] == "samYM261012_001"
                for line in generated.lines
            )
        )
        self.assertTrue(
            all(
                line["买家会员"] == "samYM261012_001"
                for line in repeated.lines
            )
        )

    def test_all_confirmed_sam_eight_plus_one_rules_expand_deterministically(self) -> None:
        sam_rules = json.loads(
            (
                PROJECT_ROOT
                / "config"
                / "platform_rules"
                / "guanyi"
                / "sam_order_v1.json"
            ).read_text(encoding="utf-8")
        )
        catalog = ProductCatalog(
            [
                ProductRecord(
                    "JTBN1E1-EH",
                    "白糯8棒家庭装",
                    "6974768564842",
                    "8袋/箱",
                    1210,
                    "上海尚舆商贸有限公司",
                ),
                ProductRecord(
                    "GMBN1",
                    "白糯彩袋单棒装",
                    "6974768564804",
                    "50根/箱",
                    180,
                    "上海尚舆商贸有限公司",
                ),
                ProductRecord(
                    "JTTZE1",
                    "花糯8棒家庭装",
                    "6974768564835",
                    "8袋/箱",
                    1210,
                    "上海尚舆商贸有限公司",
                ),
                ProductRecord(
                    "GMCN-1",
                    "花糯彩袋单棒装",
                    "6974768564798",
                    "50根/箱",
                    180,
                    "上海尚舆商贸有限公司",
                ),
                ProductRecord(
                    "JTW8E1",
                    "黄糯8棒家庭装",
                    "6974768564811",
                    "8袋/箱",
                    1210,
                    "上海尚舆商贸有限公司",
                ),
                ProductRecord(
                    "GMHN-1",
                    "黄糯彩袋单棒装",
                    "6974768564729",
                    "50根/箱",
                    180,
                    "上海尚舆商贸有限公司",
                ),
            ],
            [],
        )
        cases = [
            (
                "276388650",
                "光9-JTBN1T1-EH",
                [
                    ("JTBN1E1-EH", "6974768564842", 1),
                    ("GMBN1", "6974768564804", 1),
                ],
            ),
            (
                "276388650",
                "光9-JTTZT1",
                [
                    ("JTTZE1", "6974768564835", 1),
                    ("GMCN-1", "6974768564798", 1),
                ],
            ),
            (
                "YHGMMYQXNYM31631159",
                "光-QJTW8-F9-EH",
                [
                    ("JTW8E1", "6974768564811", 1),
                    ("GMHN-1", "6974768564729", 1),
                ],
            ),
        ]
        catalog.item_expansion_rules = tuple(
            {
                "rule_id": f"sam-central-bundle-{index}",
                "company_id": "sam",
                "source_profile_id": "sam_order_v1",
                "identifier_type": "product_code",
                "source_value": source_code,
                "source_spec": source_spec,
                "target_platform": "guanyi",
                "confirmed": True,
                "reuse_scope": "future",
                "components": [
                    {
                        "component_index": component_index,
                        "product_code": product_code,
                        "spec_code": spec_code,
                        "quantity_multiplier": multiplier,
                    }
                    for component_index, (
                        product_code,
                        spec_code,
                        multiplier,
                    ) in enumerate(expected, start=1)
                ],
            }
            for index, (source_code, source_spec, expected) in enumerate(
                cases, start=1
            )
        )

        for source_code, source_spec, expected in cases:
            with self.subTest(source_code=source_code, source_spec=source_spec):
                order = parsed_order()
                order.source = ParsedSource(
                    profile_id="sam_order_v1",
                    source_label="SAM群接龙订单",
                    order_type=order.source.order_type,
                    source_order_no=order.source.source_order_no,
                    file_name=order.source.file_name,
                    file_fingerprint=order.source.file_fingerprint,
                    sheet_name=order.source.sheet_name,
                    row_numbers=order.source.row_numbers,
                )
                order.items = [
                    ParsedItem(
                        source_product_code=source_code,
                        source_product_name="光明满元气鲜糯玉米（加送1根）",
                        source_spec=source_spec,
                        source_display_spec="8根+1根",
                        quantity=1,
                        source_line_no=2,
                    )
                ]
                rules = confirmed_rules()
                rules.pop("store_assignment_policy")
                rules["defaults"]["store"] = {"value": "sam", "confirmed": True}
                rules["item_expansion_policy"] = sam_rules["item_expansion_policy"]

                result = build_custom_import_lines(
                    [order],
                    catalog,
                    rules,
                    self.profile,
                    business_date=date(2026, 10, 15),
                    source_company_id="sam",
                    official_mapping_ready=True,
                )

                self.assertEqual(
                    result.status,
                    "ready",
                    [item.to_dict() for item in result.clarifications],
                )
                self.assertEqual(
                    [(line["商品代码"], line["规格代码"], line["数量"]) for line in result.lines],
                    expected,
                )

    def test_sam_direct_and_expanded_items_use_the_same_catalog_product_name(self) -> None:
        sam_rules = json.loads(
            (
                PROJECT_ROOT
                / "config"
                / "platform_rules"
                / "guanyi"
                / "sam_order_v1.json"
            ).read_text(encoding="utf-8")
        )
        catalog = ProductCatalog(
            [
                ProductRecord(
                    "JTTZE1",
                    "光明满元气 花糯玉米8棒家庭装",
                    "6974768564835",
                    "8袋/箱",
                    1210,
                    "上海尚舆商贸有限公司",
                ),
                ProductRecord(
                    "GMCN-1",
                    "光明满元气 花糯玉米单根装",
                    "6974768564798",
                    "50根/箱",
                    180,
                    "上海尚舆商贸有限公司",
                ),
            ],
            [
                {
                    "source_profile_id": "sam_order_v1",
                    "identifier_type": "product_code",
                    "source_value": "DIRECT-FLOWER",
                    "product_code": "JTTZE1",
                    "spec_code": "6974768564835",
                    "confirmed": True,
                }
            ],
        )
        catalog.item_expansion_rules = (
            {
                "rule_id": "sam-flower-bundle",
                "company_id": "sam",
                "source_profile_id": "sam_order_v1",
                "identifier_type": "product_code",
                "source_value": "BUNDLE-FLOWER",
                "source_spec": "8+1",
                "target_platform": "guanyi",
                "confirmed": True,
                "reuse_scope": "future",
                "components": [
                    {
                        "component_index": 1,
                        "product_code": "JTTZE1",
                        "spec_code": "6974768564835",
                        "quantity_multiplier": 1,
                    },
                    {
                        "component_index": 2,
                        "product_code": "GMCN-1",
                        "spec_code": "6974768564798",
                        "quantity_multiplier": 1,
                    },
                ],
            },
        )
        rules = confirmed_rules()
        rules.pop("store_assignment_policy")
        rules["defaults"]["store"] = {"value": "sam", "confirmed": True}
        rules["field_output_policy"]["product_name"] = sam_rules[
            "field_output_policy"
        ]["product_name"]
        rules["item_expansion_policy"] = sam_rules["item_expansion_policy"]

        direct_order = parsed_order()
        direct_order.source = ParsedSource(
            profile_id="sam_order_v1",
            source_label="SAM群接龙订单",
            order_type=direct_order.source.order_type,
            source_order_no="SAM-DIRECT",
            file_name=direct_order.source.file_name,
            file_fingerprint=direct_order.source.file_fingerprint,
            sheet_name=direct_order.source.sheet_name,
            row_numbers=direct_order.source.row_numbers,
        )
        direct_order.items = [
            ParsedItem(
                source_product_code="DIRECT-FLOWER",
                source_product_name="来源表里的花糯8根名称",
                source_display_spec="花糯8根",
                quantity=1,
                source_line_no=2,
            )
        ]
        expanded_order = parsed_order()
        expanded_order.source = ParsedSource(
            profile_id="sam_order_v1",
            source_label="SAM群接龙订单",
            order_type=expanded_order.source.order_type,
            source_order_no="SAM-BUNDLE",
            file_name=expanded_order.source.file_name,
            file_fingerprint=expanded_order.source.file_fingerprint,
            sheet_name=expanded_order.source.sheet_name,
            row_numbers=expanded_order.source.row_numbers,
        )
        expanded_order.items = [
            ParsedItem(
                source_product_code="BUNDLE-FLOWER",
                source_product_name="来源表里的花糯8+1名称",
                source_spec="8+1",
                source_display_spec="花糯8根+1根",
                quantity=1,
                source_line_no=2,
            )
        ]

        direct_result = build_custom_import_lines(
            [direct_order],
            catalog,
            rules,
            self.profile,
            source_company_id="sam",
            official_mapping_ready=True,
        )
        expanded_result = build_custom_import_lines(
            [expanded_order],
            catalog,
            rules,
            self.profile,
            source_company_id="sam",
            official_mapping_ready=True,
        )

        self.assertEqual(direct_result.status, "ready")
        self.assertEqual(expanded_result.status, "ready")
        self.assertEqual(
            direct_result.lines[0]["商品名称"],
            "光明满元气 花糯玉米8棒家庭装",
        )
        expanded_flower = next(
            line
            for line in expanded_result.lines
            if line["商品代码"] == "JTTZE1"
            and line["规格代码"] == "6974768564835"
        )
        self.assertEqual(
            expanded_flower["商品名称"],
            direct_result.lines[0]["商品名称"],
        )

    def test_missing_official_mapping_blocks_before_product_question(self) -> None:
        result = build_custom_import_lines(
            [parsed_order()],
            self.catalog,
            confirmed_rules(),
            self.profile,
            official_mapping_ready=False,
            official_mapping_error="飞书用户未授权",
        )

        codes = {request.code for request in result.clarifications}
        self.assertEqual(result.status, "needs_input")
        self.assertEqual(result.lines, [])
        self.assertIn("FEISHU_OFFICIAL_MAPPING_UNAVAILABLE", codes)
        self.assertNotIn("PROVIDE_GUANYI_PRODUCT_MAPPING", codes)
        self.assertNotIn("CONFIRM_GUANYI_PRODUCT_MAPPING", codes)
        request = next(
            item
            for item in result.clarifications
            if item.code == "FEISHU_OFFICIAL_MAPPING_UNAVAILABLE"
        )
        self.assertEqual(request.reason, "飞书用户未授权")

    def test_sam_standard_item_uses_readable_source_name_and_display_spec(self) -> None:
        order = parsed_order()
        order.items = [
            ParsedItem(
                source_product_code="EXT-1",
                source_product_name="SHOWA丝滑手感中厚型手套",
                source_spec="SWSH8072",
                source_display_spec="S/常规款/珍珠粉",
                quantity=1,
                source_line_no=2,
            )
        ]
        rules = confirmed_rules()
        rules["field_output_policy"]["product_name"]["strategy"] = (
            "source_name_with_display_spec_for_standard_items"
        )

        result = build_custom_import_lines(
            [order],
            self.catalog,
            rules,
            self.profile,
        )

        self.assertEqual(result.status, "ready")
        self.assertEqual(
            result.lines[0]["商品名称"],
            "SHOWA丝滑手感中厚型手套/S/常规款/珍珠粉",
        )

    def test_customer_note_is_written_to_seller_remark(self) -> None:
        order = parsed_order()
        order.source_note = "抓紧发货"
        result = build_custom_import_lines(
            [order],
            self.catalog,
            confirmed_rules(),
            self.profile,
        )
        self.assertEqual(result.status, "ready")
        self.assertEqual(result.lines[0]["卖家备注"], "抓紧发货")

    def test_courier_instruction_is_appended_to_recipient_address(self) -> None:
        order = parsed_order()
        order.source_note = "送货上门，不准放门口"
        result = build_custom_import_lines(
            [order],
            self.catalog,
            confirmed_rules(),
            self.profile,
        )
        self.assertEqual(result.status, "ready")
        self.assertEqual(
            result.lines[0]["收货地址"],
            "上海市宝山区测试路1号（送货上门；不准放门口）",
        )
        self.assertEqual(result.lines[0]["卖家备注"], "")

    def test_mixed_notes_are_split_between_address_and_seller_remark(self) -> None:
        order = parsed_order()
        order.source_note = "抓紧发货；不准放门口"
        result = build_custom_import_lines(
            [order],
            self.catalog,
            confirmed_rules(),
            self.profile,
        )
        self.assertEqual(result.status, "ready")
        self.assertEqual(
            result.lines[0]["收货地址"],
            "上海市宝山区测试路1号（不准放门口）",
        )
        self.assertEqual(result.lines[0]["卖家备注"], "抓紧发货")

    def test_address_instruction_is_not_duplicated(self) -> None:
        order = parsed_order()
        order.recipient = ParsedRecipient(
            name="测试用户",
            mobile="13800138000",
            province="上海",
            city="上海市",
            district="宝山区",
            raw_address="上海市宝山区测试路1号（不准放门口）",
        )
        order.source_note = "不准放门口"
        result = build_custom_import_lines(
            [order],
            self.catalog,
            confirmed_rules(),
            self.profile,
        )
        self.assertEqual(result.status, "ready")
        self.assertEqual(
            result.lines[0]["收货地址"],
            "上海市宝山区测试路1号（不准放门口）",
        )
        self.assertEqual(result.lines[0]["卖家备注"], "")

    def test_explicit_delivery_instruction_extension_is_supported(self) -> None:
        order = parsed_order()
        order.source_note = "抓紧发货"
        order.source_extensions["delivery_instruction"] = "送货上门"
        result = build_custom_import_lines(
            [order],
            self.catalog,
            confirmed_rules(),
            self.profile,
        )
        self.assertEqual(result.status, "ready")
        self.assertEqual(
            result.lines[0]["收货地址"],
            "上海市宝山区测试路1号（送货上门）",
        )
        self.assertEqual(result.lines[0]["卖家备注"], "抓紧发货")

    def test_seller_remark_separator_is_read_from_configuration(self) -> None:
        order = parsed_order()
        order.source_note = "抓紧发货"
        rules = confirmed_rules()
        rules["seller_remark_policy"]["separator"] = "/"
        result = build_custom_import_lines(
            [order],
            self.catalog,
            rules,
            self.profile,
        )
        self.assertEqual(result.status, "ready")
        self.assertEqual(result.lines[0]["卖家备注"], "抓紧发货")

    def test_unconfirmed_field_output_policy_blocks_build(self) -> None:
        rules = confirmed_rules()
        rules["field_output_policy"]["order_created_at"]["confirmed"] = False
        result = build_custom_import_lines(
            [parsed_order()],
            self.catalog,
            rules,
            self.profile,
        )
        self.assertEqual(result.status, "needs_input")
        self.assertEqual(result.lines, [])
        self.assertIn(
            "CONFIRM_GUANYI_FIELD_POLICY_ORDER_CREATED_AT",
            {request.code for request in result.clarifications},
        )

    def test_unmapped_product_store_blocks_build(self) -> None:
        rules = confirmed_rules()
        rules["store_assignment_policy"]["rules"][0]["match"] = {
            "product_codes": ["OTHER"]
        }
        result = build_custom_import_lines(
            [parsed_order()],
            self.catalog,
            rules,
            self.profile,
        )
        self.assertEqual(result.status, "needs_input")
        self.assertEqual(result.lines, [])
        self.assertIn(
            "CONFIRM_GUANYI_STORE_ASSIGNMENT",
            {request.code for request in result.clarifications},
        )

    def test_unconfirmed_required_rules_block_build(self) -> None:
        rules = confirmed_rules()
        rules["defaults"]["payment_amount"]["confirmed"] = False
        result = build_custom_import_lines(
            [parsed_order()],
            self.catalog,
            rules,
            self.profile,
        )
        self.assertEqual(result.status, "needs_input")
        self.assertEqual(result.lines, [])
        self.assertIn(
            "CONFIRM_GUANYI_PAYMENT_AMOUNT",
            {request.code for request in result.clarifications},
        )

    def test_renderer_removes_examples_and_preserves_template_headers(self) -> None:
        built = build_custom_import_lines(
            [parsed_order()],
            self.catalog,
            confirmed_rules(),
            self.profile,
        )
        built.lines[0]["规格代码"] = "4548404200030"
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "output.xlsx"
            render_custom_import(
                PROJECT_ROOT / "assets/templates/guanyi/自定义订单导入模板.xlsx",
                output,
                built.lines,
                self.profile,
            )
            workbook = load_workbook(output, read_only=False, data_only=True)
            sheet = workbook["Sheet1"]
            headers = [cell.value for cell in sheet[1] if cell.value is not None]
            self.assertEqual(headers, self.profile["columns"])
            self.assertEqual(sheet.max_row, 2)
            self.assertEqual(sheet["B2"].value, "SRC-001")
            self.assertEqual(sheet["F2"].value, "P1")
            self.assertEqual(sheet["G2"].value, "4548404200030")
            self.assertEqual(sheet["B2"].number_format, "@")
            self.assertEqual(sheet["G2"].number_format, "@")
            self.assertEqual(sheet["P2"].value, "13800138000")
            self.assertEqual(sheet["Q2"].value, "13800138000")
            self.assertEqual(sheet["AA2"].value, "中通快递（重货）")
            workbook.close()


if __name__ == "__main__":
    unittest.main()
