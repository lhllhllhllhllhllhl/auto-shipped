from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from auto_shipped.companies import load_company_registry
from auto_shipped.services.adaptive import propose_adaptive_excel_plan, write_adaptive_plan
from auto_shipped.services.batch_convert import convert_order_batch
from auto_shipped.services.convert import convert_order_file
from auto_shipped.source_adapters import (
    ADAPTIVE_ASSUMPTION_CONFIRMATION,
    AdaptiveExcelParsedAdapter,
    inspect_excel_structure,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]
NDD_PROFILE = json.loads(
    (PROJECT_ROOT / "config/source_profiles/ndd_order_v1.json").read_text(
        encoding="utf-8"
    )
)


def build_manual_address_list(path: Path, *, marker: str = "") -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(
        ["序号", "", "收件人", "收件人电话", "镇街道", "详细地址", "订单号", "备注"]
    )
    sheet.append(
        [1, 1, f"测试甲{marker}", "13800138000", "张江镇", "上海市浦东新区测试路1号", "", "尽快发货"]
    )
    sheet.append(
        [2, 2, f"测试乙{marker}", "13900139000", "周浦镇", "上海市浦东新区测试路2号", "", ""]
    )
    workbook.save(path)


class AdaptiveExcelTests(unittest.TestCase):
    def test_structure_inspection_never_returns_order_values(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "manual.xlsx"
            build_manual_address_list(source)

            structure = inspect_excel_structure(source)
            rendered = json.dumps(structure, ensure_ascii=False)

            self.assertEqual(structure["privacy"], "headers_and_type_counts_only_no_data_values")
            self.assertNotIn("测试甲", rendered)
            self.assertNotIn("13800138000", rendered)
            self.assertNotIn("测试路1号", rendered)
            headers = [
                item["header"] for item in structure["sheets"][0]["columns"]
            ]
            self.assertEqual(
                headers,
                ["序号", "", "收件人", "收件人电话", "镇街道", "详细地址", "订单号", "备注"],
            )

    def test_proposal_maps_recipient_fields_and_asks_only_missing_business_values(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "manual.xlsx"
            build_manual_address_list(source)
            registry = load_company_registry(
                PROJECT_ROOT / "config/companies/company_registry_v1.json"
            )

            proposal = propose_adaptive_excel_plan(source, "ndd", registry)

            self.assertEqual(proposal.status, "needs_input")
            self.assertIsNotNone(proposal.plan)
            assert proposal.plan is not None
            mapping = proposal.plan["field_mapping"]
            self.assertEqual(mapping["recipient_name"]["columns"], ["C"])
            self.assertEqual(mapping["contact"]["columns"], ["D"])
            self.assertEqual(mapping["address"]["columns"], ["E", "F"])
            self.assertEqual(mapping["source_note"]["columns"], ["H"])
            self.assertEqual(
                set(proposal.plan["unresolved_fields"]),
                {"source_order_no", "product_identity", "quantity"},
            )
            self.assertEqual(
                {item.code for item in proposal.clarifications},
                {
                    "ADAPTIVE_ORDER_NUMBER_REQUIRED",
                    "ADAPTIVE_PRODUCT_REQUIRED",
                    "ADAPTIVE_QUANTITY_REQUIRED",
                },
            )

    def test_confirmed_adaptive_plan_parses_without_fixed_source_headers(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "manual.xlsx"
            build_manual_address_list(source)
            registry = load_company_registry(
                PROJECT_ROOT / "config/companies/company_registry_v1.json"
            )
            proposal = propose_adaptive_excel_plan(source, "ndd", registry)
            assert proposal.plan is not None
            plan = proposal.plan
            plan["batch_defaults"] = {
                "product_name": "光明 福利套餐四",
                "quantity": 1,
            }
            plan["order_number"] = {
                "strategy": "date_sequence",
                "business_date": "2026-10-08",
                "sequence_start": 1,
                "sequence_width": 3,
            }
            plan["assumption_confirmation_token"] = ADAPTIVE_ASSUMPTION_CONFIRMATION

            parsed = AdaptiveExcelParsedAdapter().parse(source, NDD_PROFILE, plan)

            self.assertEqual(parsed.status, "parsed")
            self.assertEqual(len(parsed.orders), 2)
            self.assertEqual(
                [order.source.source_order_no for order in parsed.orders],
                ["261008001", "261008002"],
            )
            self.assertEqual(
                {order.items[0].source_product_name for order in parsed.orders},
                {"光明 福利套餐四"},
            )
            self.assertEqual(
                {order.items[0].quantity for order in parsed.orders},
                {1.0},
            )

    def test_company_hint_creates_plan_instead_of_rejecting_unknown_template(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "manual.xlsx"
            output = root / "output"
            build_manual_address_list(source)

            result = convert_order_file(
                source,
                PROJECT_ROOT / "assets/catalog/current_product_catalog.csv",
                output,
                company_hint="ndd",
            )

            self.assertEqual(result.status, "needs_input")
            self.assertIsNotNone(result.detection)
            assert result.detection is not None
            self.assertEqual(result.detection.company_id, "ndd")
            self.assertEqual(result.detection.source_profile_id, "ndd_order_v1")
            self.assertEqual(result.detection.confidence, "medium")
            codes = {item.code for item in result.clarifications}
            self.assertEqual(
                codes,
                {
                    "ADAPTIVE_ORDER_NUMBER_REQUIRED",
                    "ADAPTIVE_PRODUCT_REQUIRED",
                    "ADAPTIVE_QUANTITY_REQUIRED",
                },
            )
            plans = list(output.glob("*_自适应来源计划.json"))
            self.assertEqual(len(plans), 1)

    def test_confirmed_plan_runs_through_ndd_rules_and_renders_guanyi_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "manual.xlsx"
            output = root / "output"
            plan_path = root / "plan.json"
            build_manual_address_list(source)
            registry = load_company_registry(
                PROJECT_ROOT / "config/companies/company_registry_v1.json"
            )
            proposal = propose_adaptive_excel_plan(source, "ndd", registry)
            assert proposal.plan is not None
            plan = proposal.plan
            plan["batch_defaults"] = {
                "product_name": "光明 福利套餐四",
                "quantity": 1,
            }
            plan["order_number"] = {
                "strategy": "date_sequence",
                "business_date": "2026-10-08",
                "sequence_start": 1,
                "sequence_width": 3,
            }
            plan["assumption_confirmation_token"] = ADAPTIVE_ASSUMPTION_CONFIRMATION
            write_adaptive_plan(plan_path, plan)

            result = convert_order_file(
                source,
                PROJECT_ROOT / "assets/catalog/current_product_catalog.csv",
                output,
                company_hint="ndd",
                adaptive_plan_path=plan_path,
                official_mapping_ready=False,
            )

            self.assertEqual(result.status, "ready")
            self.assertEqual(result.parsed_order_count, 2)
            self.assertEqual(result.output_row_count, 4)
            workbook_path = next(
                Path(value) for value in result.outputs if value.endswith(".xlsx")
            )
            workbook = load_workbook(workbook_path, data_only=False)
            try:
                sheet = workbook["Sheet1"]
                platform_numbers = {
                    sheet.cell(row_no, 2).value
                    for row_no in range(2, sheet.max_row + 1)
                }
            finally:
                workbook.close()
            self.assertEqual(
                platform_numbers,
                {"NDD261008001A", "NDD261008002A"},
            )

    def test_batch_merge_accepts_adaptive_plans_per_source(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            registry = load_company_registry(
                PROJECT_ROOT / "config/companies/company_registry_v1.json"
            )
            sources = [root / "first.xlsx", root / "second.xlsx"]
            build_manual_address_list(sources[0], marker="一")
            build_manual_address_list(sources[1], marker="二")
            plan_paths: dict[str, str] = {}
            for source, prefix in zip(sources, ("X", "Y"), strict=True):
                proposal = propose_adaptive_excel_plan(source, "ndd", registry)
                assert proposal.plan is not None
                plan = proposal.plan
                plan["batch_defaults"] = {
                    "product_name": "光明 福利套餐四",
                    "quantity": 1,
                }
                plan["order_number"] = {
                    "strategy": "date_sequence",
                    "business_date": "2026-10-08",
                    "prefix": prefix,
                    "sequence_start": 1,
                    "sequence_width": 3,
                }
                plan["assumption_confirmation_token"] = (
                    ADAPTIVE_ASSUMPTION_CONFIRMATION
                )
                plan_path = root / f"{source.stem}.plan.json"
                write_adaptive_plan(plan_path, plan)
                plan_paths[source.name] = str(plan_path)

            result = convert_order_batch(
                sources,
                PROJECT_ROOT / "assets/catalog/current_product_catalog.csv",
                root / "output",
                adaptive_plan_hints=plan_paths,
                official_mapping_ready=False,
                batch_name="adaptive-batch",
            )

            self.assertEqual(result.status, "ready")
            self.assertEqual(result.order_count, 4)
            self.assertEqual(result.item_row_count, 8)
            self.assertEqual(len(result.outputs), 2)


if __name__ == "__main__":
    unittest.main()
