from __future__ import annotations

import json
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path

from openpyxl import load_workbook

from auto_shipped.platforms.guanyi import (
    preflight_custom_import,
    render_custom_import,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def profile() -> dict:
    return json.loads(
        (PROJECT_ROOT / "config/platform_profiles/guanyi_order_import_v1.json").read_text(
            encoding="utf-8"
        )
    )


def valid_line() -> dict:
    return {
        "店铺": "光明满元气",
        "平台单号": "SRC-001A",
        "买家会员": "张",
        "支付金额": 0,
        "商品名称": "管易测试商品",
        "商品代码": "P1",
        "规格代码": "S1",
        "规格名称": "测试规格",
        "是否赠品": 0,
        "数量": 2,
        "价格": 0,
        "运费": 0,
        "收货人": "测试用户",
        "联系手机": "13800138000",
        "收货地址": "上海市宝山区测试路1号",
        "省": "上海",
        "市": "上海市",
        "区": "宝山区",
        "订单创建时间": "2026-09-20 14:09:15",
        "物流公司": "韵达快递",
        "是否手机订单": 0,
        "是否货到付款": 0,
        "支付方式": "网银在线",
        "仓库名称": "上海尚舆商贸有限公司",
        "订单类型": "销售订单",
        "是否分销商订单": 0,
    }


class GuanyiPreflightTests(unittest.TestCase):
    def render(self, root: Path, lines: list[dict] | None = None) -> Path:
        output = root / "output.xlsx"
        render_custom_import(
            PROJECT_ROOT / "assets/templates/guanyi/自定义订单导入模板.xlsx",
            output,
            lines or [valid_line()],
            profile(),
        )
        return output

    def test_valid_workbook_produces_pii_free_ready_manifest(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            output = self.render(Path(tmp))
            result = preflight_custom_import(output, profile())
            self.assertEqual(result.status, "ready")
            manifest = result.to_manifest()
            self.assertEqual(manifest["counts"], {"orders": 1, "item_rows": 1})
            self.assertTrue(all(manifest["checks"].values()))
            rendered = json.dumps(manifest, ensure_ascii=False)
            self.assertNotIn("测试用户", rendered)
            self.assertNotIn("13800138000", rendered)
            self.assertNotIn("测试路1号", rendered)
            self.assertFalse(manifest["safety"]["upload_authorized"])

    def test_optional_freight_can_be_blank_but_cannot_be_negative(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            line = valid_line()
            line["运费"] = ""
            output = self.render(Path(tmp), [line])
            ready = preflight_custom_import(output, profile())
            self.assertEqual(ready.status, "ready")

            workbook = load_workbook(output)
            workbook["Sheet1"]["M2"] = -1
            workbook.save(output)
            workbook.close()
            blocked = preflight_custom_import(output, profile())
            self.assertEqual(blocked.status, "blocked")
            self.assertIn(
                "OPTIONAL_NONNEGATIVE_NUMBER_INVALID",
                {issue.code for issue in blocked.issues},
            )

    def test_missing_required_value_blocks_upload_candidate(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            output = self.render(Path(tmp))
            workbook = load_workbook(output)
            workbook["Sheet1"]["O2"] = ""
            workbook.save(output)
            workbook.close()

            result = preflight_custom_import(output, profile())
            self.assertEqual(result.status, "blocked")
            self.assertIn(
                "REQUIRED_VALUE_MISSING",
                {issue.code for issue in result.issues},
            )

    def test_formula_blocks_upload_candidate(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            output = self.render(Path(tmp))
            workbook = load_workbook(output)
            workbook["Sheet1"]["F2"] = "=1+1"
            workbook.save(output)
            workbook.close()

            result = preflight_custom_import(output, profile())
            self.assertEqual(result.status, "blocked")
            self.assertIn("FORMULA_NOT_ALLOWED", {issue.code for issue in result.issues})

    def test_invalid_contact_blocks_upload_candidate(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            line = valid_line()
            line["联系手机"] = "abc"
            output = self.render(Path(tmp), [line])
            result = preflight_custom_import(output, profile())
            self.assertEqual(result.status, "blocked")
            self.assertIn("FIELD_PATTERN_INVALID", {issue.code for issue in result.issues})

    def test_platform_hidden_mobile_suffix_is_preserved_and_allowed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            line = valid_line()
            line["联系电话"] = "13800138000-6762"
            line["联系手机"] = "13800138000-6762"
            output = self.render(Path(tmp), [line])
            result = preflight_custom_import(output, profile())
            self.assertEqual(result.status, "ready")

    def test_unknown_carrier_blocks_upload_candidate(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            line = valid_line()
            line["物流公司"] = "随便快递"
            output = self.render(Path(tmp), [line])
            result = preflight_custom_import(output, profile())
            self.assertEqual(result.status, "blocked")
            self.assertIn("FIELD_VALUE_NOT_ALLOWED", {issue.code for issue in result.issues})

    def test_jd_zhongtong_is_an_allowed_guanyi_carrier_value(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            line = valid_line()
            line["物流公司"] = "京东中通"
            output = self.render(Path(tmp), [line])
            result = preflight_custom_import(output, profile())
            self.assertEqual(result.status, "ready")

    def test_missing_automation_suffix_blocks_upload_candidate(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            line = valid_line()
            line["平台单号"] = "SRC-001"
            output = self.render(Path(tmp), [line])
            result = preflight_custom_import(output, profile())
            self.assertEqual(result.status, "blocked")
            self.assertIn(
                "REQUIRED_TEXT_SUFFIX_MISSING",
                {issue.code for issue in result.issues},
            )

    def test_inconsistent_order_fields_block_multi_line_order(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            first = valid_line()
            second = deepcopy(first)
            second["商品代码"] = "P2"
            second["规格代码"] = "S2"
            second["收货地址"] = "上海市宝山区另一地址"
            output = self.render(Path(tmp), [first, second])

            result = preflight_custom_import(output, profile())
            self.assertEqual(result.status, "blocked")
            self.assertIn(
                "ORDER_FIELD_INCONSISTENT",
                {issue.code for issue in result.issues},
            )

    def test_corrupt_xlsx_is_reported_as_blocked(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "broken.xlsx"
            output.write_bytes(b"not an xlsx zip")
            result = preflight_custom_import(output, profile())
            self.assertEqual(result.status, "blocked")
            self.assertIn("WORKBOOK_UNREADABLE", {issue.code for issue in result.issues})


if __name__ == "__main__":
    unittest.main()
